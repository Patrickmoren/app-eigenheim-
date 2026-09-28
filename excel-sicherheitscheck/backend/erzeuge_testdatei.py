# -*- coding: utf-8 -*-
"""Erzeugt eine synthetische .xlsx-Datei mit bewusst eingebauten Fehlern,
um den Analyzer zu testen. Bildet typische reale Probleme nach (siehe
docs/01-analyse-und-konzept.md in diesem Repo als Vorbild)."""
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Protection

wb = openpyxl.Workbook()

# ---- Blatt "Daten": Hauptdaten mit 300 Zeilen ----
ws = wb.active
ws.title = "Daten"
kopf = ["ID", "Name", "Datum", "Betrag", "Berechnet", "Status"]
for i, k in enumerate(kopf, start=1):
    ws.cell(1, i, k)

for r in range(2, 302):
    ws.cell(r, 1, r - 1)
    ws.cell(r, 2, f"Fall {r-1}")
    if r % 37 == 0:
        # Datum als Text statt echtes Datum (typischer Fehler)
        ws.cell(r, 3, "12.5.2026")
    else:
        ws.cell(r, 3, f"={r}.5.2026" if False else __import__("datetime").date(2026, 5, 1))
    ws.cell(r, 4, 100 + r)
    if r == 250:
        ws.cell(r, 5, 999)  # von Hand überschriebene Formel
    elif r == 180:
        ws.cell(r, 5, f"=D{r}*2+1")  # abweichende Formel (Inkonsistenz)
    else:
        ws.cell(r, 5, f"=D{r}*2")
    ws.cell(r, 6, "Ja" if r % 4 else "Nein")

# Datenprüfung, die nur bis Zeile 150 reicht, obwohl Daten bis 301 gehen
dv = DataValidation(type="list", formula1='"Ja,Nein"', showErrorMessage=True)
ws.add_data_validation(dv)
dv.add(f"F2:F150")

# ---- Blatt "Auswertung": bezieht sich fest auf Daten!A2:A150 (Datenverlust) ----
ws2 = wb.create_sheet("Auswertung")
ws2["A1"] = "Kennzahl"
ws2["B1"] = "Wert"
ws2["A2"] = "Anzahl Fälle"
ws2["B2"] = "=COUNTA(Daten!A2:A150)"
ws2["A3"] = "Summe Betrag"
ws2["B3"] = "=SUM(Daten!D2:D150)"
ws2["A4"] = "Heute"
ws2["B4"] = "=HEUTE()"

# ---- Blatt "Reserve": verstecktes, unreferenziertes Blatt mit Daten ----
ws3 = wb.create_sheet("Reserve")
ws3["A1"] = "Bewirtschafter"
ws3["A2"] = "Muster AG"
ws3["A3"] = "Beispiel GmbH"
ws3.sheet_state = "hidden"

# ---- Zirkelbezug ----
ws4 = wb.create_sheet("Zirkel")
ws4["A1"] = "=B1+1"
ws4["B1"] = "=A1+1"

# Blattschutz auf "Auswertung" setzen, aber B2 als Eingabe offen lassen
ws2.protection.sheet = True
for zelle in ["B2"]:
    ws2[zelle].protection = Protection(locked=False)

wb.save("Testdatei_kaputt.xlsx")
print("Testdatei_kaputt.xlsx erzeugt.")
