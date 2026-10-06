# -*- coding: utf-8 -*-
"""Erzeugt das «Vermieter-Toolkit Schweiz» als Excel-Arbeitsmappe.

Zielgruppe: private Vermieterinnen und Vermieter mit 1–12 Wohnungen.
Inhalt: Nebenkostenabrechnung mit Heizgradtag-Gewichtung und Leerstandsanteil,
Mietzinsanpassungs-Rechner (Referenzzinssatz, Teuerung, Kostensteigerung),
Checkliste Mieterwechsel, Wohnungsabnahmeprotokoll mit Altersentwertung.

Keine Makros, nur Formeln – läuft in Excel, LibreOffice und Numbers.

    python3 build.py                    # leere Verkaufsversion
    BEISPIEL=1 python3 build.py         # mit Beispieldaten (Demo / Tests)
"""
import os
from datetime import date

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

BEISPIEL = os.environ.get("BEISPIEL") == "1"
DATEI = os.environ.get("ZIEL", "Vermieter-Toolkit-Schweiz-Beispiel.xlsx" if BEISPIEL
                       else "Vermieter-Toolkit-Schweiz.xlsx")

N_EINHEITEN = 12
N_MIETER = 30
N_KOSTEN = 20

# ---------------------------------------------------------------- Stil
TINTE, AKZENT, HELL = "1F2A30", "1E5F74", "E8F0F2"
EINGABE, WARN, OK = "FFF8E1", "FBE4E2", "E4F0E8"
RAND = "C5CFD3"
duenn = Side(style="thin", color=RAND)
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)


def f(size=10, bold=False, color=TINTE, italic=False):
    return Font(name="Arial", size=size, bold=bold, color=color, italic=italic)


def fill(c):
    return PatternFill("solid", fgColor=c)


def titel(ws, text, sub=None):
    ws["B2"] = text
    ws["B2"].font = f(16, True, AKZENT)
    if sub:
        ws["B3"] = sub
        ws["B3"].font = f(9, italic=True, color="5A6B72")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 2


def kopf(ws, row, col, texte):
    for i, t in enumerate(texte):
        c = ws.cell(row=row, column=col + i, value=t)
        c.font = f(9, True, "FFFFFF")
        c.fill = fill(AKZENT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = rahmen


def eingabe(c, fmt=None):
    c.fill = fill(EINGABE)
    c.border = rahmen
    c.protection = Protection(locked=False)
    c.font = f()
    if fmt:
        c.number_format = fmt


def formel(c, fmt=None, bold=False):
    c.border = rahmen
    c.font = f(bold=bold)
    if fmt:
        c.number_format = fmt


def liste(ws, rng, werte, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(werte) + '"', allow_blank=True)
    if prompt:
        dv.prompt, dv.showInputMessage = prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


CHF = '#,##0.00 "CHF"'
CHF0 = '#,##0 "CHF"'
DAT = "DD.MM.YYYY"
PROZ = "0.00%"

wb = openpyxl.Workbook()

# ================================================================ Start
ws = wb.active
ws.title = "Start"
titel(ws, "Vermieter-Toolkit Schweiz", "Nebenkosten · Mietzinsanpassung · Mieterwechsel · Wohnungsabnahme")
ws.column_dimensions["B"].width = 100
texte = [
    ("So gehen Sie vor", True),
    ("1.  «Liegenschaft»: Abrechnungsperiode, Wohnungen und Verteilschlüssel erfassen (gelbe Felder).", False),
    ("2.  «Mieter»: alle Mietverhältnisse der Periode erfassen – auch Mieterwechsel unterjährig. Leere Daten = ganze Periode.", False),
    ("3.  «Kosten»: Rechnungsbeträge erfassen, Art und Schlüssel wählen. Nur im Mietvertrag ausgeschiedene Nebenkosten werden belastet.", False),
    ("4.  «Abrechnung»: Mieter-Nr. wählen, prüfen, drucken oder als PDF speichern. Leerstandskosten trägt automatisch der Vermieter.", False),
    ("5.  «Mietzins»: Anpassung nach Referenzzinssatz, Teuerung und Kostensteigerung berechnen – inkl. spätestem Zustelltermin.", False),
    ("6.  «Mieterwechsel» und «Abnahme»: Checkliste mit Terminen und druckbares Abnahmeprotokoll.", False),
    ("", False),
    ("Was dieses Toolkit automatisch berücksichtigt", True),
    ("•  Heiz- und Warmwasserkosten nach Heizgradtagen gewichtet (Mieterwechsel im Sommer ≠ im Winter). Skala anpassbar.", False),
    ("•  Andere Nebenkosten zeitanteilig nach Tagen.", False),
    ("•  Kosten leerstehender Wohnungen bleiben beim Vermieter und werden separat ausgewiesen.", False),
    ("•  Nicht im Mietvertrag ausgeschiedene oder nicht umlagefähige Positionen werden nicht belastet (Art. 257a OR).", False),
    ("•  Kontrollsumme: verteilte Kosten + Vermieteranteil = abrechenbare Gesamtkosten.", False),
    ("", False),
    ("Wichtig", True),
    ("Das Toolkit ist ein Rechenhilfsmittel und ersetzt keine Rechtsberatung. Massgebend sind Mietvertrag, OR und VMWG sowie die", False),
    ("kantonale Praxis. Mietzinserhöhungen müssen mit dem amtlichen kantonalen Formular mitgeteilt werden (Art. 269d OR).", False),
    ("Gelbe Felder sind Eingaben, alle anderen Felder rechnen automatisch. Blattschutz ohne Passwort – bei Bedarf: Überprüfen › Blatt freigeben.", False),
]
for i, (t, b) in enumerate(texte):
    c = ws.cell(row=5 + i, column=2, value=t)
    c.font = f(11 if b else 10, b, AKZENT if b else TINTE)
ws.cell(row=27, column=2, value="Version 1.0 · Stand Oktober 2026").font = f(8, italic=True, color="5A6B72")

# ================================================================ Liegenschaft
L = wb.create_sheet("Liegenschaft")
titel(L, "Liegenschaft & Abrechnungsperiode", "Gelbe Felder ausfüllen. Die Periode muss 12 Monate umfassen und am Monatsersten beginnen.")
for col, w in zip("BCDEFGHIJKLMNO", [30, 16, 10, 12, 12, 11, 3, 12, 13, 3, 6, 12, 12, 10]):
    L.column_dimensions[col].width = w
stamm = [
    (4, "Liegenschaft", None, "Musterweg 12" if BEISPIEL else None),
    (5, "Ort", None, "8000 Zürich" if BEISPIEL else None),
    (6, "Vermieter/in", None, "Anna Beispiel" if BEISPIEL else None),
    (7, "Periode von", DAT, date(2025, 7, 1) if BEISPIEL else None),
    (8, "Periode bis", DAT, date(2026, 6, 30) if BEISPIEL else None),
    (10, "Verwaltungshonorar in % der Nebenkosten", PROZ, 0.03),
    (11, "Honorar im Mietvertrag vereinbart?", None, "Ja" if BEISPIEL else "Nein"),
]
for r, label, fmt, wert in stamm:
    L.cell(row=r, column=2, value=label).font = f()
    c = L.cell(row=r, column=3, value=wert)
    eingabe(c, fmt)
L["B9"] = "Tage in Periode"
L["C9"] = '=IF(OR(C7="",C8=""),"",C8-C7+1)'
formel(L["C9"], "0")
liste(L, "C11", ["Ja", "Nein"])
L["D8"] = '=IF(OR(C7="",C8=""),"",IF(AND(DAY(C7)=1,C8=DATE(YEAR(C7),MONTH(C7)+12,0)),"✓ 12 Monate","⚠ Periode muss 12 volle Monate ab Monatserstem umfassen"))'
L["D8"].font = f(9, True)

kopf(L, 13, 2, ["Wohnung / Einheit", "Fläche m²", "Zimmer", "Wertquote ‰", "Volumen m³", "Einheiten gleich"])
bsp_einheiten = [("EG links", 68, 3, 180, 170), ("EG rechts", 82, 3.5, 210, 205), ("1. OG links", 68, 3, 180, 170),
                 ("1. OG rechts", 82, 3.5, 210, 205), ("DG", 95, 4.5, 220, 260)]
for i in range(N_EINHEITEN):
    r = 14 + i
    vals = bsp_einheiten[i] if BEISPIEL and i < len(bsp_einheiten) else (None,) * 5
    for j, v in enumerate(vals):
        eingabe(L.cell(row=r, column=2 + j, value=v), None if j == 0 else "#,##0.0")
    L.cell(row=r, column=7, value=f'=IF(B{r}="",0,1)')
    formel(L.cell(row=r, column=7), "0")
R_EINH_ERST, R_EINH_LETZT, R_EINH_SUM = 14, 13 + N_EINHEITEN, 14 + N_EINHEITEN   # 14..25, Summe 26
L.cell(row=R_EINH_SUM, column=2, value="Total").font = f(bold=True)
for col in "CDEFG":
    c = L[f"{col}{R_EINH_SUM}"]
    c.value = f"=SUM({col}{R_EINH_ERST}:{col}{R_EINH_LETZT})"
    formel(c, "#,##0.0", True)

# Heizgradtag-Skala (übliche Monatsverteilung, Summe 100 – anpassbar)
kopf(L, 13, 9, ["Monat", "Heizgradtage %"])
hgt = [18, 15, 13, 9, 4, 0, 0, 0, 3, 8, 13, 17]
monate = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September", "Oktober", "November", "Dezember"]
for i in range(12):
    L.cell(row=14 + i, column=9, value=monate[i]).border = rahmen
    eingabe(L.cell(row=14 + i, column=10, value=hgt[i]), "0.0")
L["I26"] = "Summe"
L["J26"] = "=SUM(J14:J25)"
formel(L["J26"], "0.0", True)
L["I27"] = '=IF(ABS(J26-100)<0.01,"✓","⚠ Summe ≠ 100")'
L["I27"].font = f(9, True)

# Hilfstabelle: Monate der Periode
kopf(L, 13, 12, ["Nr", "Beginn", "Ende", "Gewicht"])
for i in range(12):
    r = 14 + i
    L.cell(row=r, column=12, value=i + 1)
    L.cell(row=r, column=13, value=f'=IF($C$7="","",DATE(YEAR($C$7),MONTH($C$7)+{i},1))').number_format = DAT
    L.cell(row=r, column=14, value=f'=IF($C$7="","",DATE(YEAR($C$7),MONTH($C$7)+{i + 1},0))').number_format = DAT
    L.cell(row=r, column=15, value=f'=IF($C$7="",0,INDEX($J$14:$J$25,MONTH(M{r})))').number_format = "0.0"
    for col in range(12, 16):
        L.cell(row=r, column=col).font = f(8, color="7A8A90")
L["L12"] = "Hilfstabelle (nicht bearbeiten)"
L["L12"].font = f(8, italic=True, color="7A8A90")

# ================================================================ Mieter
M = wb.create_sheet("Mieter")
titel(M, "Mietverhältnisse in der Periode", "Eine Zeile pro Mietverhältnis. Mietbeginn/-ende leer lassen, wenn vor bzw. nach der Periode.")
kopf(M, 5, 2, ["Nr", "Wohnung", "Mieter/in", "Mietbeginn", "Mietende", "Akonto bezahlt", "von (eff.)", "bis (eff.)",
               "Tage", "Zeitanteil", "Heizanteil", "Prüfung"])
for col, w in zip("BCDEFGHIJKLM", [5, 16, 24, 12, 12, 13, 11, 11, 7, 10, 10, 26]):
    M.column_dimensions[col].width = w
bsp_mieter = [
    ("EG links", "Muster Peter", None, None, 2400),
    ("EG rechts", "Keller Sara", None, date(2025, 12, 31), 1350),
    ("EG rechts", "Rossi Marco", date(2026, 2, 1), None, 1000),
    ("1. OG links", "Huber Lea", None, None, 2400),
    ("1. OG rechts", "Meier Jonas", None, None, 2700),
    ("DG", "Brunner Eva", None, None, 3000),
]
HCOL0 = 15  # Spalte O: 12 Hilfsspalten Heizgewichtung
for i in range(N_MIETER):
    r = 6 + i
    M.cell(row=r, column=2, value=i + 1).font = f(9, color="5A6B72")
    vals = bsp_mieter[i] if BEISPIEL and i < len(bsp_mieter) else (None,) * 5
    for j, v in enumerate(vals):
        eingabe(M.cell(row=r, column=3 + j, value=v), [None, None, DAT, DAT, CHF][j])
    M[f"H{r}"] = f'=IF(D{r}="","",MAX(E{r},Liegenschaft!$C$7))'
    M[f"I{r}"] = f'=IF(D{r}="","",IF(F{r}="",Liegenschaft!$C$8,MIN(F{r},Liegenschaft!$C$8)))'
    M[f"J{r}"] = f'=IF(D{r}="",0,MAX(0,I{r}-H{r}+1))'
    M[f"K{r}"] = f'=IF(OR(D{r}="",Liegenschaft!$C$9=""),0,J{r}/Liegenschaft!$C$9)'
    hz = [openpyxl.utils.get_column_letter(HCOL0 + k) for k in range(12)]
    for k in range(12):
        lr = 14 + k
        M[f"{hz[k]}{r}"] = (f'=IF(OR(D{r}="",Liegenschaft!$C$7=""),0,MAX(0,MIN(I{r},Liegenschaft!$N${lr})'
                            f'-MAX(H{r},Liegenschaft!$M${lr})+1)/DAY(Liegenschaft!$N${lr})*Liegenschaft!$O${lr})')
        M[f"{hz[k]}{r}"].font = f(8, color="7A8A90")
    M[f"L{r}"] = f'=IF(OR(D{r}="",Liegenschaft!$J$26=0),0,SUM({hz[0]}{r}:{hz[-1]}{r})/Liegenschaft!$J$26)'
    M[f"M{r}"] = (f'=IF(D{r}="","",IF(ISNA(MATCH(C{r},Liegenschaft!$B${R_EINH_ERST}:$B${R_EINH_LETZT},0)),"⚠ Wohnung unbekannt",'
                  f'IF(J{r}=0,"⚠ nicht in Periode",IF(SUMPRODUCT(($C$6:$C${5 + N_MIETER}=C{r})*($D$6:$D${5 + N_MIETER}<>"")'
                  f'*($J$6:$J${5 + N_MIETER}>0)*($H$6:$H${5 + N_MIETER}<=I{r})*($I$6:$I${5 + N_MIETER}>=H{r}))>1,'
                  f'"⚠ Überschneidung mit anderem Mieter","✓"))))')
    for col, fmt in [("H", DAT), ("I", DAT), ("J", "0"), ("K", PROZ), ("L", PROZ), ("M", None)]:
        formel(M[f"{col}{r}"], fmt)
dv = DataValidation(type="list", formula1=f"Liegenschaft!$B${R_EINH_ERST}:$B${R_EINH_LETZT}", allow_blank=True)
M.add_data_validation(dv)
dv.add(f"C6:C{5 + N_MIETER}")
M.cell(row=5, column=HCOL0, value="Hilfsspalten Heizgewichtung (nicht bearbeiten)").font = f(8, italic=True, color="7A8A90")
M.freeze_panes = "E6"

# ================================================================ Kosten
K = wb.create_sheet("Kosten")
titel(K, "Nebenkosten der Periode", "Rechnungsbeträge der ganzen Liegenschaft erfassen. Nur Positionen mit «Ja» / «Ja» werden den Mietern belastet.")
kopf(K, 5, 2, ["Kostenposition", "Betrag CHF", "Art", "Verteilschlüssel", "Im Mietvertrag ausgeschieden?", "Umlagefähig?", "Abrechenbar", ""])
for col, w in zip("BCDEFGHIJ", [30, 14, 22, 18, 15, 13, 12, 3, 60]):
    K.column_dimensions[col].width = w
ARTEN = ["Heizung/Warmwasser", "Übrige Nebenkosten"]
SCHLUESSEL = ["Fläche m²", "Zimmer", "Wertquote", "Volumen m³", "pro Einheit gleich"]
bsp_kosten = [
    ("Heizöl / Gas inkl. Lieferung", 9800, ARTEN[0], SCHLUESSEL[3], "Ja", "Ja"),
    ("Service Heizung, Kaminfeger", 780, ARTEN[0], SCHLUESSEL[3], "Ja", "Ja"),
    ("Wasser / Abwasser", 2600, ARTEN[1], SCHLUESSEL[0], "Ja", "Ja"),
    ("Allgemeinstrom Treppenhaus", 420, ARTEN[1], SCHLUESSEL[4], "Ja", "Ja"),
    ("Hauswartung", 3200, ARTEN[1], SCHLUESSEL[0], "Ja", "Ja"),
    ("Kehrichtgebühren", 690, ARTEN[1], SCHLUESSEL[4], "Nein", "Ja"),
    ("Reparatur Storen", 1150, ARTEN[1], SCHLUESSEL[0], "Nein", "Nein"),
]
for i in range(N_KOSTEN):
    r = 6 + i
    vals = bsp_kosten[i] if BEISPIEL and i < len(bsp_kosten) else (None,) * 6
    for j, v in enumerate(vals):
        eingabe(K.cell(row=r, column=2 + j, value=v), CHF if j == 1 else None)
    K[f"H{r}"] = f'=IF(OR(B{r}="",C{r}=""),"",IF(AND(F{r}="Ja",G{r}="Ja",D{r}<>"",E{r}<>""),"Ja","Nein"))'
    formel(K[f"H{r}"])
    K[f"I{r}"] = f'=IF(E{r}="",1,MATCH(E{r},{{"{SCHLUESSEL[0]}","{SCHLUESSEL[1]}","{SCHLUESSEL[2]}","{SCHLUESSEL[3]}","{SCHLUESSEL[4]}"}},0))'
    K[f"I{r}"].font = f(8, color="FFFFFF")
liste(K, f"D6:D{5 + N_KOSTEN}", ARTEN)
liste(K, f"E6:E{5 + N_KOSTEN}", SCHLUESSEL)
liste(K, f"F6:G{5 + N_KOSTEN}", ["Ja", "Nein"])
RK_SUM = 6 + N_KOSTEN  # 26
K[f"B{RK_SUM}"] = "Total erfasste Kosten"
K[f"C{RK_SUM}"] = f"=SUM(C6:C{RK_SUM - 1})"
K[f"B{RK_SUM + 1}"] = "davon abrechenbar"
K[f"C{RK_SUM + 1}"] = f'=SUMIF(H6:H{RK_SUM - 1},"Ja",C6:C{RK_SUM - 1})'
K[f"B{RK_SUM + 2}"] = "nicht abrechenbar (trägt Vermieter)"
K[f"C{RK_SUM + 2}"] = f"=C{RK_SUM}-C{RK_SUM + 1}"
for rr in range(RK_SUM, RK_SUM + 3):
    K[f"B{rr}"].font = f(bold=True)
    formel(K[f"C{rr}"], CHF, True)
hinweise = [
    "Hinweise",
    "Umlagefähig sind nur Betriebskosten (Heizung, Warmwasser, Wasser, Abwasser, Allgemeinstrom, Hauswart, Kehricht, Lift-Service …).",
    "NICHT umlagefähig: Reparaturen, Unterhalt, Erneuerungen, Hypothekarzinsen, Gebäudeversicherung (Praxis), Abschreibungen.",
    "Nebenkosten dürfen nur belastet werden, wenn sie im Mietvertrag einzeln ausgeschieden sind (Art. 257a Abs. 2 OR).",
    "Heizung/Warmwasser wird nach Heizgradtagen verteilt, alles andere zeitanteilig nach Tagen.",
]
for i, t in enumerate(hinweise):
    K.cell(row=6 + i, column=10, value=t).font = f(9, i == 0, AKZENT if i == 0 else "4A5A60")
K.freeze_panes = "C6"

# ================================================================ Verteilung
V = wb.create_sheet("Verteilung")
titel(V, "Verteilung auf die Mietverhältnisse", "Wird automatisch berechnet. Zeilen = Mieter (wie Blatt «Mieter»), Spalten = Kostenpositionen.")
kopf(V, 5, 2, ["Nr", "Mieter/in"])
V.column_dimensions["B"].width = 5
V.column_dimensions["C"].width = 22
KCOL0 = 4  # Spalte D
kcols = [openpyxl.utils.get_column_letter(KCOL0 + k) for k in range(N_KOSTEN)]
for k in range(N_KOSTEN):
    kr = 6 + k
    c = V.cell(row=5, column=KCOL0 + k, value=f'=IF(Kosten!B{kr}="","Pos. {k + 1}",Kosten!B{kr})')
    c.font = f(8, True, "FFFFFF")
    c.fill = fill(AKZENT)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    V.column_dimensions[kcols[k]].width = 11
c_nk, c_hon, c_tot, c_ak, c_sal = [openpyxl.utils.get_column_letter(KCOL0 + N_KOSTEN + i) for i in range(5)]
kopf(V, 5, KCOL0 + N_KOSTEN, ["Nebenkosten", "Verwaltungs-honorar", "Total Anteil", "Akonto bezahlt", "Saldo (+ Nachzahlung / − Guthaben)"])
for col in (c_nk, c_hon, c_tot, c_ak, c_sal):
    V.column_dimensions[col].width = 13
V.row_dimensions[5].height = 42
ER = R_EINH_ERST
LR = R_EINH_LETZT
for i in range(N_MIETER):
    r = 6 + i
    V[f"B{r}"] = i + 1
    V[f"C{r}"] = f'=IF(Mieter!D{r}="","",Mieter!D{r})'
    for k in range(N_KOSTEN):
        kr = 6 + k
        V[f"{kcols[k]}{r}"] = (
            f'=IF(OR(Mieter!$D{r}="",Kosten!$H${kr}<>"Ja"),0,IFERROR(Kosten!$C${kr}'
            f'*INDEX(Liegenschaft!$C${ER}:$G${LR},MATCH(Mieter!$C{r},Liegenschaft!$B${ER}:$B${LR},0),Kosten!$I${kr})'
            f'/INDEX(Liegenschaft!$C${R_EINH_SUM}:$G${R_EINH_SUM},1,Kosten!$I${kr})'
            f'*IF(Kosten!$D${kr}="{ARTEN[0]}",Mieter!$L{r},Mieter!$K{r}),0))')
        formel(V[f"{kcols[k]}{r}"], "#,##0.00")
    V[f"{c_nk}{r}"] = f"=SUM({kcols[0]}{r}:{kcols[-1]}{r})"
    V[f"{c_hon}{r}"] = f'=IF(Liegenschaft!$C$11="Ja",ROUND({c_nk}{r}*Liegenschaft!$C$10,2),0)'
    V[f"{c_tot}{r}"] = f"=ROUND({c_nk}{r}+{c_hon}{r},2)"
    V[f"{c_ak}{r}"] = f"=N(Mieter!G{r})"
    V[f"{c_sal}{r}"] = f'=IF(C{r}="",0,{c_tot}{r}-{c_ak}{r})'
    for col in (c_nk, c_hon, c_tot, c_ak, c_sal):
        formel(V[f"{col}{r}"], "#,##0.00", col == c_sal)
RV_SUM = 6 + N_MIETER  # 36
V[f"C{RV_SUM}"] = "Total Mieter"
V[f"C{RV_SUM + 1}"] = "Vermieteranteil (Leerstand)"
V[f"C{RV_SUM + 2}"] = "Abrechenbar gesamt"
for k in range(N_KOSTEN):
    col = kcols[k]
    kr = 6 + k
    V[f"{col}{RV_SUM}"] = f"=SUM({col}6:{col}{RV_SUM - 1})"
    V[f"{col}{RV_SUM + 2}"] = f'=IF(Kosten!$H${kr}="Ja",Kosten!$C${kr},0)'
    V[f"{col}{RV_SUM + 1}"] = f"={col}{RV_SUM + 2}-{col}{RV_SUM}"
for col in (c_nk, c_hon, c_tot, c_ak, c_sal):
    V[f"{col}{RV_SUM}"] = f"=SUM({col}6:{col}{RV_SUM - 1})"
V[f"{c_nk}{RV_SUM + 1}"] = f"=SUM({kcols[0]}{RV_SUM + 1}:{kcols[-1]}{RV_SUM + 1})"
V[f"{c_nk}{RV_SUM + 2}"] = f"=SUM({kcols[0]}{RV_SUM + 2}:{kcols[-1]}{RV_SUM + 2})"
for rr in range(RV_SUM, RV_SUM + 3):
    V[f"C{rr}"].font = f(bold=True)
    for col in kcols + [c_nk, c_hon, c_tot, c_ak, c_sal]:
        if V[f"{col}{rr}"].value is not None:
            formel(V[f"{col}{rr}"], "#,##0.00", True)
V[f"C{RV_SUM + 4}"] = "Kontrollsumme"
V[f"D{RV_SUM + 4}"] = (f'=IF(ABS({c_nk}{RV_SUM}+{c_nk}{RV_SUM + 1}-Kosten!C{RK_SUM + 1})<0.01,'
                       f'"✓ Mieter + Vermieter = abrechenbare Kosten","⚠ Differenz – Wohnungsangaben prüfen")')
V[f"C{RV_SUM + 4}"].font = f(bold=True)
V[f"D{RV_SUM + 4}"].font = f(bold=True)
V.freeze_panes = "D6"

# ================================================================ Abrechnung (Druck)
A = wb.create_sheet("Abrechnung")
A.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [2, 34, 15, 18, 13, 12, 15]):
    A.column_dimensions[col].width = w
A["B2"] = "Mieter-Nr. wählen:"
A["B2"].font = f(10, True, AKZENT)
eingabe(A["C2"])
A["C2"] = 1
dvn = DataValidation(type="whole", operator="between", formula1="1", formula2=str(N_MIETER))
A.add_data_validation(dvn)
dvn.add("C2")
A["D2"] = "(diese Zeile wird nicht gedruckt)"
A["D2"].font = f(8, italic=True, color="7A8A90")
idx = "$C$2+5"   # Zeile im Blatt Mieter/Verteilung
A["B4"] = '=Liegenschaft!C6'
A["B5"] = '=Liegenschaft!C4&", "&Liegenschaft!C5'
A["E4"] = f'=INDEX(Mieter!D:D,{idx})'
A["E5"] = f'="Wohnung: "&INDEX(Mieter!C:C,{idx})'
for c in ("B4", "E4"):
    A[c].font = f(10, True)
A["B8"] = "Heiz- und Nebenkostenabrechnung"
A["B8"].font = f(15, True, AKZENT)
A["B9"] = '="Periode "&TEXT(Liegenschaft!C7,"DD.MM.YYYY")&" – "&TEXT(Liegenschaft!C8,"DD.MM.YYYY")'
A["B10"] = (f'="Ihre Mietdauer in der Periode: "&TEXT(INDEX(Mieter!H:H,{idx}),"DD.MM.YYYY")&" – "'
            f'&TEXT(INDEX(Mieter!I:I,{idx}),"DD.MM.YYYY")&" ("&INDEX(Mieter!J:J,{idx})&" Tage)"')
kopf(A, 12, 2, ["Kostenposition", "Kosten Liegenschaft", "Verteilschlüssel", "Anteil Wohnung", "Zeit-/Heizanteil", "Ihr Anteil CHF"])
A.row_dimensions[12].height = 30
for k in range(N_KOSTEN):
    r = 13 + k
    kr = 6 + k
    zeige = f'Kosten!$H${kr}="Ja"'
    A[f"B{r}"] = f'=IF({zeige},Kosten!B{kr},"")'
    A[f"C{r}"] = f'=IF({zeige},Kosten!C{kr},"")'
    A[f"D{r}"] = f'=IF({zeige},Kosten!E{kr},"")'
    A[f"E{r}"] = (f'=IF({zeige},IFERROR(INDEX(Liegenschaft!$C${ER}:$G${LR},MATCH(INDEX(Mieter!C:C,{idx}),Liegenschaft!$B${ER}:$B${LR},0),Kosten!$I${kr})'
                  f'/INDEX(Liegenschaft!$C${R_EINH_SUM}:$G${R_EINH_SUM},1,Kosten!$I${kr}),0),"")')
    A[f"F{r}"] = f'=IF({zeige},IF(Kosten!D{kr}="{ARTEN[0]}",INDEX(Mieter!L:L,{idx}),INDEX(Mieter!K:K,{idx})),"")'
    A[f"G{r}"] = f'=IF({zeige},INDEX(Verteilung!{kcols[k]}:{kcols[k]},{idx}),"")'
    for col, fmt in zip("BCDEFG", [None, "#,##0.00", None, PROZ, PROZ, "#,##0.00"]):
        c = A[f"{col}{r}"]
        c.number_format = fmt or "General"
        c.font = f(9)
        c.border = Border(bottom=Side(style="hair", color=RAND))
RA = 13 + N_KOSTEN  # 33
zeilen = [
    ("Total Nebenkosten", f"=INDEX(Verteilung!{c_nk}:{c_nk},{idx})"),
    ('=IF(Liegenschaft!C11="Ja","Verwaltungshonorar "&TEXT(Liegenschaft!C10,"0.0%"),"")', f"=INDEX(Verteilung!{c_hon}:{c_hon},{idx})"),
    ("Total Ihr Anteil", f"=INDEX(Verteilung!{c_tot}:{c_tot},{idx})"),
    ("abzüglich Akontozahlungen", f"=-INDEX(Verteilung!{c_ak}:{c_ak},{idx})"),
    (f'=IF(INDEX(Verteilung!{c_sal}:{c_sal},{idx})>=0,"Nachzahlung zu Ihren Lasten","Guthaben zu Ihren Gunsten")',
     f"=ABS(INDEX(Verteilung!{c_sal}:{c_sal},{idx}))"),
]
for i, (lab, val) in enumerate(zeilen):
    r = RA + 1 + i
    A[f"B{r}"] = lab
    A[f"G{r}"] = val
    A[f"G{r}"].number_format = CHF
    fett = i in (2, 4)
    A[f"B{r}"].font = f(10 if not fett else 11, fett)
    A[f"G{r}"].font = f(10 if not fett else 11, fett)
    if i == 4:
        for col in "BCDEFG":
            A[f"{col}{r}"].fill = fill(HELL)
            A[f"{col}{r}"].border = Border(top=Side(style="medium", color=AKZENT))
RT = RA + 8
fuss = [
    "Die Kosten der Liegenschaft werden nach den im Mietvertrag vereinbarten Schlüsseln verteilt. Heiz- und Warmwasserkosten",
    "werden bei Mieterwechseln nach Heizgradtagen, die übrigen Nebenkosten nach Tagen aufgeteilt. Kosten leerstehender",
    "Wohnungen trägt die Vermieterschaft. Die Belege können Sie auf Anfrage einsehen (Art. 257b Abs. 2 OR).",
    "",
    '="Bitte überweisen Sie eine Nachzahlung innert 30 Tagen. Ein Guthaben überweisen wir Ihnen auf Ihr Konto."',
    "",
    '=Liegenschaft!C5&", "&TEXT(TODAY(),"DD.MM.YYYY")',
    "",
    "_______________________________",
    "=Liegenschaft!C6",
]
for i, t in enumerate(fuss):
    A[f"B{RT + i}"] = t
    A[f"B{RT + i}"].font = f(9, color="4A5A60")
A.print_area = f"B4:G{RT + len(fuss)}"
A.page_setup.paperSize = A.PAPERSIZE_A4
A.page_setup.fitToWidth = 1
A.page_setup.fitToHeight = 1
A.sheet_properties.pageSetUpPr.fitToPage = True

# ================================================================ Mietzins
Z = wb.create_sheet("Mietzins")
titel(Z, "Mietzinsanpassung berechnen", "Referenzzinssatz (Art. 13 VMWG), Teuerung (Art. 16 VMWG) und allgemeine Kostensteigerung.")
Z.column_dimensions["B"].width = 52
Z.column_dimensions["C"].width = 16
Z.column_dimensions["D"].width = 70
ein = [
    (5, "Aktueller Nettomietzins pro Monat", CHF, 1850 if BEISPIEL else None, ""),
    (6, "Referenzzinssatz bei der letzten Anpassung", PROZ, 0.0175 if BEISPIEL else None, "Steht im letzten Mietzinsformular bzw. im Mietvertrag."),
    (7, "Aktueller Referenzzinssatz", PROZ, 0.0125, "Stand Okt. 2026: 1.25 %. Quartalsweise auf bwo.admin.ch prüfen."),
    (8, "Landesindex (LIK) bei der letzten Anpassung", "0.0", 106.2 if BEISPIEL else None, "Gleiche Basis wie aktueller Wert verwenden (z. B. Dez. 2020 = 100)."),
    (9, "Landesindex (LIK) aktuell", "0.0", 107.4 if BEISPIEL else None, "bfs.admin.ch › Landesindex der Konsumentenpreise."),
    (10, "Stichtag der letzten Anpassung", DAT, date(2023, 10, 1) if BEISPIEL else None, "Ab diesem Datum wird die Kostensteigerung gerechnet."),
    (11, "Wirksam werden der neuen Miete (erster Tag)", DAT, date(2027, 4, 1) if BEISPIEL else None, "Nächster vertraglicher Kündigungstermin + 1 Tag."),
    (12, "Allgemeine Kostensteigerung pro Jahr", PROZ, 0.005, "Pauschale gemäss kantonaler Praxis (häufig 0.5 %–1 %). Belegbar halten."),
    (13, "Kündigungsfrist in Monaten", "0", 3, "Wohnungen: mindestens 3 Monate (Art. 266c OR)."),
]
for r, lab, fmt, val, hint in ein:
    Z[f"B{r}"] = lab
    Z[f"B{r}"].font = f()
    eingabe(Z.cell(row=r, column=3, value=val), fmt)
    Z[f"D{r}"] = hint
    Z[f"D{r}"].font = f(8, italic=True, color="5A6B72")
kopf(Z, 15, 2, ["Komponente", "Veränderung", "Erläuterung"])
Z["B16"] = "Referenzzinssatz"
Z["C16"] = ('=IF(OR(C6="",C7=""),0,IF(ROUND((C7-C6)/0.0025,0)>=0,ROUND((C7-C6)/0.0025,0)*0.03,'
            '1/(1+ROUND((C6-C7)/0.0025,0)*0.03)-1))')
Z["D16"] = ('=IF(OR(C6="",C7=""),"",ROUND((C7-C6)/0.0025,0)&" Schritt(e) à 0.25 %: Erhöhung +3 % je Schritt, Senkung −2.91 % / −5.66 % / …"'
            '&IF(MAX(C6,C7)>0.05," ⚠ über 5 %: andere Sätze (Art. 13 VMWG)",""))')
Z["B17"] = "Teuerung (40 % der LIK-Veränderung)"
Z["C17"] = '=IF(OR(N(C8)=0,N(C9)=0),0,(C9/C8-1)*0.4)'
Z["D17"] = '=IF(OR(N(C8)=0,N(C9)=0),"","LIK "&TEXT(C9/C8-1,"0.00%")&" × 40 %")'
Z["B18"] = "Allgemeine Kostensteigerung"
Z["C18"] = '=IF(OR(C10="",C11=""),0,C12*MAX(0,(YEAR(C11)-YEAR(C10))*12+MONTH(C11)-MONTH(C10))/12)'
Z["D18"] = '=IF(OR(C10="",C11=""),"",MAX(0,(YEAR(C11)-YEAR(C10))*12+MONTH(C11)-MONTH(C10))&" Monate × "&TEXT(C12,"0.00%")&" / 12")'
Z["B19"] = "Total Anpassung"
Z["C19"] = "=C16+C17+C18"
for r in range(16, 20):
    formel(Z[f"C{r}"], "+0.00%;-0.00%;0.00%", r == 19)
    Z[f"B{r}"].border = rahmen
    Z[f"D{r}"].font = f(8, color="5A6B72")
Z["B19"].font = f(bold=True)
Z["B21"] = "Neuer Nettomietzins pro Monat"
Z["C21"] = "=IF(N(C5)=0,0,ROUND(C5*(1+C19)*20,0)/20)"
Z["B22"] = "Differenz pro Monat"
Z["C22"] = "=C21-N(C5)"
Z["B23"] = "Differenz pro Jahr"
Z["C23"] = "=C22*12"
for r in (21, 22, 23):
    Z[f"B{r}"].font = f(11 if r == 21 else 10, r == 21)
    formel(Z[f"C{r}"], CHF, r == 21)
Z["C21"].fill = fill(HELL)
Z["B25"] = "Ergebnis"
Z["B25"].font = f(10, True, AKZENT)
Z["B26"] = ('=IF(N(C5)=0,"Bitte Eingaben ausfüllen.",IF(C19>0,"Erhöhung möglich. Mitteilung zwingend mit amtlichem kantonalem Formular und Begründung.",'
            'IF(C19<0,"Die Mieterschaft hat Anspruch auf eine Senkung, wenn sie diese verlangt. Freiwillige Senkung möglich.","Keine Anpassung.")))')
Z.merge_cells("B26:D26")
Z["B26"].alignment = Alignment(wrap_text=True, vertical="top")
Z.row_dimensions[26].height = 28
Z.print_area = "B2:D30"
Z.page_setup.paperSize = Z.PAPERSIZE_A4
Z.page_setup.orientation = "landscape"
Z.sheet_properties.pageSetUpPr.fitToPage = True
Z.page_setup.fitToHeight = 1
Z["B27"] = "Zugang beim Mieter spätestens (inkl. 1 Tag Reserve)"
Z["C27"] = '=IF(C11="","",EDATE(C11,-C13)-11)'
formel(Z["C27"], DAT, True)
Z["D27"] = "Mitteilung muss mind. 10 Tage vor Beginn der Kündigungsfrist eintreffen (Art. 269d OR). Eingeschrieben senden."
Z["D27"].font = f(8, italic=True, color="5A6B72")
Z["B29"] = ("Methode: Die Komponenten werden addiert (gängige Praxis). Gerundet auf 5 Rappen. Ob eine Anpassung im Einzelfall zulässig ist "
            "(z. B. bei übersetztem Ertrag oder vorbehaltenen Reserven), prüfen Sie mit Ihrem Verband oder einer Fachperson.")
Z["B29"].alignment = Alignment(wrap_text=True, vertical="top")
Z.merge_cells("B29:D30")
Z.row_dimensions[29].height = 30
Z["B29"].font = f(8, italic=True, color="5A6B72")

# ================================================================ Mieterwechsel
W = wb.create_sheet("Mieterwechsel")
titel(W, "Checkliste Mieterwechsel", "Auszugsdatum eintragen – die Termine rechnen sich selbst. Überfällige offene Schritte werden rot.")
for col, w in zip("BCDEFG", [18, 64, 12, 12, 10, 30]):
    W.column_dimensions[col].width = w
W["B4"] = "Wohnung"
eingabe(W["C4"], None)
W["B5"] = "Auszug Vormieter"
eingabe(W["C5"], DAT)
W["B6"] = "Einzug Nachmieter"
eingabe(W["C6"], DAT)
if BEISPIEL:
    W["C4"], W["C5"], W["C6"] = "EG rechts", date(2025, 12, 31), date(2026, 2, 1)
kopf(W, 8, 2, ["Phase", "Schritt", "Tage rel. Auszug", "Fällig am", "Erledigt", "Bemerkung"])
schritte = [
    ("Kündigung", "Kündigung prüfen: Schriftform, beide Ehegatten/Partner unterschrieben, Frist und Termin (Art. 266a ff., 266n OR)", -90),
    ("Kündigung", "Kündigung schriftlich bestätigen, bei ausserterminlicher Kündigung auf Ersatzmieter-Regel hinweisen (Art. 264 OR)", -88),
    ("Kündigung", "Mietzins für Neuvermietung festlegen; in Kantonen mit Formularpflicht Anfangsmietzins-Formular vorbereiten (Art. 270 OR)", -80),
    ("Vermarktung", "Fotos, Inserat und Besichtigungstermine mit Vormieter vereinbaren (Besichtigungsrecht Art. 257h OR)", -75),
    ("Vermarktung", "Bewerbungen prüfen: Betreibungsregisterauszug, Einkommen/Miete-Verhältnis, Referenzen", -55),
    ("Vermarktung", "Mietvertrag erstellen und versenden – Nebenkosten einzeln aufführen (Art. 257a OR)", -45),
    ("Vermarktung", "Unterzeichneten Vertrag erhalten; Mietzinsdepot einfordern (max. 3 Monatsmieten, Sperrkonto, Art. 257e OR)", -30),
    ("Abnahme", "Vorbesichtigung mit Vormieter: absehbare Mängel und Reinigung besprechen", -21),
    ("Abnahme", "Termin Wohnungsabnahme festlegen und schriftlich bestätigen", -14),
    ("Abnahme", "Handwerker (Maler, Bodenleger, Reinigung) provisorisch reservieren", -10),
    ("Abnahme", "Wohnungsabnahme durchführen, Protokoll von beiden Seiten unterzeichnen lassen, Schlüssel zählen", 0),
    ("Abnahme", "Zählerstände ablesen; Strom/Gas für Leerstand auf Vermieter melden", 0),
    ("Abnahme", "Mängelrüge an Vormieter sofort schriftlich (innert 2–3 Werktagen, sonst Verwirkung – Art. 267a OR)", 2),
    ("Instandstellung", "Instandstellungsarbeiten beauftragen und Termine koordinieren", 3),
    ("Instandstellung", "Endreinigung / Nachreinigung organisieren, Namensschilder Briefkasten und Klingel bestellen", 7),
    ("Übergabe", "Prüfen, ob Mietzinsdepot des Nachmieters eingegangen ist – erst dann Schlüssel übergeben", "E"),
    ("Übergabe", "Übergabe an Nachmieter mit Antrittsprotokoll; Frist für Mängelliste nennen (üblich 10–30 Tage)", "E"),
    ("Übergabe", "Zählerstände an Werke melden (Wechsel auf Nachmieter), Mieterliste/Hauswart informieren", "E"),
    ("Abrechnung", "Schadensabrechnung Vormieter nach Lebensdauertabelle (Altersentwertung, Blatt «Abnahme»)", 30),
    ("Abrechnung", "Nebenkosten-Teilabrechnung Vormieter erstellen oder in der Jahresabrechnung berücksichtigen", 45),
    ("Abrechnung", "Mietzinsdepot Vormieter freigeben bzw. mit Forderungen verrechnen", 60),
]
for i, (ph, txt, off) in enumerate(schritte):
    r = 9 + i
    W[f"B{r}"] = ph
    W[f"C{r}"] = txt
    if off == "E":
        W[f"D{r}"] = "Einzug"
        W[f"E{r}"] = '=IF($C$6="","",$C$6)'
    else:
        W[f"D{r}"] = off
        W[f"E{r}"] = f'=IF($C$5="","",$C$5+D{r})'
    eingabe(W[f"F{r}"])
    eingabe(W[f"G{r}"])
    for col in "BCDE":
        W[f"{col}{r}"].border = rahmen
        W[f"{col}{r}"].font = f(9)
    W[f"E{r}"].number_format = DAT
    W[f"C{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    if BEISPIEL and i < 12:
        W[f"F{r}"] = "Ja"
RW = 8 + len(schritte)
liste(W, f"F9:F{RW}", ["Ja", "n. erf."])
W.conditional_formatting.add(f"B9:G{RW}", FormulaRule(formula=[f'AND($E9<>"",$E9<TODAY(),$F9="")'], fill=fill(WARN)))
W.conditional_formatting.add(f"B9:G{RW}", FormulaRule(formula=['$F9<>""'], fill=fill(OK)))
W[f"B{RW + 2}"] = "Offene Schritte"
W[f"C{RW + 2}"] = f'=COUNTIF(F9:F{RW},"")&" von {len(schritte)} offen, davon überfällig: "&SUMPRODUCT((F9:F{RW}="")*ISNUMBER(E9:E{RW})*(E9:E{RW}<TODAY()))'
W[f"B{RW + 2}"].font = f(bold=True)

# ================================================================ Abnahme
B = wb.create_sheet("Abnahme")
titel(B, "Wohnungsabnahmeprotokoll", "Ausdrucken oder am Tablet ausfüllen. Mängel konkret beschreiben (Ort, Grösse, Foto-Nr.).")
for col, w in zip("BCDEFG", [22, 26, 14, 38, 12, 12]):
    B.column_dimensions[col].width = w
kopfdaten = ["Liegenschaft / Wohnung", "Mieter/in (ausziehend)", "Mieter/in (einziehend)", "Datum / Uhrzeit", "Anwesend"]
for i, t in enumerate(kopfdaten):
    B[f"B{4 + i}"] = t
    B[f"B{4 + i}"].font = f(9, True)
    B.merge_cells(f"C{4 + i}:G{4 + i}")
    eingabe(B[f"C{4 + i}"])
kopf(B, 10, 2, ["Raum", "Bauteil", "Zustand", "Mangel / Bemerkung", "zulasten Mieter?", "Foto-Nr."])
raeume = ["Eingang/Korridor", "Wohnzimmer", "Zimmer 1", "Zimmer 2", "Küche", "Bad/WC", "Balkon/Keller"]
bauteile = ["Boden", "Wände/Decke", "Fenster/Storen", "Türen/Schlösser", "Elektro/Lampen", "Einbauten/Geräte"]
r = 11
for raum in raeume:
    for j, bt in enumerate(bauteile):
        B[f"B{r}"] = raum if j == 0 else ""
        B[f"C{r}"] = bt
        for col in "DEFG":
            eingabe(B[f"{col}{r}"])
        for col in "BC":
            B[f"{col}{r}"].border = rahmen
            B[f"{col}{r}"].font = f(9, col == "B" and j == 0)
        r += 1
RB = r - 1
liste(B, f"D11:D{RB}", ["in Ordnung", "normale Abnützung", "Mangel"])
liste(B, f"F11:F{RB}", ["Ja", "Nein", "strittig"])
B.conditional_formatting.add(f"B11:G{RB}", FormulaRule(formula=['$D11="Mangel"'], fill=fill(WARN)))
r = RB + 2
B[f"B{r}"] = "Schlüssel"
B[f"B{r}"].font = f(10, True, AKZENT)
for i, t in enumerate(["Wohnung", "Briefkasten", "Keller/Estrich", "Garage/Velo", "Badges"]):
    B[f"B{r + 1 + i}"] = t
    eingabe(B[f"C{r + 1 + i}"], "0")
B[f"D{r}"] = "Zählerstände"
B[f"D{r}"].font = f(10, True, AKZENT)
for i, t in enumerate(["Strom", "Gas", "Wasser", "Wärme/HKV"]):
    B[f"D{r + 1 + i}"] = t
    eingabe(B[f"E{r + 1 + i}"])
r += 8
B[f"B{r}"] = "Altersentwertung (Lebensdauer gemäss paritätischer Lebensdauertabelle eintragen)"
B[f"B{r}"].font = f(10, True, AKZENT)
kopf(B, r + 1, 2, ["Bauteil", "Neuwert/Kosten CHF", "Lebensdauer J.", "Alter J.", "Restwert %", "zulasten Mieter CHF"])
for i in range(6):
    rr = r + 2 + i
    eingabe(B[f"B{rr}"])
    eingabe(B[f"C{rr}"], CHF)
    eingabe(B[f"D{rr}"], "0")
    eingabe(B[f"E{rr}"], "0.0")
    B[f"F{rr}"] = f'=IF(N(D{rr})=0,"",MAX(0,1-N(E{rr})/D{rr}))'
    B[f"G{rr}"] = f'=IF(F{rr}="","",ROUND(N(C{rr})*F{rr}*20,0)/20)'
    formel(B[f"F{rr}"], "0%")
    formel(B[f"G{rr}"], CHF)
if BEISPIEL:
    B[f"B{r + 2}"], B[f"C{r + 2}"], B[f"D{r + 2}"], B[f"E{r + 2}"] = "Wände Wohnzimmer streichen", 1200, 8, 5
rs = r + 8
B[f"F{rs}"] = "Total"
B[f"G{rs}"] = f"=SUM(G{r + 2}:G{r + 7})"
formel(B[f"G{rs}"], CHF, True)
rs += 2
B[f"B{rs}"] = "Die Unterzeichnenden bestätigen den festgehaltenen Zustand. Mit der Unterschrift anerkennt der Mieter nur die mit «Ja» bezeichneten Mängel."
B[f"B{rs}"].font = f(8, italic=True)
for i, t in enumerate(["Vermieter/in", "Mieter/in ausziehend", "Mieter/in einziehend"]):
    B.cell(row=rs + 3, column=2 + i * 2, value="______________________").font = f()
    B.cell(row=rs + 4, column=2 + i * 2, value=t).font = f(8)
B.print_area = f"B2:G{rs + 4}"
B.page_setup.paperSize = B.PAPERSIZE_A4
B.page_setup.fitToWidth = 1
B.page_setup.fitToHeight = 0
B.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------------------------------------------------- Schutz
for s in (L, M, K, V, A, Z, W, B):
    s.protection.sheet = True
    s.protection.formatColumns = False
    s.protection.formatRows = False
wb.active = 0
wb.save(DATEI)
print("geschrieben:", DATEI)
