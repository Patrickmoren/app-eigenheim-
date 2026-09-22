# -*- coding: utf-8 -*-
"""Rechenprobe: rechnet alle Formeln der Arbeitsmappe wirklich durch und
vergleicht jeden Wert mit einer unabhängig in Python berechneten Erwartung.

Die Erwartungswerte werden bewusst ein zweites Mal ausformuliert und nicht aus
build.py importiert. Weichen Aufbau und Erwartung voneinander ab, fällt der Test
durch – genau das soll er leisten.
"""
import warnings, sys, json
from datetime import date, datetime, timedelta
warnings.filterwarnings("ignore")
import openpyxl
import formulas

DATEI = sys.argv[1] if len(sys.argv) > 1 else "test-klein.xlsx"
BLATTNAME = sys.argv[2] if len(sys.argv) > 2 else "Leerstände"
KOPF, ERSTE = 4, 5
HEUTE = date.today()
EPOCHE = date(1899, 12, 30)

def kauf_runden(x):
    """Excel rundet 0,5 stets vom Nullpunkt weg – Pythons round() nicht."""
    import decimal
    return float(decimal.Decimal(repr(x)).quantize(decimal.Decimal("1"),
                 rounding=decimal.ROUND_HALF_UP))

def seriell(d):
    if isinstance(d, datetime):
        d = d.date()
    return float((d - EPOCHE).days)

# ---------------------------------------------------------------- Eingaben lesen
wb = openpyxl.load_workbook(DATEI)
ws = wb[BLATTNAME]
SPALTEN = [ws.cell(KOPF, i).value for i in range(1, ws.max_column + 1)]
IDX = {n: i + 1 for i, n in enumerate(SPALTEN)}
LETZTE = ws.max_row
STAMM = {}
wst = wb["Stammdaten"]
for r in range(5, 25):
    n = wst.cell(r, 1).value
    if n:
        STAMM[n] = wst.cell(r, 2).value

def roh(r, name):
    v = ws.cell(r, IDX[name]).value
    return "" if v is None else v

def dat(r, name):
    """Datumswert als Serienzahl, sonst None (Text wie «n.e.» zählt nicht als Zahl)."""
    v = roh(r, name)
    if isinstance(v, (date, datetime)):
        return seriell(v)
    return None

def gesetzt(r, name):
    """Belegt im Sinne von <>"" – Datum und «n.e.» gelten beide als erledigt."""
    return roh(r, name) != ""

SCHRITTE = [
    "Kündigung bestätigt / G-Rem", "Meldung Werke / EGT", "Vorbesichtigung",
    "VMZ aktualisiert", "Inserat online", "Wohnungsabnahme durchgeführt",
    "Handwerker aufgeboten", "Zustimmungserklärung", "Instandstellungen beauftragt",
    "Vertrag versendet", "Vertrag retour", "Info Werke / EGT",
    "Schlüsselübergabe erfolgt", "Reinigung veranlasst", "Namensschilder",
    "Kaution / G-Rem mutiert", "Instandstellungs-RG", "Schlussabrechnung", "Zahlungseingang",
]

NAECHSTER = [
    ("Haftungsdatum", "Haftungsdatum erfassen", None),
    ("ex-Mieter", "ex-Mieter erfassen", None),
    ("Kündigung bestätigt / G-Rem", "Kündigung bestätigen / G-Rem", None),
    ("Meldung Werke / EGT", "Werke / EGT melden", None),
    ("Vorbesichtigung", "Vorbesichtigung (oder n.e.)", None),
    ("VMZ aktualisiert", "VMZ aktualisieren", None),
    ("Inserat online", "Inserat aufschalten", None),
    ("WA-Termin", "Abnahmetermin vereinbaren", "Wohnungsabnahme durchgeführt"),
    ("Wohnungsabnahme durchgeführt", "Wohnungsabnahme durchführen", None),
    ("Handwerker aufgeboten", "Handwerker aufbieten", None),
    ("Zustimmungserklärung", "Zustimmungserklärung einholen", None),
    ("Instandstellungen beauftragt", "Instandstellungen beauftragen", None),
    ("neuer Mieter", "Nachmieter erfassen", None),
    ("Vertrag versendet", "Vertrag versenden", None),
    ("Vertrag retour", "Vertrag retour erfassen", None),
    ("Vermietet per", "Vermietet per erfassen", None),
    ("Info Werke / EGT", "Werke / EGT über Einzug informieren", None),
    ("Schlüsselübergabe-Termin", "Schlüsselübergabe terminieren", "Schlüsselübergabe erfolgt"),
    ("Schlüsselübergabe erfolgt", "Schlüsselübergabe durchführen", None),
    ("Reinigung veranlasst", "Reinigung veranlassen", None),
    ("Namensschilder", "Namensschilder bestellen", None),
    ("Kaution / G-Rem mutiert", "Kaution prüfen / G-Rem mutieren", None),
    ("Instandstellungs-RG", "Instandstellungs-Rechnung erfassen", None),
    ("Schlussabrechnung", "Schlussabrechnung erstellen", None),
    ("Zahlungseingang", "Zahlungseingang kontrollieren", None),
]

# ---------------------------------------------------------------- Erwartung je Zeile
def erwartet(r):
    """Liefert {Spaltenname: erwarteter Wert} für eine Zeile."""
    leer = {n: "" for n in SPALTEN}
    if not gesetzt(r, "Obj.-Nr."):
        return leer

    e = {}
    haft = dat(r, "Haftungsdatum")
    kue  = dat(r, "gekündigt per")
    verm = dat(r, "Vermietet per")
    wa   = dat(r, "WA-Termin")
    abg  = gesetzt(r, "Fall abgeschlossen")
    heute_s = seriell(HEUTE)

    # Team
    bw = roh(r, "Bewirtschafter")
    e["Team"] = "" if bw == "" else STAMM.get(bw, "unbekannt")

    # Leerstandsbeginn / Vorlauftage / Berichtsmonat
    e["Leerstandsbeginn"] = "" if haft is None else haft + 1
    e["Vorlauftage"] = "" if (kue is None or haft is None) else haft - kue
    if haft is None:
        e["Berichtsmonat"] = ""
    else:
        b = EPOCHE + timedelta(days=int(haft) + 1)
        e["Berichtsmonat"] = "{}-{:02d}".format(b.year, b.month)

    # Leerstandstage
    if haft is None:
        e["Leerstandstage"] = ""
    elif verm is not None:
        e["Leerstandstage"] = max(0.0, verm - haft - 1)
    else:
        e["Leerstandstage"] = max(0.0, heute_s - haft)

    # Sollfristen
    e["Sollfrist Abnahmetermin"] = "" if haft is None else haft - 30
    wd = dat(r, "Wohnungsabnahme durchgeführt")
    e["Sollfrist Handwerker"] = "" if wd is None else wd + 3

    # Phase / Arbeitsbereich
    if abg:
        ph = "6 · abgeschlossen"
    elif gesetzt(r, "Zahlungseingang"):
        ph = "5 · bereit zum Abschluss"
    elif gesetzt(r, "Kaution / G-Rem mutiert"):
        ph = "5 · Schlussabrechnung"
    elif gesetzt(r, "Vertrag retour"):
        ph = "4 · Übergabe"
    elif gesetzt(r, "Instandstellungen beauftragt"):
        ph = "3 · Neuvermietung"
    elif gesetzt(r, "Haftungsdatum"):
        ph = "2 · Abnahme & Vermarktung"
    else:
        ph = "1 · Erfassung"
    e["Phase"] = ph
    k = ph[0]
    e["Arbeitsbereich"] = {"6": "abgeschlossen", "1": "Kündigung & Abnahme",
                           "2": "Kündigung & Abnahme", "3": "Wiedervermietung",
                           "4": "Wiedervermietung"}.get(k, "Schlussabrechnung")

    # Status
    sa = e["Sollfrist Abnahmetermin"]
    sh = e["Sollfrist Handwerker"]
    st = dat(r, "Schlüsselübergabe-Termin")
    vb = dat(r, "Vertrag retour bis")
    zb = dat(r, "zurückgestellt bis")
    ueberfaellig = (
        (wa is not None and wa < heute_s and not gesetzt(r, "Wohnungsabnahme durchgeführt")) or
        (sa != "" and sa < heute_s and not gesetzt(r, "WA-Termin")
                                   and not gesetzt(r, "Wohnungsabnahme durchgeführt")) or
        (sh != "" and sh < heute_s and not gesetzt(r, "Handwerker aufgeboten")) or
        (st is not None and st < heute_s and not gesetzt(r, "Schlüsselübergabe erfolgt")) or
        (vb is not None and vb < heute_s and not gesetzt(r, "Vertrag retour"))
    )
    if abg:
        e["Status"] = "abgeschlossen"
    elif zb is not None and zb >= heute_s:
        e["Status"] = "zurückgestellt"
    else:
        e["Status"] = "überfällig" if ueberfaellig else "offen"

    # Nächster Schritt – verstrichene Fristen zuerst
    if abg:
        e["Nächster Schritt"] = "—"
    else:
        e["Nächster Schritt"] = "Fall abschliessen"
        frist_texte = [
            (sa != "" and sa < heute_s and not gesetzt(r, "WA-Termin")
                       and not gesetzt(r, "Wohnungsabnahme durchgeführt"),
             "Abnahmetermin vereinbaren – Sollfrist verstrichen"),
            (wa is not None and wa < heute_s and not gesetzt(r, "Wohnungsabnahme durchgeführt"),
             "Wohnungsabnahme nachtragen – Termin war fällig"),
            (sh != "" and sh < heute_s and not gesetzt(r, "Handwerker aufgeboten"),
             "Handwerker aufbieten – Frist verstrichen"),
            (vb is not None and vb < heute_s and not gesetzt(r, "Vertrag retour"),
             "Vertrag retour mahnen – Frist verstrichen"),
            (st is not None and st < heute_s and not gesetzt(r, "Schlüsselübergabe erfolgt"),
             "Schlüsselübergabe nachtragen – Termin war fällig"),
        ]
        for traf, text in frist_texte:
            if traf:
                e["Nächster Schritt"] = text
                break
        else:
         for spalte, text, zusatz in NAECHSTER:
             offen = not gesetzt(r, spalte)
             if zusatz:
                 offen = offen and not gesetzt(r, zusatz)
             if offen:
                 e["Nächster Schritt"] = text
                 break

    # Prüfhinweis
    obj = roh(r, "Obj.-Nr.")
    fnr = roh(r, "Fall-Nr.")
    doppel_obj = sum(1 for x in range(ERSTE, LETZTE + 1)
                     if roh(x, "Obj.-Nr.") == obj and not gesetzt(x, "Fall abgeschlossen"))
    doppel_fn = sum(1 for x in range(ERSTE, LETZTE + 1) if roh(x, "Fall-Nr.") == fnr)
    lt = e["Leerstandstage"]
    hin = []
    if kue is not None and haft is not None and haft < kue:
        hin.append("Haftungsdatum vor Kündigung.")
    if verm is not None and haft is not None and verm < haft:
        hin.append("Vermietung vor Leerstandsbeginn.")
    if wa is not None and kue is not None and wa < kue:
        hin.append("WA-Termin vor Kündigung.")
    if doppel_obj > 1:
        hin.append("Zweiter offener Fall zum selben Objekt.")
    if fnr != "" and doppel_fn > 1:
        hin.append("Fall-Nr. doppelt vergeben.")
    if not gesetzt(r, "ex-Mieter"):
        hin.append("ex-Mieter fehlt.")
    if not gesetzt(r, "Haftungsdatum"):
        hin.append("Haftungsdatum fehlt.")
    if not gesetzt(r, "Bewirtschafter"):
        hin.append("Bewirtschafter fehlt.")
    if abg and not gesetzt(r, "Zahlungseingang"):
        hin.append("Abgeschlossen trotz offener Schlussabrechnung.")
    if (not abg) and verm is None and lt != "" and lt > 60:
        hin.append("Langläufer über 60 Tage.")
    e["Prüfhinweis"] = " ".join(hin)
    return e

# ---------------------------------------------------------------- durchrechnen
xl = formulas.ExcelModel().loads(DATEI).finish()
erg = xl.calculate()
W = {}
for k, v in erg.items():
    try:
        W[k.upper()] = v.value[0, 0]
    except Exception:
        pass
BLATT = "'[{}]{}'!".format(DATEI.upper(), BLATTNAME.upper())
KENN  = "'[{}]KENNZAHLEN'!".format(DATEI.upper())
FEHLT = object()

def wert(praefix, zelle):
    return W.get(praefix + zelle, FEHLT)

def gleich(a, b):
    if isinstance(a, float) or isinstance(b, float):
        try:
            return abs(float(a) - float(b)) < 1e-6
        except (TypeError, ValueError):
            return False
    return a == b

FORMELSPALTEN = ["Team", "Status", "Nächster Schritt", "Leerstandstage", "Prüfhinweis",
                 "Leerstandsbeginn", "Vorlauftage", "Berichtsmonat", "Arbeitsbereich",
                 "Phase", "Sollfrist Abnahmetermin", "Sollfrist Handwerker"]
from openpyxl.utils import get_column_letter
fehler = []
geprueft = 0
for r in range(ERSTE, LETZTE + 1):
    soll = erwartet(r)
    for name in FORMELSPALTEN:
        zelle = "{}{}".format(get_column_letter(IDX[name]), r)
        ist = wert(BLATT, zelle)
        geprueft += 1
        if ist is FEHLT:
            fehler.append("{} ({}, Zeile {}): nicht berechnet".format(zelle, name, r))
        elif not gleich(ist, soll[name]):
            fehler.append("{} ({}, Zeile {}): ist {!r}, erwartet {!r}".format(
                zelle, name, r, ist, soll[name]))

# ---------------------------------------------------------------- Kennzahlen
def zaehle(pruef):
    return float(sum(1 for r in range(ERSTE, LETZTE + 1)
                     if gesetzt(r, "Obj.-Nr.") and pruef(r)))

ST = {r: erwartet(r)["Status"] for r in range(ERSTE, LETZTE + 1)}
BW = {r: roh(r, "Bewirtschafter") for r in range(ERSTE, LETZTE + 1)}
TM = {r: erwartet(r)["Team"] for r in range(ERSTE, LETZTE + 1)}

NUR_LISTE = BLATTNAME != "Leerstände"
wk = wb["Kennzahlen"]
namen = []
for r in range(6, 6 + len(STAMM)):
    namen.append(wst.cell(r - 1, 1).value)

kz_fehler = []
ERW = {r: erwartet(r) for r in range(ERSTE, LETZTE + 1)}

def pruefe_kz(zeile, filt, etikett):
    beendet = [ERW[r]["Leerstandstage"] for r in range(ERSTE, LETZTE + 1)
               if gesetzt(r, "Obj.-Nr.") and filt(r) and dat(r, "Vermietet per") is not None]
    mittel = "–" if not beendet else kauf_runden(sum(beendet) / len(beendet))
    paare = [
        (3, zaehle(filt)),
        (4, zaehle(lambda r: filt(r) and ST[r] == "offen")),
        (5, zaehle(lambda r: filt(r) and ST[r] == "überfällig")),
        (6, zaehle(lambda r: filt(r) and ST[r] == "zurückgestellt")),
        (7, zaehle(lambda r: filt(r) and ST[r] == "abgeschlossen")),
        (8, mittel),
    ]
    for sp, soll in paare:
        z = "{}{}".format(get_column_letter(sp), zeile)
        ist = wert(KENN, z)
        if not gleich(ist, soll):
            kz_fehler.append("Kennzahlen!{} ({}, {}): ist {!r}, erwartet {!r}".format(
                z, etikett, wk.cell(5, sp).value, ist, soll))

if NUR_LISTE:
    namen = []
for i, n in enumerate(namen):
    pruefe_kz(6 + i, (lambda nn: (lambda r: BW[r] == nn))(n), n)
zeile_team = 6 + len(STAMM) + 1
for i, tm in enumerate([] if NUR_LISTE else ["BS 01", "BS 02"]):
    pruefe_kz(zeile_team + i, (lambda tt: (lambda r: TM[r] == tt))(tm), "Total " + tm)
if not NUR_LISTE:
    pruefe_kz(zeile_team + 2, lambda r: True, "Gesamt")
ZEILE_GESAMT = zeile_team + 2

# Offene Fälle nach Arbeitsbereich
zb = ZEILE_GESAMT + 4
for i, bn in enumerate([] if NUR_LISTE else["Kündigung & Abnahme", "Wiedervermietung", "Schlussabrechnung"]):
    soll = float(sum(1 for r in range(ERSTE, LETZTE + 1)
                     if gesetzt(r, "Obj.-Nr.") and ERW[r]["Arbeitsbereich"] == bn))
    z = "C{}".format(zb + i)
    if wk.cell(zb + i, 1).value != bn:
        kz_fehler.append("Kennzahlen: Zeile {} erwartet «{}», steht «{}»".format(
            zb + i, bn, wk.cell(zb + i, 1).value))
    ist = wert(KENN, z)
    if not gleich(ist, soll):
        kz_fehler.append("Kennzahlen!{} ({}): ist {!r}, erwartet {!r}".format(z, bn, ist, soll))

# Prüfhinweise – Häufigkeit
zp = zb + 3 + 3
PRUEFUNGEN = [
    "Haftungsdatum vor Kündigung", "Vermietung vor Leerstandsbeginn", "WA-Termin vor Kündigung",
    "Zweiter offener Fall zum selben Objekt", "Fall-Nr. doppelt vergeben", "ex-Mieter fehlt",
    "Haftungsdatum fehlt", "Bewirtschafter fehlt",
    "Abgeschlossen trotz offener Schlussabrechnung", "Langläufer über 60 Tage",
]
for i, txt in enumerate([] if NUR_LISTE else PRUEFUNGEN):
    soll = float(sum(1 for r in range(ERSTE, LETZTE + 1) if txt in ERW[r]["Prüfhinweis"]))
    z = "C{}".format(zp + i)
    if wk.cell(zp + i, 1).value != txt:
        kz_fehler.append("Kennzahlen: Zeile {} erwartet «{}», steht «{}»".format(
            zp + i, txt, wk.cell(zp + i, 1).value))
    ist = wert(KENN, z)
    if not gleich(ist, soll):
        kz_fehler.append("Kennzahlen!{} ({}): ist {!r}, erwartet {!r}".format(z, txt, ist, soll))

# Fälle mit mindestens einem Hinweis
soll = float(sum(1 for r in range(ERSTE, LETZTE + 1) if ERW[r]["Prüfhinweis"] != ""))
z = "C{}".format(zp + len(PRUEFUNGEN) + 1)
ist = wert(KENN, z)
if not NUR_LISTE and not gleich(ist, soll):
    kz_fehler.append("Kennzahlen!{} (mind. ein Hinweis): ist {!r}, erwartet {!r}".format(z, ist, soll))

# ---------------------------------------------------------------- Ausgabe
print("Datei:", DATEI, "· Blatt:", BLATTNAME)
print("berechnete Zellen:", len(W))
print("verglichene Formelzellen:", geprueft)
if fehler:
    print("\nABWEICHUNGEN in der Liste:", len(fehler))
    for f in fehler[:40]:
        print("  ", f)
if kz_fehler:
    print("\nABWEICHUNGEN in den Kennzahlen:", len(kz_fehler))
    for f in kz_fehler[:40]:
        print("  ", f)
if not fehler and not kz_fehler:
    print("\nErgebnis: alle Werte stimmen mit der unabhängigen Berechnung überein.")
json.dump({"zellen": len(W), "verglichen": geprueft,
           "abweichungen": len(fehler) + len(kz_fehler)},
          open("rechenprobe-umfang.json", "w"))
sys.exit(1 if (fehler or kz_fehler) else 0)
