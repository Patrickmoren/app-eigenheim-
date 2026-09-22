# -*- coding: utf-8 -*-
"""Abdeckungsprüfung: belegt, dass jede Verzweigung in der Prüfdatei
tatsächlich auslöst. Ein Test, der nur bestätigt, dass nichts passiert,
beweist nichts."""
import warnings, sys
warnings.filterwarnings("ignore")
import formulas, openpyxl
from openpyxl.utils import get_column_letter

DATEI = sys.argv[1] if len(sys.argv) > 1 else "test-lastfaelle.xlsx"
BLATTNAME = sys.argv[2] if len(sys.argv) > 2 else "Leerstände"
KOPF, ERSTE = 4, 5
xl = formulas.ExcelModel().loads(DATEI).finish()
erg = xl.calculate()
W = {}
for k, v in erg.items():
    try:
        W[k.upper()] = v.value[0, 0]
    except Exception:
        pass
B = "'[{}]{}'!".format(DATEI.upper(), BLATTNAME.upper())
wb = openpyxl.load_workbook(DATEI)
ws = wb[BLATTNAME]
SP = [ws.cell(KOPF, i).value for i in range(1, ws.max_column + 1)]
IDX = {n: i + 1 for i, n in enumerate(SP)}
LETZTE = ws.max_row

def v(r, name):
    return W.get(B + "{}{}".format(get_column_letter(IDX[name]), r), "<fehlt>")

zeilen = [r for r in range(ERSTE, LETZTE + 1) if ws.cell(r, IDX["Obj.-Nr."]).value]
print("Zeilen mit Inhalt:", len(zeilen), "\n")
kopf = "{:<4} {:<52} {:<15} {:<30} {:>5} {}".format(
    "Zl", "Lastfall", "Status", "Nächster Schritt", "Tage", "Prüfhinweis")
print(kopf); print("-" * len(kopf))
for r in zeilen:
    print("{:<4} {:<52} {:<15} {:<30} {:>5} {}".format(
        r, str(ws.cell(r, IDX["Notiz"]).value)[:52], str(v(r, "Status"))[:15],
        str(v(r, "Nächster Schritt"))[:30], str(v(r, "Leerstandstage"))[:5],
        str(v(r, "Prüfhinweis"))[:70]))

STATUS = {"offen", "überfällig", "zurückgestellt", "abgeschlossen"}
HINWEISE = [
    "Haftungsdatum vor Kündigung.", "Vermietung vor Leerstandsbeginn.",
    "WA-Termin vor Kündigung.", "Zweiter offener Fall zum selben Objekt.",
    "Fall-Nr. doppelt vergeben.", "ex-Mieter fehlt.", "Haftungsdatum fehlt.",
    "Bewirtschafter fehlt.", "Abgeschlossen trotz offener Schlussabrechnung.",
    "Langläufer über 60 Tage.",
]
gesehen_status = {v(r, "Status") for r in zeilen}
alle_hinweise = " ".join(str(v(r, "Prüfhinweis")) for r in zeilen)
luecken = []
for st in STATUS:
    if st not in gesehen_status:
        luecken.append("Status «{}» wird von keinem Lastfall ausgelöst".format(st))
for h in HINWEISE:
    if h not in alle_hinweise:
        luecken.append("Prüfhinweis «{}» wird von keinem Lastfall ausgelöst".format(h))
if any(str(v(r, "Prüfhinweis")).count(".") >= 3 for r in zeilen):
    pass
else:
    luecken.append("Kein Lastfall mit drei oder mehr Hinweisen in einer Zelle")
for r in zeilen:
    p = str(v(r, "Prüfhinweis"))
    if p != p.strip() or "  " in p:
        luecken.append("Zeile {}: Prüfhinweis mit überzähligem Leerzeichen: {!r}".format(r, p))
    if p.startswith(".") or ". ." in p:
        luecken.append("Zeile {}: leerer Hinweis eingefügt: {!r}".format(r, p))
print()
if luecken:
    print("LÜCKEN:", len(luecken))
    for x in luecken:
        print("  ", x)
    sys.exit(1)
print("Abdeckung vollständig: jeder Status und jeder Prüfhinweis wurde mindestens einmal ausgelöst.")
