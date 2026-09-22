# -*- coding: utf-8 -*-
"""Statische Prüfung: Verweise, Zyklen, Struktur, Schutz."""
import re, sys
import openpyxl
from openpyxl.utils import get_column_letter, column_index_from_string

DATEI = sys.argv[1] if len(sys.argv) > 1 else "Leerstandsliste.xlsx"
wb = openpyxl.load_workbook(DATEI)
ok, fehler = 0, []
def P(name, bed, info=""):
    global ok
    if bed: ok += 1; print("  ✓ " + name + ("  " + info if info else ""))
    else:   fehler.append(name + "  " + info); print("  ✗ " + name + "  " + info)

ws = wb["Leerstände"]
KOPF, ERSTE = 4, 5
letzte_sp = ws.max_column
kopf = {ws.cell(KOPF, i).value: i for i in range(1, letzte_sp + 1)}
LETZTE = max(r for r in range(ERSTE, ws.max_row + 1) if ws.cell(r, 1).value or ws.cell(r, 7).value)

print("\n1 · STRUKTUR")
P("Blätter vorhanden", set(["Anleitung","Stammdaten","Leerstände","Kennzahlen","Prozessschritte"]) <= set(wb.sheetnames),
  " / ".join(wb.sheetnames))
P("52 Spalten", letzte_sp == 52, str(letzte_sp))
P("keine doppelten Überschriften", len(kopf) == letzte_sp, "{} eindeutig".format(len(kopf)))
P("keine leere Überschrift", all(ws.cell(KOPF, i).value for i in range(1, letzte_sp + 1)))
tabs = list(ws.tables.values()) if hasattr(ws.tables, "values") else []
P("Tabelle definiert", len(tabs) == 1 and tabs[0].ref.startswith("A4:"), tabs[0].ref if tabs else "keine")
P("Fenster fixiert", ws.freeze_panes == "D5", str(ws.freeze_panes))
P("Blattschutz an", ws.protection.sheet is True)
P("Zeilen einfügen erlaubt", ws.protection.insertRows is False, "(False = erlaubt)")

print("\n2 · SPALTENSCHUTZ")
formel_sp = [i for i in range(1, letzte_sp + 1)
             if isinstance(ws.cell(ERSTE, i).value, str) and str(ws.cell(ERSTE, i).value).startswith("=")]
eingabe_sp = [i for i in range(1, letzte_sp + 1) if i not in formel_sp]
P("Formelspalten gesperrt", all(ws.cell(ERSTE, i).protection.locked for i in formel_sp),
  "{} Spalten".format(len(formel_sp)))
P("Eingabespalten offen", all(ws.cell(ERSTE, i).protection.locked is False for i in eingabe_sp),
  "{} Spalten".format(len(eingabe_sp)))

print("\n3 · FORMELVERWEISE")
ZELLE = re.compile(r"(?<![A-Za-z0-9_.!])(\$?)([A-Z]{1,3})(\$?)(\d{1,5})(?![\(\d])")
BLATT = re.compile(r"([A-Za-zÄÖÜäöüß]+)!\$?([A-Z]{1,3})\$?(\d+)")
schlecht = []
for r in (ERSTE, ERSTE + 1, LETZTE):
    for i in formel_sp:
        f = ws.cell(r, i).value
        for m in BLATT.finditer(f):
            blatt, spb, zl = m.group(1), m.group(2), int(m.group(3))
            if blatt not in wb.sheetnames:
                schlecht.append("{}{} verweist auf unbekanntes Blatt {}".format(get_column_letter(i), r, blatt))
        ohne = BLATT.sub("", f)
        for m in ZELLE.finditer(ohne):
            spb, zl = m.group(2), int(m.group(4))
            try: idx = column_index_from_string(spb)
            except Exception:
                schlecht.append("{}{}: unlesbarer Spaltenname {}".format(get_column_letter(i), r, spb)); continue
            if idx > letzte_sp:
                schlecht.append("{}{} verweist auf Spalte {} (ausserhalb)".format(get_column_letter(i), r, spb))
            if not (ERSTE <= zl <= LETZTE):
                schlecht.append("{}{} verweist auf Zeile {} (ausserhalb {}–{})".format(
                    get_column_letter(i), r, zl, ERSTE, LETZTE))
P("alle Verweise gültig", not schlecht, "; ".join(schlecht[:4]))
P("kein #REF!", not any("#REF" in str(ws.cell(r, i).value) for r in (ERSTE, LETZTE) for i in formel_sp))
P("keine verbotenen Funktionen",
  not any(re.search(r"\b(XLOOKUP|FILTER|UNIQUE|SORT|SEQUENCE|XMATCH|LET|TEXTJOIN)\s*\(", str(ws.cell(ERSTE, i).value))
          for i in formel_sp))

print("\n4 · ZYKLEN")
graph = {}
for i in formel_sp:
    f = ws.cell(ERSTE, i).value
    ohne = BLATT.sub("", f)
    ziele = set()
    for m in ZELLE.finditer(ohne):
        try: idx = column_index_from_string(m.group(2))
        except Exception: continue
        if int(m.group(4)) == ERSTE and idx in formel_sp:
            ziele.add(idx)
    graph[i] = ziele
    if i in ziele:
        fehler.append("Selbstbezug in Spalte " + get_column_letter(i))
def zyklus(start):
    weg, stapel = set(), [(start, [start])]
    while stapel:
        k, pfad = stapel.pop()
        for z in graph.get(k, ()):
            if z == start: return pfad + [z]
            if z not in weg:
                weg.add(z); stapel.append((z, pfad + [z]))
    return None
gefunden = [zyklus(i) for i in graph]
gefunden = [g for g in gefunden if g]
P("keine Zirkelbezüge", not gefunden,
  " → ".join(get_column_letter(x) for x in gefunden[0]) if gefunden else "")

print("\n5 · PRÜFREGELN UND EINFÄRBUNG")
dvs = ws.data_validations.dataValidation
P("Datenprüfungen vorhanden", len(dvs) >= 20, "{} Regeln".format(len(dvs)))
arten = {}
for d in dvs: arten[d.type] = arten.get(d.type, 0) + 1
P("Liste für Bewirtschafter", arten.get("list", 0) >= 1)
P("Datumsprüfung", arten.get("date", 0) >= 1)
P("Schrittprüfung (Datum oder n.e.)", arten.get("custom", 0) == 19, "{} Spalten".format(arten.get("custom", 0)))
P("alle Prüfungen melden Fehler", all(d.showErrorMessage for d in dvs))
anz_regeln = sum(len(v) for v in ws.conditional_formatting._cf_rules.values())
P("bedingte Formatierung", anz_regeln >= 4, "{} Regeln auf {} Bereichen".format(
    anz_regeln, len(ws.conditional_formatting._cf_rules)))

print("\n6 · KENNZAHLEN")
wk = wb["Kennzahlen"]
kf = [(c.coordinate, c.value) for row in wk.iter_rows() for c in row
      if isinstance(c.value, str) and c.value.startswith("=")]
P("Kennzahlenformeln vorhanden", len(kf) >= 60, "{} Formeln".format(len(kf)))
falsch = [k for k, v in kf if "Leerstände!" in v and not re.search(r"Leerstände!\$[A-Z]{1,2}\$\d+:\$[A-Z]{1,2}\$\d+", v)]
P("Bezüge auf «Leerstände» absolut", not falsch, "; ".join(falsch[:3]))
ausserhalb = []
for k, v in kf:
    for m in re.finditer(r"Leerstände!\$([A-Z]{1,2})\$(\d+):\$([A-Z]{1,2})\$(\d+)", v):
        if column_index_from_string(m.group(1)) > letzte_sp or int(m.group(4)) != LETZTE:
            ausserhalb.append(k)
P("Kennzahlen greifen auf gültige Bereiche", not ausserhalb, "; ".join(sorted(set(ausserhalb))[:3]))

print("\n7 · BEISPIELDATEN")
wbsp = wb["Beispiele"]
bsp = [r for r in range(ERSTE, ERSTE + 20) if str(wbsp.cell(r, 1).value or "").startswith("BSP-")]
P("Beispiele auf eigenem Blatt", len(bsp) == 10, "{} Zeilen".format(len(bsp)))
P("Arbeitsblatt «Leerstände» ist leer",
  not any(ws.cell(r, 1).value or ws.cell(r, 2).value for r in range(ERSTE, LETZTE + 1)))
P("Verweis auf das Beispielblatt", "Beispiele" in str(ws["A3"].value))
P("Beispielblatt hat dieselben Spalten",
  [wbsp.cell(KOPF, i).value for i in range(1, letzte_sp + 1)] ==
  [ws.cell(KOPF, i).value for i in range(1, letzte_sp + 1)])
P("Beispielblatt wirkt nicht auf die Kennzahlen",
  not any("Beispiele!" in str(v) for _, v in kf))

print("\n8 · VERHALTEN BEIM ÖFFNEN")
P("Neuberechnung beim Öffnen erzwungen", bool(wb.calculation and wb.calculation.fullCalcOnLoad))
P("Berechnung auf automatisch", (wb.calculation.calcMode or "auto") == "auto")
P("keine Makros (.xlsx)", not DATEI.lower().endswith(".xlsm"))
P("keine Verknüpfung zu anderen Dateien",
  not any("[" in str(c.value) for zeile in ws.iter_rows(min_row=ERSTE, max_row=min(LETZTE, ERSTE + 40))
          for c in zeile if str(c.value or "").startswith("=")))

print("\n" + ("═" * 46))
print("{} bestanden, {} fehlgeschlagen".format(ok, len(fehler)))
sys.exit(1 if fehler else 0)
