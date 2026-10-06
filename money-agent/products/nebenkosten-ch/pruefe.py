"""Rechenprobe: vergleicht die von LibreOffice berechnete Arbeitsmappe mit
einer unabhaengigen Python-Rechnung der Beispieldaten.

Aufruf: python3 pruefe.py <neu-berechnete.xlsx>
"""
import calendar
import sys
from datetime import date, timedelta

import openpyxl

from build import BEISPIEL_EINHEITEN, BEISPIEL_KOSTEN, BEISPIEL_MIETER, GRADTAGE

START = date(2025, 1, 1)
ENDE = date(2025, 12, 31)


def referenz():
    kosten = list(BEISPIEL_KOSTEN)
    kosten.append(("Verwaltungsaufwand 3 %", round(sum(k[1] for k in kosten) * 0.03, 2), "Fläche", "Linear"))
    einh = {nr: {"Fläche": fl, "Anteil": an, "Einheiten": 1} for nr, _, fl, an in BEISPIEL_EINHEITEN}
    tot = {s: sum(e[s] for e in einh.values()) for s in ("Fläche", "Anteil", "Einheiten")}
    tage = (ENDE - START).days + 1
    erg = []
    for name, nr, beg, end, akonto in BEISPIEL_MIETER:
        b, e = max(beg or START, START), min(end or ENDE, ENDE)
        lin = ((e - b).days + 1) / tage
        heiz = 0
        for mo in range(1, 13):
            ms = date(2025, mo, 1)
            me = date(2025, mo, calendar.monthrange(2025, mo)[1])
            belegt = max(0, (min(e, me) - max(b, ms)).days + 1) / me.day
            heiz += belegt * GRADTAGE[mo - 1]
        heiz /= sum(GRADTAGE)
        summe = sum(betrag * einh[nr][schl] * (heiz if vert == "Heizung" else lin) / tot[schl]
                    for _, betrag, schl, vert in kosten)
        total = round(summe * 20) / 20
        erg.append((name, total, akonto, total - akonto))
    return erg, sum(k[1] for k in kosten)


def main(pfad):
    wb = openpyxl.load_workbook(pfad, data_only=True)
    v, k = wb["Verteilung"], wb["Kontrolle"]
    erg, gesamt = referenz()
    fehler = 0
    for i, (name, total, akonto, saldo) in enumerate(erg):
        r = 5 + i
        xl_total, xl_saldo = v.cell(r, 28).value, v.cell(r, 30).value
        ok = abs(xl_total - total) < 0.001 and abs(xl_saldo - saldo) < 0.001
        fehler += not ok
        print(f"{'OK ' if ok else 'ERR'} {name:16} Excel {xl_total:10.2f} / Python {total:10.2f}  Saldo {xl_saldo:9.2f}")
    print("Kontrolle Total Kosten:", k["B4"].value, "erwartet", round(gesamt, 2))
    print("Kontrolle Differenz:", k["B7"].value, "Status:", k["B13"].value)
    fehler += k["B13"].value != "OK" or abs(k["B4"].value - gesamt) > 0.001
    b = wb["Abrechnung"]
    print("Abrechnungsblatt Mieter 1:", b["B5"].value, "|", b["B6"].value, "|", b["B8"].value, "| Saldo", b["G39"].value, b["E39"].value)
    print("FEHLER:", fehler)
    sys.exit(1 if fehler else 0)


if __name__ == "__main__":
    main(sys.argv[1])
