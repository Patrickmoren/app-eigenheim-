"""Testfaelle fuer die Nebenkostenabrechnung: fuellt die leere Vorlage,
laesst LibreOffice rechnen und vergleicht mit einer unabhaengigen Python-Rechnung.

Aufruf: python3 pruefe_faelle.py
"""
import calendar
import os
import subprocess
import sys
import tempfile
from datetime import date, timedelta

import openpyxl

from build import GRADTAGE

HIER = os.path.dirname(os.path.abspath(__file__))
E = [("1", 80, 300), ("2", 60, 200), ("3", 100, 500)]

FAELLE = {
    "einfach": dict(start=date(2025, 1, 1), einheiten=[("1", 80, 300)],
                    kosten=[("Heizung", 5000, "Fläche", "Heizung")],
                    mieter=[("A", "1", None, None, 4800)]),
    "mehrere Wohnungen, Schlüssel": dict(start=date(2025, 1, 1), einheiten=E,
                    kosten=[("Heizung", 9000, "Fläche", "Heizung"), ("Wasser", 3000, "Anteil", "Linear"),
                            ("Kehricht", 600, "Einheiten", "Linear")],
                    mieter=[("A", "1", None, None, 4000), ("B", "2", None, None, 3000), ("C", "3", None, None, 5000)]),
    "mehrere Wechsel + Leerstand": dict(start=date(2025, 1, 1), einheiten=E,
                    kosten=[("Heizung", 9000, "Fläche", "Heizung"), ("Hauswart", 2400, "Fläche", "Linear")],
                    mieter=[("A", "1", None, date(2025, 2, 28), 500), ("B", "1", date(2025, 3, 1), date(2025, 6, 15), 900),
                            ("C", "1", date(2025, 8, 1), None, 1200), ("D", "2", None, None, 3000)]),
    "Jahreswechsel + Schaltjahr": dict(start=date(2023, 7, 1), einheiten=E[:2],
                    kosten=[("Heizung", 7000, "Fläche", "Heizung"), ("Strom", 500, "Einheiten", "Linear")],
                    mieter=[("A", "1", date(2020, 1, 1), date(2024, 1, 31), 1500), ("B", "1", date(2024, 2, 1), None, 2000),
                            ("C", "2", None, date(2025, 12, 31), 3000)]),
    "Nullwerte + Gutschrift": dict(start=date(2025, 1, 1), einheiten=[("1", 80, 0), ("2", 0, 200)],
                    kosten=[("Heizung", 0, "Fläche", "Heizung"), ("Rückvergütung", -300, "Fläche", "Linear"),
                            ("Wasser", 1000, "Anteil", "Linear")],
                    mieter=[("A", "1", None, None, 0), ("B", "2", None, None, 500)]),
}
# Fehlerfaelle: erwarteter Status in der Kontrolle
FEHLER = {
    "unbekannte Einheit": dict(start=date(2025, 1, 1), einheiten=E[:1], kosten=[("Wasser", 1000, "Fläche", "Linear")],
                               mieter=[("A", "9", None, None, 0)]),
    "Auszug vor Einzug": dict(start=date(2025, 1, 1), einheiten=E[:1], kosten=[("Wasser", 1000, "Fläche", "Linear")],
                              mieter=[("A", "1", date(2025, 6, 1), date(2025, 3, 1), 0)]),
    "Überlappung": dict(start=date(2025, 1, 1), einheiten=E[:1], kosten=[("Wasser", 1000, "Fläche", "Linear")],
                        mieter=[("A", "1", None, None, 0), ("B", "1", date(2025, 6, 1), None, 0)]),
    "Schlüssel-Total 0": dict(start=date(2025, 1, 1), einheiten=[("1", 80, 0)], kosten=[("Wasser", 1000, "Anteil", "Linear")],
                              mieter=[("A", "1", None, None, 0)]),
    "Kosten ohne Schlüssel": dict(start=date(2025, 1, 1), einheiten=E[:1], kosten=[("Wasser", 1000, None, "Linear")],
                                  mieter=[("A", "1", None, None, 0)]),
}


def referenz(f):
    start = f["start"]
    ende = date(start.year + 1, start.month, 1) - timedelta(days=1)
    tage = (ende - start).days + 1
    einh = {nr: {"Fläche": fl, "Anteil": an, "Einheiten": 1} for nr, fl, an in f["einheiten"]}
    tot = {s: sum(e[s] for e in einh.values()) for s in ("Fläche", "Anteil", "Einheiten")}
    monate = [date(start.year + (start.month - 1 + j) // 12, (start.month - 1 + j) % 12 + 1, 1) for j in range(12)]
    gew = [GRADTAGE[m.month - 1] for m in monate]
    erg = []
    for name, nr, beg, end, ak in f["mieter"]:
        b, e = max(beg or start, start), min(end or ende, ende)
        lin = max(0, (e - b).days + 1) / tage
        heiz = 0
        for m, w in zip(monate, gew):
            me = date(m.year, m.month, calendar.monthrange(m.year, m.month)[1])
            heiz += max(0, (min(e, me) - max(b, m)).days + 1) / me.day * w
        heiz /= sum(GRADTAGE)
        s = sum(betrag * einh[nr][schl] * (heiz if vert == "Heizung" else lin) / tot[schl]
                for _, betrag, schl, vert in f["kosten"] if tot[schl])
        total = round(s * 20) / 20
        erg.append((name, total, total - ak))
    return erg


def fuelle(f, ziel):
    wb = openpyxl.load_workbook(os.path.join(HIER, "Nebenkostenabrechnung-CH.xlsx"))
    s, k, m = wb["Stammdaten"], wb["Kosten"], wb["Mieter"]
    s["B6"] = f["start"]
    for i, (nr, fl, an) in enumerate(f["einheiten"]):
        s.cell(22 + i, 1, nr); s.cell(22 + i, 3, fl); s.cell(22 + i, 4, an)
    for i, (art, betrag, schl, vert) in enumerate(f["kosten"]):
        k.cell(5 + i, 1, art); k.cell(5 + i, 2, betrag); k.cell(5 + i, 3).value = schl; k.cell(5 + i, 4, vert)
    for i, (name, nr, beg, end, ak) in enumerate(f["mieter"]):
        r = 5 + i
        m.cell(r, 2, name); m.cell(r, 3, nr); m.cell(r, 4).value = beg; m.cell(r, 5).value = end; m.cell(r, 6, ak)
    wb.save(ziel)


def rechne(faelle, t):
    for i, f in enumerate(faelle.values()):
        fuelle(f, os.path.join(t, f"f{i}.xlsx"))
    out = os.path.join(t, "out")
    os.makedirs(out, exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out]
                   + [os.path.join(t, f"f{i}.xlsx") for i in range(len(faelle))], check=True, capture_output=True)
    return [openpyxl.load_workbook(os.path.join(out, f"f{i}.xlsx"), data_only=True) for i in range(len(faelle))]


def main():
    subprocess.run([sys.executable, "build.py", "leer"], cwd=HIER, check=True, capture_output=True)
    fehler = 0
    with tempfile.TemporaryDirectory() as t:
        for (name, f), wb in zip(FAELLE.items(), rechne(FAELLE, t)):
            v, k = wb["Verteilung"], wb["Kontrolle"]
            for i, (mieter, total, saldo) in enumerate(referenz(f)):
                ist_t, ist_s = v.cell(5 + i, 28).value, v.cell(5 + i, 30).value
                ok = abs(ist_t - total) < 0.001 and abs(ist_s - saldo) < 0.001
                fehler += not ok
                print(f"{'OK ' if ok else 'ERR'} {name:30} {mieter}: Excel {ist_t:9.2f} Python {total:9.2f}")
            ok = k["B15"].value == "OK" and k["B7"].value == 0
            fehler += not ok
            print(f"{'OK ' if ok else 'ERR'} {name:30} Kontrolle {k['B15'].value}, Differenz {k['B7'].value}")
        for (name, f), wb in zip(FEHLER.items(), rechne(FEHLER, os.path.join(t))):
            status = wb["Kontrolle"]["B15"].value
            ok = status != "OK"
            fehler += not ok
            print(f"{'OK ' if ok else 'ERR'} Fehlerfall {name:22} erkannt: Status «{status}»")
    print("FEHLER:", fehler)
    sys.exit(1 if fehler else 0)


if __name__ == "__main__":
    main()
