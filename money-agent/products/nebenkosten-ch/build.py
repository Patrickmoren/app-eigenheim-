# -*- coding: utf-8 -*-
"""Erzeugt die Verkaufsvorlage «Nebenkostenabrechnung Schweiz» (Excel).

Zielgruppe: private Vermieter mit 1-30 Einheiten.
Eigenstaendiges Produkt, unabhaengig vom Leerstandsmanager im selben Repository.

Fachliche Regeln
- Abrechnungsperiode: 12 Monate ab dem Ersten eines Monats.
- Schluessel je Kostenart: Flaeche, Anteil (z.B. Wertquote/Volumen) oder Einheiten.
- Verteilung je Kostenart: «Linear» (taggenau) oder «Heizung» (Monatsgewichte,
  damit ein Auszug im Sommer nicht gleich viel Heizkosten traegt wie im Winter).
- Leerstand: der nicht belegte Anteil bleibt beim Eigentuemer und wird
  ausgewiesen, statt auf die uebrigen Mieter abgewaelzt zu werden.

Aufruf: python3 build.py [beispiel|leer]   (ohne Argument: beide)
"""
import sys
from datetime import date

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as col

ANZ_EINHEITEN = 30
ANZ_KOSTEN = 25
ANZ_MIETER = 40

TINTE, AKZENT, HELL, EINGABE, RAND = "1A1A1A", "1F4E79", "EAF1F8", "FFF8DC", "BFC9D3"
duenn = Side(style="thin", color=RAND)
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
CHF = '#,##0.00'
DATUM = 'DD.MM.YYYY'


def schrift(g=10, fett=False, farbe=TINTE, kursiv=False):
    return Font(name="Arial", size=g, bold=fett, color=farbe, italic=kursiv)


def titel(ws, text, unter=None):
    ws["A1"] = text
    ws["A1"].font = schrift(16, True, AKZENT)
    if unter:
        ws["A2"] = unter
        ws["A2"].font = schrift(9, kursiv=True, farbe="555555")


def kopf(ws, zeile, texte, start=1):
    for i, t in enumerate(texte):
        c = ws.cell(zeile, start + i, t)
        c.font = schrift(10, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AKZENT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = rahmen


def eingabe(c, fmt=None):
    c.fill = PatternFill("solid", fgColor=EINGABE)
    c.protection = Protection(locked=False)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def formel(c, f, fmt=None):
    c.value = f
    c.fill = PatternFill("solid", fgColor=HELL)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def schuetzen(ws):
    ws.protection.sheet = True          # ohne Passwort: Nutzer kann bei Bedarf aufheben
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


BEISPIEL_EINHEITEN = [
    ("1", "EG links, 3.5 Zi.", 82, 250),
    ("2", "EG rechts, 2.5 Zi.", 61, 190),
    ("3", "1. OG, 4.5 Zi.", 104, 320),
    ("4", "DG, 3.5 Zi.", 78, 240),
]
BEISPIEL_KOSTEN = [
    ("Heizung und Warmwasser (Gas)", 9600, "Fläche", "Heizung"),
    ("Hauswartung", 3600, "Fläche", "Linear"),
    ("Treppenhausreinigung", 1200, "Einheiten", "Linear"),
    ("Allgemeinstrom", 640, "Anteil", "Linear"),
    ("Wasser und Abwasser", 2350, "Anteil", "Linear"),
    ("Kehrichtgrundgebühr", 480, "Einheiten", "Linear"),
    ("Service Heizanlage", 520, "Fläche", "Heizung"),
]
# Mieterwechsel in Einheit 2, Leerstand in Einheit 4 (Okt-Nov)
BEISPIEL_MIETER = [
    ("Muster Anna", "1", None, None, 4600),
    ("Beispiel Marco", "2", None, date(2025, 3, 31), 1200),
    ("Probst Lea", "2", date(2025, 4, 1), None, 2400),
    ("Keller Familie", "3", None, None, 5800),
    ("Weber Tim", "4", None, date(2025, 9, 30), 3150),
    ("Frei Sara", "4", date(2025, 12, 1), None, 600),
]
# Monatsgewichte in Promille ab Januar (Standardwerte, editierbar)
GRADTAGE = [170, 150, 130, 90, 40, 10, 0, 0, 30, 80, 140, 160]


def baue(beispiel=True):
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------ Anleitung
    a = wb.active
    a.title = "Anleitung"
    titel(a, "Nebenkostenabrechnung Schweiz", "Vorlage für private Vermieter · ohne Makros · Excel und LibreOffice")
    schritte = [
        "So gehen Sie vor",
        "1. «Stammdaten»: Liegenschaft, Beginn der Abrechnungsperiode (immer der 1. eines Monats) und Ihre Einheiten mit Fläche und Anteil erfassen.",
        "2. «Kosten»: je Kostenart Betrag, Verteilschlüssel und Verteilungsart wählen. «Heizung» verteilt nach Monatsgewichten (Heizgradtage), «Linear» taggenau.",
        "3. «Mieter»: alle Mietverhältnisse der Periode erfassen, auch ausgezogene. Leeres Mietende = Mietverhältnis läuft weiter.",
        "4. «Abrechnung»: rechts oben die Mieter-Nr. wählen und das Blatt drucken oder als PDF speichern. Eine Seite pro Mieter.",
        "5. «Kontrolle»: zeigt, dass jeder Franken verteilt ist. Leerstand bleibt beim Eigentümer und wird separat ausgewiesen.",
        "",
        "Gelbe Felder sind Eingaben. Blaue Felder rechnen automatisch und sind geschützt (Blattschutz ohne Passwort).",
        "",
        "Rechtliche Hinweise (Stand der Vorlage, ohne Gewähr)",
        "• Nebenkosten dürfen nur separat verrechnet werden, wenn sie im Mietvertrag einzeln vereinbart sind (Art. 257a OR).",
        "• Der Mieter kann Einsicht in die Belege verlangen (Art. 257b Abs. 2 OR).",
        "• Heiz- und Warmwasserkosten: anrechenbare Positionen siehe Art. 5 und 6 VMWG; Unterhalt und Reparaturen sind nicht anrechenbar.",
        "• Die Monatsgewichte der Heizung sind Standardwerte. Wenn Ihr Vertrag eine andere Tabelle vorsieht, Werte in «Stammdaten» anpassen.",
        "• Diese Vorlage ersetzt keine Rechtsberatung. Bei Streitfällen: Schlichtungsbehörde Ihres Kantons oder Verband (HEV / SVIT).",
    ]
    for i, t in enumerate(schritte, 4):
        a.cell(i, 1, t).font = schrift(11 if t in ("So gehen Sie vor", "Rechtliche Hinweise (Stand der Vorlage, ohne Gewähr)") else 10,
                                       fett=t in ("So gehen Sie vor", "Rechtliche Hinweise (Stand der Vorlage, ohne Gewähr)"))
    a.column_dimensions["A"].width = 130

    # ------------------------------------------------------------ Stammdaten
    s = wb.create_sheet("Stammdaten")
    titel(s, "Stammdaten")
    felder = [("Liegenschaft", "Musterstrasse 12, 8004 Zürich" if beispiel else None, None),
              ("Vermieter / Verwaltung", "Hans Beispiel, Seeweg 3, 8800 Thalwil" if beispiel else None, None),
              ("Beginn Abrechnungsperiode", date(2025, 1, 1) if beispiel else None, DATUM)]
    for i, (k, v, fmt) in enumerate(felder, 4):
        s.cell(i, 1, k).font = schrift(10, True)
        c = s.cell(i, 2, v)
        eingabe(c, fmt)
    s["A7"] = "Ende Abrechnungsperiode"
    s["A7"].font = schrift(10, True)
    formel(s["B7"], '=IF(B6="","",EDATE(B6,12)-1)', DATUM)
    s["A8"] = "Tage in der Periode"
    s["A8"].font = schrift(10, True)
    formel(s["B8"], '=IF(B6="","",B7-B6+1)', "0")
    dv_tag = DataValidation(type="custom", formula1="DAY(B6)=1", showErrorMessage=True,
                            errorTitle="Periodenbeginn", error="Die Periode beginnt immer am 1. eines Monats.")
    s.add_data_validation(dv_tag)
    dv_tag.add("B6")

    # Heizgewichte, Spalten D:E ab Zeile 4
    s["D3"] = "Heizung: Monatsgewichte (Promille)"
    s["D3"].font = schrift(10, True, AKZENT)
    kopf(s, 4, ["Monat", "Gewicht ‰"], start=4)
    monate = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
              "September", "Oktober", "November", "Dezember"]
    for m in range(12):
        s.cell(5 + m, 4, monate[m]).font = schrift()
        s.cell(5 + m, 4).border = rahmen
        eingabe(s.cell(5 + m, 5, GRADTAGE[m]), "0")
    s["D17"] = "Summe"
    s["D17"].font = schrift(10, True)
    formel(s["E17"], "=SUM(E5:E16)", "0")

    # Einheiten ab Zeile 21
    s["A20"] = "Einheiten"
    s["A20"].font = schrift(12, True, AKZENT)
    kopf(s, 21, ["Nr.", "Bezeichnung", "Fläche m²", "Anteil (Wertquote, ‰ o.ä.)", "Zähler Einheiten"])
    E0 = 22
    for r in range(E0, E0 + ANZ_EINHEITEN):
        for c_ in range(1, 5):
            eingabe(s.cell(r, c_), "0.00" if c_ in (3, 4) else None)
        formel(s.cell(r, 5), f'=IF(A{r}="","",1)', "0")
    if beispiel:
        for i, (nr, bez, fl, ant) in enumerate(BEISPIEL_EINHEITEN):
            r = E0 + i
            s.cell(r, 1, nr); s.cell(r, 2, bez); s.cell(r, 3, fl); s.cell(r, 4, ant)
    EL = E0 + ANZ_EINHEITEN - 1
    s.cell(EL + 1, 2, "Total").font = schrift(10, True)
    for c_ in (3, 4, 5):
        L = col(c_)
        formel(s.cell(EL + 1, c_), f"=SUM({L}{E0}:{L}{EL})", "0.00")
    for L, w in zip("ABCDE", (26, 34, 14, 22, 14)):
        s.column_dimensions[L].width = w
    s.column_dimensions["D"].width = 22
    schuetzen(s)
    EINH = f"Stammdaten!$A${E0}:$A${EL}"
    TOT = {"Fläche": f"Stammdaten!$C${EL+1}", "Anteil": f"Stammdaten!$D${EL+1}",
           "Einheiten": f"Stammdaten!$E${EL+1}"}

    # ------------------------------------------------------------ Kosten
    k = wb.create_sheet("Kosten")
    titel(k, "Kosten der Periode", "Nur Kostenarten erfassen, die im Mietvertrag als Nebenkosten vereinbart sind.")
    kopf(k, 4, ["Kostenart", "Betrag CHF", "Schlüssel", "Verteilung", "Total Schlüssel"])
    K0 = 5
    dv_s = DataValidation(type="list", formula1='"Fläche,Anteil,Einheiten"', allow_blank=True)
    dv_v = DataValidation(type="list", formula1='"Linear,Heizung"', allow_blank=True)
    k.add_data_validation(dv_s); k.add_data_validation(dv_v)
    for r in range(K0, K0 + ANZ_KOSTEN):
        eingabe(k.cell(r, 1)); eingabe(k.cell(r, 2), CHF); eingabe(k.cell(r, 3)); eingabe(k.cell(r, 4))
        dv_s.add(k.cell(r, 3)); dv_v.add(k.cell(r, 4))
        formel(k.cell(r, 5), f'=IF(C{r}="Fläche",{TOT["Fläche"]},IF(C{r}="Anteil",{TOT["Anteil"]},'
                             f'IF(C{r}="Einheiten",{TOT["Einheiten"]},0)))', "0.00")
    if beispiel:
        for i, (art, betrag, schl, vert) in enumerate(BEISPIEL_KOSTEN):
            r = K0 + i
            k.cell(r, 1, art); k.cell(r, 2, betrag); k.cell(r, 3, schl); k.cell(r, 4, vert)
        r = K0 + len(BEISPIEL_KOSTEN)
        k.cell(r, 1, "Verwaltungsaufwand 3 %")
        k.cell(r, 2, f"=ROUND(SUM(B{K0}:B{r-1})*3%,2)")
        k.cell(r, 3, "Fläche"); k.cell(r, 4, "Linear")
    KL = K0 + ANZ_KOSTEN - 1
    k.cell(KL + 1, 1, "Total").font = schrift(10, True)
    formel(k.cell(KL + 1, 2), f"=SUM(B{K0}:B{KL})", CHF)
    for L, w in zip("ABCDE", (40, 16, 14, 14, 16)):
        k.column_dimensions[L].width = w
    schuetzen(k)

    # ------------------------------------------------------------ Mieter
    m = wb.create_sheet("Mieter")
    titel(m, "Mietverhältnisse der Periode", "Gelb erfassen. Hellblaue Spalten sind Hilfsrechnungen.")
    kopf(m, 4, ["Nr.", "Mieter", "Einheit Nr.", "Mietbeginn", "Mietende", "Akonto bezahlt CHF",
                "Beginn in Periode", "Ende in Periode", "Tage", "Faktor linear", "Faktor Heizung",
                "Fläche", "Anteil", "Einheiten"] + [f"Belegung M{i+1}" for i in range(12)])
    M0 = 5
    dv_e = DataValidation(type="list", formula1=EINH, allow_blank=True)
    m.add_data_validation(dv_e)
    for i in range(ANZ_MIETER):
        r = M0 + i
        formel(m.cell(r, 1), i + 1, "0")
        eingabe(m.cell(r, 2)); eingabe(m.cell(r, 3)); eingabe(m.cell(r, 4), DATUM)
        eingabe(m.cell(r, 5), DATUM); eingabe(m.cell(r, 6), CHF)
        dv_e.add(m.cell(r, 3))
        aktiv = f'C{r}<>""'
        formel(m.cell(r, 7), f'=IF({aktiv},MAX(IF(D{r}="",Stammdaten!$B$6,D{r}),Stammdaten!$B$6),"")', DATUM)
        formel(m.cell(r, 8), f'=IF({aktiv},MIN(IF(E{r}="",Stammdaten!$B$7,E{r}),Stammdaten!$B$7),"")', DATUM)
        formel(m.cell(r, 9), f'=IF({aktiv},MAX(0,H{r}-G{r}+1),0)', "0")
        formel(m.cell(r, 10), f'=IF({aktiv},I{r}/Stammdaten!$B$8,0)', "0.0000")
        formel(m.cell(r, 11), f'=IF({aktiv},SUMPRODUCT(O{r}:Z{r},$O$3:$Z$3)/Stammdaten!$E$17,0)', "0.0000")
        for c_, spalte in ((12, "C"), (13, "D"), (14, "E")):
            formel(m.cell(r, c_), f'=IF({aktiv},IFERROR(INDEX(Stammdaten!${spalte}${E0}:${spalte}${EL},'
                                  f'MATCH(C{r},{EINH},0)),0),0)', "0.00")
        for j in range(12):
            L = col(15 + j)
            ms = f"EDATE(Stammdaten!$B$6,{j})"
            formel(m.cell(r, 15 + j),
                   f'=IF({aktiv},MAX(0,MIN(H{r},EOMONTH({ms},0))-MAX(G{r},{ms})+1)/DAY(EOMONTH({ms},0)),0)',
                   "0.00")
    # Zeile 3: Heizgewicht des jeweiligen Periodenmonats
    m.cell(3, 14, "Gewicht ‰ →").font = schrift(9, kursiv=True)
    for j in range(12):
        formel(m.cell(3, 15 + j), f'=INDEX(Stammdaten!$E$5:$E$16,MONTH(EDATE(Stammdaten!$B$6,{j})))', "0")
    if beispiel:
        for i, (name, einh, beg, ende, ak) in enumerate(BEISPIEL_MIETER):
            r = M0 + i
            m.cell(r, 2, name); m.cell(r, 3, einh); m.cell(r, 4, beg); m.cell(r, 5, ende); m.cell(r, 6, ak)
    ML = M0 + ANZ_MIETER - 1
    m.column_dimensions["A"].width = 6
    m.column_dimensions["B"].width = 26
    for L in "CDEFGH":
        m.column_dimensions[L].width = 13
    m.freeze_panes = "C5"
    # fehlerhafte Eingaben markieren: Einheit unbekannt oder Ende vor Beginn
    m.conditional_formatting.add(f"C{M0}:C{ML}", FormulaRule(
        formula=[f'AND(C{M0}<>"",ISNA(MATCH(C{M0},{EINH},0)))'], fill=PatternFill("solid", fgColor="F8CBAD")))
    m.conditional_formatting.add(f"E{M0}:E{ML}", FormulaRule(
        formula=[f'AND(E{M0}<>"",D{M0}<>"",E{M0}<D{M0})'], fill=PatternFill("solid", fgColor="F8CBAD")))
    schuetzen(m)

    # ------------------------------------------------------------ Verteilung
    v = wb.create_sheet("Verteilung")
    titel(v, "Verteilung je Mieter und Kostenart", "Vollständig berechnet. Keine Eingaben nötig.")
    kopf(v, 4, ["Nr.", "Mieter"] + [f"K{i+1}" for i in range(ANZ_KOSTEN)] + ["Total", "Akonto", "Saldo"])
    for i in range(ANZ_KOSTEN):
        formel(v.cell(3, 3 + i), f'=IF(Kosten!A{K0+i}="","",Kosten!A{K0+i})')
        v.cell(3, 3 + i).alignment = Alignment(wrap_text=True, vertical="bottom")
    v.row_dimensions[3].height = 60
    V0 = 5
    for i in range(ANZ_MIETER):
        r, mr = V0 + i, M0 + i
        formel(v.cell(r, 1), f"=Mieter!A{mr}", "0")
        formel(v.cell(r, 2), f'=IF(Mieter!C{mr}="","",Mieter!B{mr})')
        for j in range(ANZ_KOSTEN):
            kr = K0 + j
            schl = (f'IF(Kosten!$C${kr}="Fläche",Mieter!$L{mr},IF(Kosten!$C${kr}="Anteil",Mieter!$M{mr},'
                    f'Mieter!$N{mr}))')
            fak = f'IF(Kosten!$D${kr}="Heizung",Mieter!$K{mr},Mieter!$J{mr})'
            formel(v.cell(r, 3 + j),
                   f'=IF(OR(Kosten!$B${kr}="",Kosten!$E${kr}=0,Mieter!$C{mr}=""),0,'
                   f'Kosten!$B${kr}*{schl}*{fak}/Kosten!$E${kr})', CHF)
        tc = 3 + ANZ_KOSTEN
        formel(v.cell(r, tc), f"=ROUND(SUM({col(3)}{r}:{col(tc-1)}{r})*20,0)/20", CHF)   # auf 5 Rappen
        formel(v.cell(r, tc + 1), f"=N(Mieter!F{mr})", CHF)
        formel(v.cell(r, tc + 2), f"={col(tc)}{r}-{col(tc+1)}{r}", CHF)
    VL = V0 + ANZ_MIETER - 1
    v.cell(VL + 1, 2, "Summe Mieter").font = schrift(10, True)
    v.cell(VL + 2, 2, "Leerstand / Eigentümer").font = schrift(10, True)
    for j in range(ANZ_KOSTEN + 1):
        L = col(3 + j)
        formel(v.cell(VL + 1, 3 + j), f"=SUM({L}{V0}:{L}{VL})", CHF)
    for j in range(ANZ_KOSTEN):
        L = col(3 + j)
        formel(v.cell(VL + 2, 3 + j), f"=N(Kosten!B{K0+j})-{L}{VL+1}", CHF)
    v.column_dimensions["B"].width = 24
    for j in range(ANZ_KOSTEN + 3):
        v.column_dimensions[col(3 + j)].width = 12
    v.freeze_panes = "C5"
    schuetzen(v)
    TOTAL_C, AKONTO_C, SALDO_C = col(3 + ANZ_KOSTEN), col(4 + ANZ_KOSTEN), col(5 + ANZ_KOSTEN)

    # ------------------------------------------------------------ Abrechnung
    b = wb.create_sheet("Abrechnung")
    b["A1"] = "Heiz- und Nebenkostenabrechnung"
    b["A1"].font = schrift(16, True, AKZENT)
    b["I1"] = "Mieter-Nr. wählen (wird nicht gedruckt):"
    b["I1"].font = schrift(9, kursiv=True)
    b["I2"] = 1
    eingabe(b["I2"], "0")
    dv_n = DataValidation(type="whole", operator="between", formula1="1", formula2=str(ANZ_MIETER))
    b.add_data_validation(dv_n); dv_n.add("I2")
    zeile = f"MATCH($I$2,Mieter!$A${M0}:$A${ML},0)"
    info = [("Liegenschaft", "=Stammdaten!B4", None),
            ("Vermieter", "=Stammdaten!B5", None),
            ("Mieter", f"=INDEX(Mieter!$B${M0}:$B${ML},{zeile})", None),
            ("Einheit", f'=INDEX(Mieter!$C${M0}:$C${ML},{zeile})&" · "&IFERROR(INDEX(Stammdaten!$B${E0}:$B${EL},'
                        f'MATCH(INDEX(Mieter!$C${M0}:$C${ML},{zeile}),{EINH},0)),"")', None),
            ("Abrechnungsperiode", '=TEXT(Stammdaten!B6,"DD.MM.YYYY")&" – "&TEXT(Stammdaten!B7,"DD.MM.YYYY")', None),
            ("Ihre Mietdauer in der Periode",
             f'=TEXT(INDEX(Mieter!$G${M0}:$G${ML},{zeile}),"DD.MM.YYYY")&" – "&'
             f'TEXT(INDEX(Mieter!$H${M0}:$H${ML},{zeile}),"DD.MM.YYYY")&" ("&INDEX(Mieter!$I${M0}:$I${ML},{zeile})&" Tage)"',
             None)]
    for i, (lab, f, _) in enumerate(info, 3):
        b.cell(i, 1, lab).font = schrift(10, True)
        b.cell(i, 2, f).font = schrift()
    kopf(b, 10, ["Kostenart", "Gesamtkosten CHF", "Schlüssel", "Verteilung", "Ihr Anteil am Schlüssel",
                 "Zeitanteil", "Ihr Betrag CHF"])
    B0 = 11
    for j in range(ANZ_KOSTEN):
        r, kr = B0 + j, K0 + j
        vz = f"INDEX(Verteilung!{col(3+j)}${V0}:{col(3+j)}${VL},{zeile})"
        b.cell(r, 1, f'=IF(Kosten!A{kr}="","",Kosten!A{kr})')
        b.cell(r, 2, f'=IF(Kosten!A{kr}="","",Kosten!B{kr})').number_format = CHF
        b.cell(r, 3, f'=IF(Kosten!A{kr}="","",Kosten!C{kr})')
        b.cell(r, 4, f'=IF(Kosten!A{kr}="","",Kosten!D{kr})')
        anteil = (f'IF(Kosten!C{kr}="Fläche",INDEX(Mieter!$L${M0}:$L${ML},{zeile}),'
                  f'IF(Kosten!C{kr}="Anteil",INDEX(Mieter!$M${M0}:$M${ML},{zeile}),INDEX(Mieter!$N${M0}:$N${ML},{zeile})))')
        b.cell(r, 5, f'=IF(OR(Kosten!A{kr}="",Kosten!E{kr}=0),"",TEXT({anteil},"0.##")&" / "&TEXT(Kosten!E{kr},"0.##"))')
        zeit = (f'IF(Kosten!D{kr}="Heizung",INDEX(Mieter!$K${M0}:$K${ML},{zeile}),'
                f'INDEX(Mieter!$J${M0}:$J${ML},{zeile}))')
        b.cell(r, 6, f'=IF(Kosten!A{kr}="","",{zeit})').number_format = "0.0%"
        b.cell(r, 7, f'=IF(Kosten!A{kr}="","",{vz})').number_format = CHF
        for c_ in range(1, 8):
            b.cell(r, c_).font = schrift()
    BL = B0 + ANZ_KOSTEN - 1
    s1, s2, s3 = BL + 2, BL + 3, BL + 4
    b.cell(s1, 5, "Total Ihr Anteil (gerundet auf 5 Rp.)").font = schrift(10, True)
    b.cell(s1, 7, f"=INDEX(Verteilung!{TOTAL_C}${V0}:{TOTAL_C}${VL},{zeile})").number_format = CHF
    b.cell(s2, 5, "abzüglich Akontozahlungen").font = schrift(10)
    b.cell(s2, 7, f"=-INDEX(Verteilung!{AKONTO_C}${V0}:{AKONTO_C}${VL},{zeile})").number_format = CHF
    saldo = f"INDEX(Verteilung!{SALDO_C}${V0}:{SALDO_C}${VL},{zeile})"
    b.cell(s3, 5, f'=IF({saldo}>=0,"Nachzahlung zu Ihren Lasten","Guthaben zu Ihren Gunsten")').font = schrift(11, True)
    b.cell(s3, 7, f"=ABS({saldo})").number_format = CHF
    b.cell(s3, 7).font = schrift(11, True)
    b.cell(s3 + 2, 1, "Die Belege können Sie nach Vereinbarung einsehen (Art. 257b Abs. 2 OR). "
                      "Einwände bitte innert 30 Tagen schriftlich.").font = schrift(9, kursiv=True)
    b.cell(s3 + 3, 1, "Nachzahlung zahlbar innert 30 Tagen. Ein Guthaben wird mit der nächsten Miete verrechnet "
                      "oder überwiesen.").font = schrift(9, kursiv=True)
    # leere Kostenzeilen beim Druck ausblenden geht ohne Makro nicht; Druckbereich deckt alle Zeilen ab
    for L, w in zip("ABCDEFGHI", (34, 14, 11, 11, 22, 11, 15, 2, 34)):
        b.column_dimensions[L].width = w
    b.print_area = f"A1:G{s3+3}"
    b.page_setup.orientation = "portrait"
    b.page_setup.fitToWidth = 1
    b.page_setup.fitToHeight = 1
    b.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(b)

    # ------------------------------------------------------------ Kontrolle
    c = wb.create_sheet("Kontrolle")
    titel(c, "Kontrolle", "Jeder Franken muss entweder einem Mieter oder dem Leerstand zugeordnet sein.")
    rows = [("Total Kosten", f"=Kosten!B{KL+1}"),
            ("Verteilt an Mieter (ungerundet)", f"=SUM(Verteilung!C{VL+1}:{col(2+ANZ_KOSTEN)}{VL+1})"),
            ("Leerstand / Eigentümer", f"=SUM(Verteilung!C{VL+2}:{col(2+ANZ_KOSTEN)}{VL+2})"),
            ("Differenz (muss 0 sein)", "=ROUND(B4-B5-B6,2)"),
            ("Summe Monatsgewichte (soll 1000)", "=Stammdaten!E17"),
            ("Kostenzeilen ohne Schlüssel/Verteilung",
             f'=SUMPRODUCT((Kosten!B{K0}:B{KL}<>"")*((Kosten!C{K0}:C{KL}="")+(Kosten!D{K0}:D{KL}="")))'),
            ("Mieterzeilen mit unbekannter Einheit",
             f'=SUMPRODUCT((Mieter!C{M0}:C{ML}<>"")*(Mieter!L{M0}:L{ML}+Mieter!N{M0}:N{ML}=0))'),
            ("Überlappende Mietverhältnisse (Tage > Periode je Einheit)",
             f'=SUMPRODUCT(--(SUMIF(Mieter!C{M0}:C{ML},Stammdaten!A{E0}:A{EL},Mieter!I{M0}:I{ML})>Stammdaten!B8))')]
    for i, (lab, f) in enumerate(rows, 4):
        c.cell(i, 1, lab).font = schrift(10, True)
        formel(c.cell(i, 2), f, CHF if i <= 7 else "0")
    c.cell(13, 1, "Status").font = schrift(11, True)
    formel(c.cell(13, 2), '=IF(AND(B7=0,B8=1000,B9=0,B10=0,B11=0),"OK","Bitte prüfen")')
    c.conditional_formatting.add("B13", FormulaRule(formula=['B13="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    c.conditional_formatting.add("B13", FormulaRule(formula=['B13<>"OK"'], fill=PatternFill("solid", fgColor="F8CBAD")))
    c.column_dimensions["A"].width = 52
    c.column_dimensions["B"].width = 18
    schuetzen(c)

    wb.active = 0
    name = "Nebenkostenabrechnung-CH-Beispiel.xlsx" if beispiel else "Nebenkostenabrechnung-CH.xlsx"
    wb.save(name)
    print("geschrieben:", name)


if __name__ == "__main__":
    art = sys.argv[1] if len(sys.argv) > 1 else "beide"
    if art in ("beispiel", "beide"):
        baue(True)
    if art in ("leer", "beide"):
        baue(False)
