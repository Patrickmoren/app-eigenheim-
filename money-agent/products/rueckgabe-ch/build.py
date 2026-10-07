# -*- coding: utf-8 -*-
"""Erzeugt das Modul «Wohnungsrückgabe: Schaden & Kaution» (Excel).

Vom Abnahmeprotokoll zur fertigen Kautionsabrechnung:
- Schäden erfassen mit Ursache (Mieter / normale Abnutzung / vorbestehend)
- Mieteranteil nach Lebensdauer: Kosten × max(0, 1 − Alter / Lebensdauer)
  (Methode der paritätischen Lebensdauertabelle HEV/MV; die Lebensdauer trägt
  der Nutzer aus der Tabelle ein, die Tabelle selbst ist nicht enthalten)
- Kautionsabrechnung und Freigabe-Erklärung für die Bank

Bewusst ohne TEXT(), FILTER() und Matrixformeln, damit die Datei auch in
LibreOffice und Apple Numbers rechnet.
Aufruf: python3 build.py [beispiel|leer]   (ohne Argument: beide)
"""
import sys
from datetime import date

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

ANZ = 25
S0 = 24                      # erste Schadenzeile
SL = S0 + ANZ - 1
TINTE, AKZENT, HELL, EINGABE, RAND = "1A1A1A", "1F4E79", "EAF1F8", "FFF8DC", "BFC9D3"
duenn = Side(style="thin", color=RAND)
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
DATUM, CHF = "DD.MM.YYYY", "#,##0.00"
URSACHEN = ("Mieter", "Normale Abnutzung", "Vorbestehend")

KOPF = [  # Zeile, Bezeichnung, Beispielwert, Format
    (4, "Liegenschaft", "Musterstrasse 12, 8004 Zürich", None),
    (5, "Wohnung", "1. OG links, 3.5 Zimmer", None),
    (6, "Mieter", "Muster Anna", None),
    (7, "Neue Adresse Mieter", "Seeweg 3, 8800 Thalwil", None),
    (8, "Vermieter", "Hans Beispiel, Bergstrasse 1, 8004 Zürich", None),
    (9, "Mietende", date(2026, 9, 30), DATUM),
    (10, "Abnahme am", date(2026, 9, 30), DATUM),
    (11, "Mietkaution (Betrag)", 4950, CHF),
    (12, "Aufgelaufener Kautionszins", 12.40, CHF),
    (13, "Bank / Kautionskonto", "Musterbank, Konto 123.456.789", None),
]
FORDERUNGEN = [  # Zeile, Bezeichnung, Beispielwert
    (16, "Ausstehende Mietzinse", 0),
    (17, "Nebenkosten-Saldo (+ Nachzahlung / − Guthaben)", 168.35),
    (18, "Weitere Forderungen (z.B. Reinigung, Entsorgung)", 0),
]
BEISPIEL = [  # Raum, Bauteil, Mangel, Ursache, Kosten, Einbau, Lebensdauer
    ("Wohnzimmer", "Wandanstrich", "Starke Nikotinverfärbung", "Mieter", 1800, date(2019, 4, 1), 8),
    ("Küche", "Backofen", "Türglas gesprungen", "Mieter", 950, date(2014, 6, 1), 15),
    ("Wohnzimmer", "Parkett", "Wasserflecken vor Balkontür", "Mieter", 2400, date(2005, 1, 1), 40),
    ("Schlafzimmer", "Teppich", "Laufspuren", "Normale Abnutzung", 1200, date(2016, 3, 1), 10),
    ("Bad", "Spiegelschrank", "Kratzer (im Eintrittsprotokoll vermerkt)", "Vorbestehend", 300, None, None),
    ("Eingang", "Schlüssel", "1 Wohnungsschlüssel fehlt", "Mieter", 180, None, None),
]


def schrift(g=10, fett=False, farbe=TINTE, kursiv=False):
    return Font(name="Arial", size=g, bold=fett, color=farbe, italic=kursiv)


def eingabe(c, fmt=None):
    c.fill = PatternFill("solid", fgColor=EINGABE)
    c.protection = Protection(locked=False)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def formel(c, f, fmt=None, fett=False):
    c.value = f
    c.fill = PatternFill("solid", fgColor=HELL)
    c.border = rahmen
    c.font = schrift(10, fett)
    if fmt:
        c.number_format = fmt


def kopf(ws, zeile, texte, start=1):
    for i, t in enumerate(texte):
        c = ws.cell(zeile, start + i, t)
        c.font = schrift(9, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AKZENT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = rahmen


def text(ws, zelle, t, g=10, fett=False, kursiv=False, farbe=TINTE):
    ws[zelle] = t
    ws[zelle].font = schrift(g, fett, farbe, kursiv)
    ws[zelle].alignment = Alignment(wrap_text=False)


def baue(beispiel=True):
    wb = openpyxl.Workbook()

    # ------------------------------------------------------------ Start
    a = wb.active
    a.title = "Start"
    text(a, "A1", "Wohnungsrückgabe: Schaden & Kaution", 16, True, farbe=AKZENT)
    zeilen = [
        ("START – was Sie eingeben", True),
        ("Blatt «Rückgabe»: Mietende, Kaution und offene Beträge. Dann jeden Mangel aus dem Abnahmeprotokoll: Kosten laut Offerte, Einbaudatum und Lebensdauer.", False),
        ("BERECHNUNG – was die Vorlage macht", True),
        ("Der Mieter zahlt nur den Restwert: Kosten × (1 − Alter ÷ Lebensdauer). Ist ein Bauteil älter als seine Lebensdauer, zahlt er nichts. Normale Abnutzung und vorbestehende Mängel gehen nie zu seinen Lasten.", False),
        ("ERGEBNIS – was Sie jetzt tun", True),
        ("«Abrechnung» zeigt, wie viel von der Kaution Ihnen zusteht und wie viel der Mieter zurückerhält. «Kontrolle» meldet fehlende oder unplausible Angaben.", False),
        ("AUSGABE – was Sie versenden", True),
        ("«Abrechnung» an den Mieter senden. «Freigabe Kaution» von beiden Parteien unterschreiben lassen und der Bank schicken.", False),
        ("", False),
        ("Wo finde ich die Lebensdauer?", True),
        ("In der paritätischen Lebensdauertabelle von Hauseigentümerverband (HEV) und Mieterverband (MV). Ein Auszug ist online gratis, die vollständige Tabelle verkaufen beide Verbände.", False),
        ("Leere Lebensdauer = keine Altersentwertung (z.B. fehlender Schlüssel, Reinigung). Die Beispielwerte im Beispiel sind vor Gebrauch in der Tabelle zu prüfen.", False),
        ("", False),
        ("Rechtliche Hinweise (ohne Gewähr, ersetzt keine Rechtsberatung)", True),
        ("• Der Mieter haftet nicht für Abnutzung durch vertragsgemässen Gebrauch (Art. 267 Abs. 1 OR).", False),
        ("• Mängel bei der Rückgabe sofort und konkret rügen, sonst sind die Ansprüche verwirkt – ausser bei versteckten Mängeln (Art. 267a OR).", False),
        ("• Die Lebensdauertabelle ist kein Gesetz, wird aber von Schlichtungsbehörden und Gerichten regelmässig angewendet.", False),
        ("• Kaution: Die Bank zahlt mit Zustimmung beider Parteien aus. Macht der Vermieter innert eines Jahres nach Mietende keinen Anspruch geltend, kann der Mieter die Kaution allein herausverlangen (Art. 257e Abs. 3 OR).", False),
        ("• Berechnen Sie Kosten nur mit Offerten oder Rechnungen. Pauschalen ohne Beleg werden häufig nicht anerkannt.", False),
    ]
    for i, (t, fett) in enumerate(zeilen, 3):
        a.cell(i, 1, t).font = schrift(11 if fett else 10, fett, AKZENT if fett else TINTE)
    a.column_dimensions["A"].width = 150

    # ------------------------------------------------------------ Rückgabe
    r = wb.create_sheet("Rückgabe")
    text(r, "A1", "Wohnungsrückgabe", 16, True, farbe=AKZENT)
    text(r, "A2", "Gelb = Eingabe · Blau = rechnet automatisch", 9, kursiv=True, farbe="555555")
    for z, lab, bsp, fmt in KOPF:
        r.cell(z, 1, lab).font = schrift(10, True)
        c = r.cell(z, 2, bsp if beispiel else None)
        eingabe(c, fmt)
    text(r, "A15", "Weitere Forderungen an den Mieter", 11, True, farbe=AKZENT)
    for z, lab, bsp in FORDERUNGEN:
        r.cell(z, 1, lab).font = schrift(10)
        eingabe(r.cell(z, 2, bsp if beispiel else None), CHF)
    text(r, "C17", "← aus der Nebenkostenabrechnung übernehmen", 9, kursiv=True, farbe="555555")
    text(r, "A21", "Mängel aus dem Abnahmeprotokoll", 11, True, farbe=AKZENT)
    text(r, "A22", "Ursache «Mieter» nur bei Schäden über die normale Abnutzung hinaus.", 9, kursiv=True, farbe="555555")
    kopf(r, 23, ["Raum", "Bauteil", "Mangel", "Ursache", "Kosten laut Offerte CHF", "Einbau / letzte Erneuerung",
                 "Lebensdauer Jahre", "Alter bei Mietende", "Restwert", "Zulasten Mieter CHF", "Prüfen"])
    r.row_dimensions[23].height = 40
    dv = DataValidation(type="list", formula1='"' + ",".join(URSACHEN) + '"', allow_blank=True)
    r.add_data_validation(dv)
    for z in range(S0, SL + 1):
        for c_, fmt in ((1, None), (2, None), (3, None), (4, None), (5, CHF), (6, DATUM), (7, "0")):
            eingabe(r.cell(z, c_), fmt)
        dv.add(r.cell(z, 4))
        formel(r.cell(z, 8), f'=IF(OR(F{z}="",$B$9=""),"",($B$9-F{z})/365.25)', "0.0")
        formel(r.cell(z, 9), f'=IF(E{z}="","",IF(D{z}<>"Mieter",0,IF(G{z}="",1,IF(OR(F{z}="",$B$9=""),"",'
                             f'MAX(0,1-H{z}/G{z})))))', "0%")
        formel(r.cell(z, 10), f'=IF(OR(I{z}="",E{z}=""),0,ROUND(E{z}*I{z}*20,0)/20)', CHF)
        formel(r.cell(z, 11),
               f'=IF(E{z}="","",IF(E{z}<0,"Kosten negativ",IF(D{z}="","Ursache wählen",'
               f'IF(AND(F{z}<>"",$B$9<>"",F{z}>$B$9),"Einbau nach Mietende",'
               f'IF(AND(D{z}="Mieter",G{z}<>"",F{z}=""),"Einbaudatum fehlt",'
               f'IF(AND(D{z}="Mieter",G{z}<>"",G{z}<=0),"Lebensdauer ungültig",""))))))')
    if beispiel:
        for i, (raum, teil, mangel, urs, kosten, einbau, ld) in enumerate(BEISPIEL):
            z = S0 + i
            for c_, v in zip(range(1, 8), (raum, teil, mangel, urs, kosten, einbau, ld)):
                r.cell(z, c_).value = v
    r.cell(SL + 1, 9, "Total").font = schrift(10, True)
    formel(r.cell(SL + 1, 10), f"=SUM(J{S0}:J{SL})", CHF, True)
    rot = PatternFill("solid", fgColor="F8CBAD")
    r.conditional_formatting.add(f"K{S0}:K{SL}", FormulaRule(formula=[f'K{S0}<>""'], fill=rot))
    for col_, w in zip("ABCDEFGHIJK", (30, 16, 32, 17, 13, 13, 10, 10, 9, 13, 20)):
        r.column_dimensions[col_].width = w
    r.freeze_panes = f"A{S0}"
    r.protection.sheet = True
    r.protection.formatColumns = False

    # ------------------------------------------------------------ Abrechnung
    b = wb.create_sheet("Abrechnung")
    text(b, "A1", "Abrechnung Wohnungsrückgabe und Mietkaution", 15, True, farbe=AKZENT)
    for z, (lab, ref, fmt) in enumerate([("Liegenschaft / Wohnung", '=Rückgabe!B4&", "&Rückgabe!B5', None),
                                         ("Mieter", '=Rückgabe!B6&", "&Rückgabe!B7', None),
                                         ("Mietende", "=Rückgabe!B9", DATUM),
                                         ("Abnahme am", "=Rückgabe!B10", DATUM)], 3):
        b.cell(z, 1, lab).font = schrift(10, True)
        b.cell(z, 2, ref).font = schrift()
        b.cell(z, 2).alignment = Alignment(horizontal="left")
        if fmt:
            b.cell(z, 2).number_format = fmt
    kopf(b, 8, ["Raum / Bauteil", "Mangel", "Ursache", "Kosten CHF", "Alter / Lebensdauer", "Restwert", "Zulasten Mieter CHF"])
    B0 = 9
    for i in range(ANZ):
        z, q = B0 + i, S0 + i
        leer = f'Rückgabe!E{q}=""'
        b.cell(z, 1, f'=IF({leer},"",Rückgabe!A{q}&" · "&Rückgabe!B{q})')
        b.cell(z, 2, f'=IF({leer},"",Rückgabe!C{q})')
        b.cell(z, 3, f'=IF({leer},"",Rückgabe!D{q})')
        b.cell(z, 4, f'=IF({leer},"",Rückgabe!E{q})').number_format = CHF
        b.cell(z, 5, f'=IF(OR({leer},Rückgabe!H{q}="",Rückgabe!G{q}=""),"",'
                     f'ROUND(Rückgabe!H{q},1)&" J. / "&Rückgabe!G{q}&" J.")')
        b.cell(z, 6, f'=IF({leer},"",Rückgabe!I{q})').number_format = "0%"
        b.cell(z, 7, f'=IF({leer},"",Rückgabe!J{q})').number_format = CHF
        for c_ in range(1, 8):
            b.cell(z, c_).font = schrift(9)
    BL = B0 + ANZ - 1
    posten = [("Schäden zulasten Mieter", f"=Rückgabe!J{SL+1}"),
              ("Ausstehende Mietzinse", "=N(Rückgabe!B16)"),
              ("Nebenkosten-Saldo", "=N(Rückgabe!B17)"),
              ("Weitere Forderungen", "=N(Rückgabe!B18)"),
              ("Total Forderungen des Vermieters", f"=SUM(G{BL+2}:G{BL+5})"),
              ("Mietkaution inkl. Zins", "=N(Rückgabe!B11)+N(Rückgabe!B12)"),
              ("Aus der Kaution an den Vermieter", f"=MAX(0,MIN(G{BL+6},G{BL+7}))"),
              ("Aus der Kaution an den Mieter", f"=G{BL+7}-G{BL+8}"),
              ("Restforderung über die Kaution hinaus", f"=MAX(0,G{BL+6}-G{BL+7})"),
              ("Zusätzlich vom Vermieter zu bezahlen", f"=MAX(0,-G{BL+6})")]
    for i, (lab, f) in enumerate(posten):
        z = BL + 2 + i
        fett = lab.startswith(("Total", "Aus der Kaution"))
        b.cell(z, 4, lab).font = schrift(10, fett)
        c = b.cell(z, 7, f)
        c.number_format = CHF
        c.font = schrift(10, fett)
    ER = BL + 2   # erste Zeile der Zusammenfassung
    b.cell(ER + 11, 1, "Berechnung: Kosten × (1 − Alter ÷ Lebensdauer) gemäss paritätischer Lebensdauertabelle HEV/MV. "
                       "Normale Abnutzung geht nicht zulasten des Mieters (Art. 267 OR).").font = schrift(8, kursiv=True)
    b.cell(ER + 12, 1, "Belege und Offerten können eingesehen werden. Einwände bitte innert 30 Tagen schriftlich.").font = schrift(8, kursiv=True)
    for col_, w in zip("ABCDEFG", (28, 30, 15, 13, 15, 9, 15)):
        b.column_dimensions[col_].width = w
    b.print_area = f"A1:G{ER+12}"
    b.page_setup.fitToWidth = 1
    b.page_setup.fitToHeight = 1
    b.sheet_properties.pageSetUpPr.fitToPage = True
    b.protection.sheet = True

    # ------------------------------------------------------------ Freigabe Kaution
    f = wb.create_sheet("Freigabe Kaution")
    text(f, "A1", "Freigabe der Mietkaution", 15, True, farbe=AKZENT)
    f["A3"] = "An:"
    f["B3"] = "=Rückgabe!B13"
    f["A5"] = "Mietobjekt:"
    f["B5"] = '=Rückgabe!B4&", "&Rückgabe!B5'
    f["A6"] = "Mietende:"
    f["B6"] = "=Rückgabe!B9"
    f["B6"].number_format = DATUM
    f["B6"].alignment = Alignment(horizontal="left")
    f["A8"] = "Die Unterzeichnenden ersuchen Sie, das Mietkautionskonto zu saldieren und wie folgt auszuzahlen:"
    f["A10"] = "An den Mieter:"
    f["B10"] = "=Rückgabe!B6"
    f["C10"] = f"=Abrechnung!G{ER+7}"
    f["A11"] = "IBAN Mieter:"
    eingabe(f["B11"])
    f["A12"] = "An den Vermieter:"
    f["B12"] = "=Rückgabe!B8"
    f["C12"] = f"=Abrechnung!G{ER+6}"
    f["A13"] = "IBAN Vermieter:"
    eingabe(f["B13"])
    f["A14"] = "Allfälliger Zins bis zur Saldierung geht an den Mieter."
    for z in (10, 12):
        f.cell(z, 3).number_format = CHF
        f.cell(z, 3).font = schrift(11, True)
    f["A17"] = "Ort, Datum: ______________________"
    f["A20"] = "Mieter:"
    f["B20"] = "______________________________"
    f["A23"] = "Vermieter:"
    f["B23"] = "______________________________"
    for z in (3, 5, 6, 8, 10, 11, 12, 13, 14, 17, 20, 23):
        for c_ in ("A", "B"):
            if f[f"{c_}{z}"].value and not f[f"{c_}{z}"].font.b:
                f[f"{c_}{z}"].font = schrift(10, c_ == "A" and z in (3, 5, 6, 10, 11, 12, 13))
    f["A14"].font = schrift(9, kursiv=True)
    f.column_dimensions["A"].width = 22
    f.column_dimensions["B"].width = 44
    f.column_dimensions["C"].width = 16
    f.print_area = "A1:C24"
    f.protection.sheet = True

    # ------------------------------------------------------------ Kontrolle
    k = wb.create_sheet("Kontrolle")
    text(k, "A1", "Kontrolle", 16, True, farbe=AKZENT)
    pruef = [("Mietende erfasst (1 = ja)", '=IF(Rückgabe!B9="",0,1)', 1),
             ("Kaution nicht negativ (1 = ja)", "=IF(N(Rückgabe!B11)+N(Rückgabe!B12)>=0,1,0)", 1),
             ("Mängelzeilen mit Hinweis «Prüfen»", f'=COUNTIF(Rückgabe!K{S0}:K{SL},"?*")', 0),
             ("Aufteilung der Kaution geht auf (Differenz)",
              f"=ROUND(Abrechnung!G{ER+6}+Abrechnung!G{ER+7}-Abrechnung!G{ER+5},2)", 0),
             ("Abnahme mehr als 5 Tage nach Mietende (1 = ja, Rügefrist prüfen)",
              '=IF(OR(Rückgabe!B9="",Rückgabe!B10=""),0,IF(Rückgabe!B10-Rückgabe!B9>5,1,0))', 0)]
    for i, (lab, form, soll) in enumerate(pruef, 3):
        k.cell(i, 1, lab).font = schrift(10, True)
        formel(k.cell(i, 2), form, "0.00" if "Differenz" in lab else "0")
        k.cell(i, 3, soll).font = schrift(9, kursiv=True, farbe="555555")
    k["C2"] = "Soll"
    k["C2"].font = schrift(9, True)
    k["A9"] = "Status"
    k["A9"].font = schrift(11, True)
    formel(k["B9"], '=IF(AND(B3=C3,B4=C4,B5=C5,B6=C6,B7=C7),"OK","Bitte prüfen")', None, True)
    k.conditional_formatting.add("B9", FormulaRule(formula=['B9="OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    k.conditional_formatting.add("B9", FormulaRule(formula=['B9<>"OK"'], fill=rot))
    k.column_dimensions["A"].width = 60
    k.column_dimensions["B"].width = 16
    k.protection.sheet = True

    wb.active = 0
    name = "Wohnungsrueckgabe-CH-Beispiel.xlsx" if beispiel else "Wohnungsrueckgabe-CH.xlsx"
    wb.save(name)
    print("geschrieben:", name)
    return ER


if __name__ == "__main__":
    art = sys.argv[1] if len(sys.argv) > 1 else "beide"
    if art in ("beispiel", "beide"):
        baue(True)
    if art in ("leer", "beide"):
        baue(False)
