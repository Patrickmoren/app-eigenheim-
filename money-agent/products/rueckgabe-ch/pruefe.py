"""Rechenprobe fuer «Wohnungsrückgabe: Schaden & Kaution».

Fuellt die leere Vorlage mit Testfaellen, laesst LibreOffice rechnen und
vergleicht Mieteranteile und Kautionsaufteilung mit einer Python-Rechnung.
Aufruf: python3 pruefe.py
"""
import os
import subprocess
import sys
import tempfile
from datetime import date

import openpyxl

from build import BEISPIEL, S0, SL, baue

HIER = os.path.dirname(os.path.abspath(__file__))
ENDE = date(2026, 9, 30)

FAELLE = {
    "Beispiel": dict(ende=ENDE, kaution=4950, zins=12.40, ford=(0, 168.35, 0), maengel=BEISPIEL, ok=True),
    "älter als Lebensdauer": dict(ende=ENDE, kaution=3000, zins=0, ford=(0, 0, 0),
                                  maengel=[("K", "Herd", "defekt", "Mieter", 1500, date(2005, 1, 1), 15)], ok=True),
    "neu eingebaut": dict(ende=ENDE, kaution=3000, zins=0, ford=(0, 0, 0),
                          maengel=[("W", "Anstrich", "Kritzeleien", "Mieter", 800, ENDE, 8)], ok=True),
    "Forderungen > Kaution": dict(ende=ENDE, kaution=1000, zins=0, ford=(1800, 250, 0),
                                  maengel=[("B", "Lavabo", "Riss", "Mieter", 600, date(2020, 1, 1), 25)], ok=True),
    "Guthaben > Forderungen": dict(ende=ENDE, kaution=2000, zins=5, ford=(0, -320, 0),
                                   maengel=[("W", "Wand", "Dübellöcher", "Normale Abnutzung", 200, None, None)], ok=True),
    "ohne Kaution, Nullwerte": dict(ende=ENDE, kaution=0, zins=0, ford=(0, 0, 0),
                                    maengel=[("W", "Wand", "nichts", "Mieter", 0, date(2020, 1, 1), 8)], ok=True),
    # Fehlerfaelle: muessen «Bitte prüfen» ergeben
    "Einbaudatum fehlt": dict(ende=ENDE, kaution=1000, zins=0, ford=(0, 0, 0),
                              maengel=[("W", "Parkett", "Kratzer", "Mieter", 900, None, 40)], ok=False),
    "Einbau nach Mietende": dict(ende=ENDE, kaution=1000, zins=0, ford=(0, 0, 0),
                                 maengel=[("W", "Parkett", "Kratzer", "Mieter", 900, date(2027, 1, 1), 40)], ok=False),
    "negative Kosten": dict(ende=ENDE, kaution=1000, zins=0, ford=(0, 0, 0),
                            maengel=[("W", "Parkett", "Kratzer", "Mieter", -900, date(2010, 1, 1), 40)], ok=False),
    "Ursache fehlt": dict(ende=ENDE, kaution=1000, zins=0, ford=(0, 0, 0),
                          maengel=[("W", "Parkett", "Kratzer", None, 900, date(2010, 1, 1), 40)], ok=False),
    "Mietende fehlt": dict(ende=None, kaution=1000, zins=0, ford=(0, 0, 0),
                           maengel=[("W", "Parkett", "Kratzer", "Mieter", 900, date(2010, 1, 1), 40)], ok=False),
}


def referenz(f):
    total = 0
    for _, _, _, urs, kosten, einbau, ld in f["maengel"]:
        if urs != "Mieter" or kosten is None:
            continue
        if ld is None:
            rest = 1
        elif einbau is None or f["ende"] is None:
            continue
        else:
            rest = max(0, 1 - ((f["ende"] - einbau).days / 365.25) / ld)
        total += round(kosten * rest * 20) / 20
    ford = total + sum(f["ford"])
    k = f["kaution"] + f["zins"]
    an_v = max(0, min(ford, k))
    return round(total, 2), round(an_v, 2), round(k - an_v, 2), round(max(0, ford - k), 2), round(max(0, -ford), 2)


def main():
    os.chdir(HIER)
    ER = baue(False)
    fehler = 0
    with tempfile.TemporaryDirectory() as t:
        pfade = []
        for i, f in enumerate(FAELLE.values()):
            wb = openpyxl.load_workbook("Wohnungsrueckgabe-CH.xlsx")
            r = wb["Rückgabe"]
            r["B9"].value, r["B10"].value = f["ende"], f["ende"]
            r["B11"].value, r["B12"].value = f["kaution"], f["zins"]
            for z, v in zip((16, 17, 18), f["ford"]):
                r.cell(z, 2).value = v
            for j, m in enumerate(f["maengel"]):
                for c_, v in enumerate(m, 1):
                    r.cell(S0 + j, c_).value = v
            p = os.path.join(t, f"f{i}.xlsx")
            wb.save(p)
            pfade.append(p)
        out = os.path.join(t, "out")
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out] + pfade,
                       check=True, capture_output=True)
        for i, (name, f) in enumerate(FAELLE.items()):
            wb = openpyxl.load_workbook(os.path.join(out, f"f{i}.xlsx"), data_only=True)
            a, k, r = wb["Abrechnung"], wb["Kontrolle"], wb["Rückgabe"]
            ist = (r.cell(SL + 1, 10).value, a.cell(ER + 6, 7).value, a.cell(ER + 7, 7).value,
                   a.cell(ER + 8, 7).value, a.cell(ER + 9, 7).value)
            ist = tuple(round(x or 0, 2) for x in ist)
            status = k["B9"].value
            fehler_werte = [z for z in r.iter_rows(min_row=S0, max_row=SL) for c in z
                            if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
            if f["ok"]:
                soll = referenz(f)
                ok = ist == soll and status == "OK" and not fehler_werte
                print(f"{'OK ' if ok else 'ERR'} {name:24} Mieter {ist[0]:8.2f} | an Vermieter {ist[1]:8.2f} | "
                      f"an Mieter {ist[2]:8.2f} | Rest {ist[3]:7.2f} | Zusatz {ist[4]:6.2f} | {status}"
                      + ("" if ok else f"   SOLL {soll}"))
            else:
                ok = status != "OK" and not fehler_werte
                hinweis = r.cell(S0, 11).value
                print(f"{'OK ' if ok else 'ERR'} Fehlerfall {name:22} -> «{status}» {hinweis or ''}")
            fehler += not ok
    print("FEHLER:", fehler)
    sys.exit(1 if fehler else 0)


if __name__ == "__main__":
    main()
