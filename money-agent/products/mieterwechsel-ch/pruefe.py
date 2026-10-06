"""Rechenprobe fuer den Mieterwechsel-Planer.

Fuellt die leere Vorlage mit Grenzfaellen, laesst LibreOffice rechnen und
vergleicht Mietende, Abnahme, Maengelruege und Leerstand mit Python.
Aufruf: python3 pruefe.py
"""
import calendar
import os
import subprocess
import sys
import tempfile
from datetime import date, timedelta

import openpyxl

FEIERTAGE = {date(2026, 12, 25), date(2026, 12, 26), date(2027, 1, 1)}


def monatsende(d, plus):
    m = d.month - 1 + plus
    j, m = d.year + m // 12, m % 12 + 1
    return date(j, m, calendar.monthrange(j, m)[1])


def ende(eingang, frist, termine):
    for k in range(24):
        e = monatsende(eingang, frist + k)
        if not termine or e.month in termine:
            return e


def werktag(d, n):
    schritt = 1 if n > 0 else -1
    while n:
        d += timedelta(days=schritt)
        if d.weekday() < 5 and d not in FEIERTAGE:
            n -= schritt
    return d


FAELLE = [  # Eingang, Frist, vereinbartes Ende, neuer Mieter ab
    (date(2026, 3, 31), 3, None, None),     # letzter Tag fuer 30.6.
    (date(2026, 4, 1), 3, None, None),      # einen Tag zu spaet -> 30.9.
    (date(2026, 6, 30), 3, None, date(2026, 10, 1)),
    (date(2026, 7, 1), 3, None, None),      # -> 31.3. Folgejahr (Dezember kein Termin)
    (date(2026, 1, 15), 6, None, None),
    (date(2026, 9, 2), 3, date(2026, 10, 31), date(2026, 11, 1)),
    (date(2026, 2, 28), 1, None, date(2026, 4, 15)),
    (date(2025, 12, 31), 3, None, None),
]


def main():
    hier = os.path.dirname(os.path.abspath(__file__))
    subprocess.run([sys.executable, "build.py", "leer"], cwd=hier, check=True, capture_output=True)
    fehler = 0
    for termine in ({3, 6, 9}, set(range(1, 12)), set()):
        wb = openpyxl.load_workbook(os.path.join(hier, "Mieterwechsel-Planer-CH.xlsx"))
        e, w = wb["Einstellungen"], wb["Wechsel"]
        for m in range(12):
            e.cell(5 + m, 2).value = "x" if (m + 1) in termine else None
        for i, d in enumerate(sorted(FEIERTAGE)):
            e.cell(5 + i, 4, d)
        for i, (eing, fr, vend, neu) in enumerate(FAELLE):
            r = 5 + i
            w.cell(r, 1, f"T{i}"); w.cell(r, 4, eing); w.cell(r, 5, fr); w.cell(r, 6, vend)
            w.cell(r, 8, 1825); w.cell(r, 9, neu)
        with tempfile.TemporaryDirectory() as t:
            quelle = os.path.join(t, "in.xlsx")
            wb.save(quelle)
            os.makedirs(os.path.join(t, "out"))
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", os.path.join(t, "out"), quelle],
                           check=True, capture_output=True)
            res = openpyxl.load_workbook(os.path.join(t, "out", "in.xlsx"), data_only=True)["Wechsel"]
            for i, (eing, fr, vend, neu) in enumerate(FAELLE):
                r = 5 + i
                soll_e = vend or ende(eing, fr, termine)
                soll_a = soll_e if soll_e.weekday() < 5 and soll_e not in FEIERTAGE else werktag(soll_e, 1)
                soll_m = werktag(soll_a, 2)
                soll_l = max(0, (neu - soll_e).days - 1) if neu else None
                ist = [res.cell(r, c).value for c in (7, 16, 18, 10)]
                ist = [v.date() if hasattr(v, "date") else v for v in ist]
                soll = [soll_e, soll_a, soll_m, soll_l]
                ist[3] = None if ist[3] in ("", None) else int(ist[3])
                ok = ist == soll
                fehler += not ok
                print(f"{'OK ' if ok else 'ERR'} Termine {sorted(termine) or 'alle'} Eingang {eing} -> {ist}"
                      + ("" if ok else f"  SOLL {soll}"))
    print("FEHLER:", fehler)
    sys.exit(1 if fehler else 0)


if __name__ == "__main__":
    main()
