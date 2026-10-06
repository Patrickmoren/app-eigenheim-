# -*- coding: utf-8 -*-
"""Erzeugt die Verkaufsvorlage «Mieterwechsel- und Fristenplaner Schweiz» (Excel).

Neu geschrieben fuer private Vermieter. Verwendet keinen Code und keine
Unterlagen des Leerstandsmanagers in diesem Repository, nur allgemeines
Schweizer Mietrecht.

Fachliche Regeln
- Ende Mietverhaeltnis: erster ortsueblicher (bzw. vertraglicher) Termin, fuer den
  die Kuendigung rechtzeitig eingegangen ist. Fuer ein Monatsende E muss die
  Kuendigung spaetestens am Ende des Monats E minus Frist eintreffen.
  Kandidaten: EOMONTH(Eingang, Frist + k), k = 0..23.
- Ein vereinbartes Ende (z.B. ausserterminliche Kuendigung mit Nachmieter) geht vor.
- Abnahme: Ende oder naechster Werktag (Wochenende/Feiertage).
- Maengelruege: sofort (Art. 267a OR); Planer setzt 2 Werktage nach Abnahme.
- Kaution: Forderung innert 1 Jahr nach Mietende geltend machen (Art. 257e Abs. 3 OR).

Aufruf: python3 build.py [beispiel|leer]   (ohne Argument: beide)
"""
import sys
from datetime import date

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation

ANZ = 30
KANDIDATEN = 24
TINTE, AKZENT, HELL, EINGABE, RAND = "1A1A1A", "1F4E79", "EAF1F8", "FFF8DC", "BFC9D3"
duenn = Side(style="thin", color=RAND)
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
DATUM, CHF = "DD.MM.YYYY", "#,##0.00"
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
          "September", "Oktober", "November", "Dezember"]
STANDARD_TERMINE = {3, 6, 9}          # verbreitet; Gemeinde/Vertrag pruefen

BEISPIEL = [
    # Einheit, Mieter, durch, Eingang, Frist, vereinbartes Ende, Miete, neuer Mieter ab, erledigte Spalten
    ("EG links", "Muster Anna", "Mieter", date(2026, 5, 20), 3, None, 1650, date(2026, 11, 1), (13, 15, 17, 19)),
    ("1. OG", "Beispiel Marco", "Ausserterminlich", date(2026, 9, 2), 3, date(2026, 10, 31), 1890,
     date(2026, 11, 1), (13,)),
    ("DG", "Weber Tim", "Vermieter", date(2026, 6, 30), 3, None, 2100, None, ()),
]
FEIERTAGE = [date(2026, 12, 25), date(2026, 12, 26), date(2027, 1, 1), date(2027, 1, 2)]


def schrift(g=10, fett=False, farbe=TINTE, kursiv=False):
    return Font(name="Arial", size=g, bold=fett, color=farbe, italic=kursiv)


def eingabe(c, fmt=None):
    c.fill = PatternFill("solid", fgColor=EINGABE)
    c.protection = Protection(locked=False)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def formel(c, f, fmt=None):
    c.value = f
    c.fill = PatternFill("solid", fgColor=HELL)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def kopf(ws, zeile, texte, start=1):
    for i, t in enumerate(texte):
        c = ws.cell(zeile, start + i, t)
        c.font = schrift(9, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AKZENT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = rahmen


def baue(beispiel=True):
    wb = openpyxl.Workbook()

    a = wb.active
    a.title = "Anleitung"
    a["A1"] = "Mieterwechsel- und Fristenplaner Schweiz"
    a["A1"].font = schrift(16, True, AKZENT)
    texte = [
        ("So gehen Sie vor", True),
        ("1. «Einstellungen»: ortsübliche Kündigungstermine Ihrer Gemeinde bzw. Ihres Mietvertrags ankreuzen und Feiertage eintragen.", False),
        ("2. «Wechsel»: pro Kündigung eine Zeile. Der Planer rechnet das Ende des Mietverhältnisses, alle Fristen, den Leerstand und dessen Kosten.", False),
        ("3. Erledigte Schritte mit «x» markieren. Spalte «Nächste Frist» zeigt, was als Nächstes fällig ist (rot = überfällig, gelb = innert 7 Tagen).", False),
        ("", False),
        ("Rechtliche Hinweise (ohne Gewähr, ersetzt keine Rechtsberatung)", True),
        ("• Wohnungen: Kündigungsfrist mindestens 3 Monate auf einen ortsüblichen Termin (Art. 266c OR). Der Mietvertrag kann längere Fristen oder andere Termine vorsehen – dann diese eintragen.", False),
        ("• Massgebend ist der Eingang der Kündigung beim Empfänger, nicht der Versand.", False),
        ("• Kündigung durch den Vermieter nur mit amtlichem Formular (Art. 266l OR), bei Familienwohnungen separat an beide Ehegatten (Art. 266n OR).", False),
        ("• Ausserterminliche Kündigung: Mieter wird frei, wenn er einen zumutbaren Nachmieter stellt (Art. 264 OR). Dann «vereinbartes Ende» eintragen.", False),
        ("• Mängel bei der Rückgabe sofort rügen (Art. 267a OR). Der Planer setzt dafür 2 Werktage nach der Abnahme.", False),
        ("• Mietkaution: Ohne Forderung innert eines Jahres nach Mietende kann der Mieter die Kaution bei der Bank herausverlangen (Art. 257e Abs. 3 OR).", False),
        ("• Fällt das Mietende auf ein Wochenende oder einen Feiertag, findet die Rückgabe am nächsten Werktag statt.", False),
    ]
    for i, (t, fett) in enumerate(texte, 3):
        a.cell(i, 1, t).font = schrift(11 if fett else 10, fett)
    a.column_dimensions["A"].width = 140

    e = wb.create_sheet("Einstellungen")
    e["A1"] = "Einstellungen"
    e["A1"].font = schrift(16, True, AKZENT)
    e["A3"] = "Kündigungstermine (Monatsende)"
    e["A3"].font = schrift(11, True, AKZENT)
    kopf(e, 4, ["Monat", "Termin? (x)"])
    for m in range(12):
        e.cell(5 + m, 1, MONATE[m]).font = schrift()
        eingabe(e.cell(5 + m, 2, "x" if (m + 1) in STANDARD_TERMINE else None))
    e["A18"] = ("Voreinstellung: Ende März, Juni, September (verbreitet). Viele Städte (z.B. Zürich) erlauben jedes "
                "Monatsende ausser Dezember. Massgebend sind Gemeinde und Mietvertrag.")
    e["A18"].font = schrift(9, kursiv=True)
    e["D3"] = "Feiertage (für Werktagsrechnung)"
    e["D3"].font = schrift(11, True, AKZENT)
    kopf(e, 4, ["Datum"], start=4)
    for i in range(30):
        eingabe(e.cell(5 + i, 4), DATUM)
    if beispiel:
        for i, d in enumerate(FEIERTAGE):
            e.cell(5 + i, 4, d)
    e.column_dimensions["A"].width = 16
    e.column_dimensions["B"].width = 14
    e.column_dimensions["D"].width = 14
    e.protection.sheet = True
    FEIER = "Einstellungen!$D$5:$D$34"
    FLAGS = "Einstellungen!$B$5:$B$16"

    w = wb.create_sheet("Wechsel")
    w["A1"] = "Mieterwechsel"
    w["A1"].font = schrift(16, True, AKZENT)
    w["A2"] = "Heute:"
    formel(w["B2"], "=TODAY()", DATUM)
    spalten = ["Einheit", "Mieter bisher", "Kündigung durch", "Eingang Kündigung", "Frist (Monate)",
               "Vereinbartes Ende (optional)", "Ende Mietverhältnis", "Nettomiete / Monat CHF",
               "Neuer Mieter ab", "Leerstand Tage", "Leerstandskosten CHF",
               "Nachmieter suchen ab", "erl.", "Vorabnahme", "erl.", "Abnahme / Übergabe", "erl.",
               "Mängelrüge bis", "erl.", "Kaution: Forderung bis", "erl.",
               "Nächste Frist", "Tage bis", "Hinweis"]
    kopf(w, 4, spalten)
    w.row_dimensions[4].height = 42
    # Hilfsspalten ab AA: Kandidaten fuer das Ende, danach offene Fristen
    H0 = 27
    O0 = H0 + KANDIDATEN
    w.cell(3, H0, "Hilfsspalten (Kandidaten Mietende, offene Fristen)").font = schrift(8, kursiv=True)
    dv_d = DataValidation(type="list", formula1='"Mieter,Vermieter,Ausserterminlich"', allow_blank=True)
    dv_x = DataValidation(type="list", formula1='"x"', allow_blank=True)
    w.add_data_validation(dv_d); w.add_data_validation(dv_x)
    Z0 = 5
    for i in range(ANZ):
        r = Z0 + i
        for c_, fmt in ((1, None), (2, None), (3, None), (4, DATUM), (5, "0"), (6, DATUM), (8, CHF), (9, DATUM)):
            eingabe(w.cell(r, c_), fmt)
        dv_d.add(w.cell(r, 3))
        aktiv = f'D{r}<>""'
        frist = f'IF(E{r}="",3,E{r})'
        # Kandidaten: erstes zulaessiges Monatsende; ohne angekreuzte Termine jedes Monatsende
        for k in range(KANDIDATEN):
            kc = w.cell(r, H0 + k)
            kc.value = (f'=IF({aktiv},IF(OR(COUNTIF({FLAGS},"x")=0,'
                        f'INDEX({FLAGS},MONTH(EOMONTH(D{r},{frist}+{k})))="x"),EOMONTH(D{r},{frist}+{k}),""),"")')
            kc.number_format = DATUM
            kc.font = schrift(8, farbe="888888")
        kand = f"{col(H0)}{r}:{col(H0 + KANDIDATEN - 1)}{r}"
        formel(w.cell(r, 7), f'=IF(F{r}<>"",F{r},IF({aktiv},IF(COUNT({kand})=0,"",MIN({kand})),""))', DATUM)
        formel(w.cell(r, 10), f'=IF(OR(G{r}="",I{r}=""),"",MAX(0,I{r}-G{r}-1))', "0")
        formel(w.cell(r, 11), f'=IF(OR(J{r}="",H{r}=""),"",ROUND(H{r}*12/365*J{r},2))', CHF)
        fristen = {12: f'=IF({aktiv},D{r},"")',
                   14: f'=IF(G{r}="","",WORKDAY(G{r},-10,{FEIER}))',
                   16: f'=IF(G{r}="","",WORKDAY(G{r}-1,1,{FEIER}))',
                   18: f'=IF(P{r}="","",WORKDAY(P{r},2,{FEIER}))',
                   20: f'=IF(G{r}="","",EDATE(G{r},12))'}
        for j, (c_, f) in enumerate(fristen.items()):
            formel(w.cell(r, c_), f, DATUM)
            eingabe(w.cell(r, c_ + 1))
            dv_x.add(w.cell(r, c_ + 1))
            oc = w.cell(r, O0 + j)
            oc.value = f'=IF(AND({col(c_)}{r}<>"",{col(c_ + 1)}{r}=""),{col(c_)}{r},"")'
            oc.number_format = DATUM
            oc.font = schrift(8, farbe="888888")
        offen = f"{col(O0)}{r}:{col(O0 + 4)}{r}"
        formel(w.cell(r, 22), f'=IF(COUNT({offen})=0,"",MIN({offen}))', DATUM)
        formel(w.cell(r, 23), f'=IF(V{r}="","",V{r}-$B$2)', "0")
        formel(w.cell(r, 24),
               f'=IF(C{r}="Vermieter","Amtliches Formular verwendet? Bei Familienwohnung an beide Ehegatten.",'
               f'IF(C{r}="Ausserterminlich","Nachmieter zumutbar? Vereinbartes Ende eintragen.",'
               f'IF(AND(G{r}<>"",I{r}=""),"Noch kein Nachmieter.","")))')
    ZL = Z0 + ANZ - 1
    if beispiel:
        for i, (einh, name, durch, eing, fr, vend, miete, neu, erl) in enumerate(BEISPIEL):
            r = Z0 + i
            for c_, v in zip((1, 2, 3, 4, 5, 6, 8, 9), (einh, name, durch, eing, fr, vend, miete, neu)):
                w.cell(r, c_, v)
            for c_ in erl:
                w.cell(r, c_, "x")
    w.cell(ZL + 1, 10, "Total").font = schrift(10, True)
    formel(w.cell(ZL + 1, 11), f"=SUM(K{Z0}:K{ZL})", CHF)
    rot, gelb = PatternFill("solid", fgColor="F8CBAD"), PatternFill("solid", fgColor="FFE699")
    w.conditional_formatting.add(f"V{Z0}:W{ZL}", FormulaRule(formula=[f'AND($W{Z0}<>"",$W{Z0}<0)'], fill=rot))
    w.conditional_formatting.add(f"V{Z0}:W{ZL}", FormulaRule(formula=[f'AND($W{Z0}<>"",$W{Z0}>=0,$W{Z0}<=7)'],
                                                             fill=gelb))
    breiten = [12, 18, 14, 12, 8, 13, 13, 12, 12, 9, 12, 12, 4, 12, 4, 12, 4, 12, 4, 13, 4, 12, 7, 46]
    for i, b in enumerate(breiten, 1):
        w.column_dimensions[col(i)].width = b
    for k in range(H0, O0 + 5):
        w.column_dimensions[col(k)].hidden = True
    w.freeze_panes = "C5"
    w.protection.sheet = True
    w.protection.formatColumns = False

    wb.active = 0
    name = "Mieterwechsel-Planer-CH-Beispiel.xlsx" if beispiel else "Mieterwechsel-Planer-CH.xlsx"
    wb.save(name)
    print("geschrieben:", name)


if __name__ == "__main__":
    art = sys.argv[1] if len(sys.argv) > 1 else "beide"
    if art in ("beispiel", "beide"):
        baue(True)
    if art in ("leer", "beide"):
        baue(False)
