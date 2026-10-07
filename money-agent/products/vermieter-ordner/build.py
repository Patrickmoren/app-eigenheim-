# -*- coding: utf-8 -*-
"""Erzeugt den «Vermieter-Ordner Schweiz»: eine Arbeitsmappe, sechs Teile.

  Start           Übersicht, Workflow, nächste Fristen, Kennzahlen
  Mein Haus       Einheiten und Mietverhältnisse – einmal erfassen, überall verwendet
  Rückgabe        Abnahmeprotokoll → Schaden nach Lebensdauer → Kautionsabrechnung
  Nebenkosten     Verteilung nach Fläche/Anteil/Einheiten, Heizgradtage, Leerstand
  Mietzins        Referenzzinssatz, Teuerung (40 %), Kostensteigerung
  Mieterwechsel   Mietende, zehn Schritte mit Datum (die frühere Checkliste), Ampel

Kompatibilität: keine Makros, kein TEXT(), kein FILTER(), keine Matrixformeln.
Aufruf: python3 build.py [beispiel|leer]   (ohne Argument: beide)
"""
import sys
from datetime import date

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter as col
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# ------------------------------------------------------------------ Layout
EIN0, EIN_N = 11, 30                 # Einheiten
EINL = EIN0 + EIN_N - 1              # 40
ETOT = EINL + 1                      # 41
MV0, MV_N = 45, 60                   # Mietverhältnisse, ID = Zeile - 44
MVL = MV0 + MV_N - 1                 # 104
K0, K_N = 10, 25                     # Kostenarten
KL = K0 + K_N - 1                    # 34
NV0 = 5                              # NK Verteilung: Zeile = 4 + ID
NVL = NV0 + MV_N - 1                 # 64
W0, W_N = 9, 20                      # Mieterwechsel
WL = W0 + W_N - 1                    # 28
F0, F_N = 8, 20                      # Rückgabe-Fälle
FL = F0 + F_N - 1                    # 27
M0, M_N = 32, 100                    # Mängel
ML = M0 + M_N - 1                    # 131
DRUCK = 30                           # Zeilen in Protokoll/Abrechnung

MH = "'Mein Haus'!"
SALDO = "BB"                         # Saldo-Spalte in «NK Verteilung» (27 + 25 + 2)
NV = "'NK Verteilung'!"
EIN = "Einstellungen!"
FLAGS = f"{EIN}$B$5:$B$16"
FEIER = f"{EIN}$D$5:$D$34"
GEW = f"{EIN}$F$5:$F$16"


def mh(spalte):
    """Spalte der Mietverhältnisse in «Mein Haus» als absoluter Bereich."""
    return f"{MH}${spalte}${MV0}:${spalte}${MVL}"


EINH_NR = f"{MH}$A${EIN0}:$A${EINL}"

# Mieterwechsel: Spalten
STEP0 = 12                           # erste Schritt-Spalte (L)
SCHRITTE = [
    ("Kündigung schriftlich bestätigen", "Eingang bestätigen; bei Mieterkündigung Form und Termin prüfen."),
    ("Nachmieter suchen / Inserat", "Inserat schalten, Besichtigungen planen. Bei ausserterminlicher Kündigung: Nachmieter prüfen (Art. 264 OR)."),
    ("Vorabnahme", "Wohnung mit dem Mieter begehen, notwendige Arbeiten absprechen."),
    ("Neuen Mietvertrag abschliessen", "Sollfrist 30 Tage vor Mietende, damit kein Leerstand entsteht."),
    ("Abnahme · Protokoll · Zähler · Schlüssel", "Blatt «Rückgabe»: Zählerstände, Schlüssel und Mängel erfassen, «Protokoll» drucken und unterschreiben."),
    ("Mängelrüge schriftlich", "Sofort nach der Abnahme rügen (Art. 267a OR), sonst verwirkt."),
    ("Übergabe an Neumieter", "Mit Protokoll übergeben; Neumieter kann das Rückgabeprotokoll des Vormieters verlangen (Art. 256a OR)."),
    ("Kaution abrechnen", "Blatt «Kautionsabrechnung» und «Freigabe Kaution»."),
    ("Nebenkosten-Schlussabrechnung", "Nach Ende der Abrechnungsperiode, Blatt «Nebenkosten»."),
    ("Kautionsfrist (1 Jahr nach Mietende)", "Ohne geltend gemachten Anspruch kann der Mieter die Kaution allein herausverlangen (Art. 257e OR)."),
]
NEXT_C = STEP0 + 2 * len(SCHRITTE)   # 32 AF
TAGE_C, SCHRITT_C, HINW_C = NEXT_C + 1, NEXT_C + 2, NEXT_C + 3
KAND0 = HINW_C + 2                   # 37 Kandidaten Mietende (24)
OFFEN0 = KAND0 + 24                  # 61 offene Fristen (10)
SCHLUESSEL_C = OFFEN0 + len(SCHRITTE)  # 71 eindeutiger Sortierschlüssel

TINTE, AKZENT, HELL, EINGABE, RAND = "1A1A1A", "1F4E79", "EAF1F8", "FFF8DC", "BFC9D3"
duenn = Side(style="thin", color=RAND)
rahmen = Border(left=duenn, right=duenn, top=duenn, bottom=duenn)
DATUM, CHF, PROZ = "DD.MM.YYYY", "#,##0.00", "0.00%"
ROT = PatternFill("solid", fgColor="F8CBAD")
GELB = PatternFill("solid", fgColor="FFE699")
GRUEN = PatternFill("solid", fgColor="C6EFCE")
MONATE = ["Januar", "Februar", "März", "April", "Mai", "Juni", "Juli", "August",
          "September", "Oktober", "November", "Dezember"]
GRADTAGE = [170, 150, 130, 90, 40, 10, 0, 0, 30, 80, 140, 160]
FEIERTAGE = [date(2025, 12, 25), date(2025, 12, 26), date(2026, 1, 1), date(2026, 1, 2),
             date(2026, 12, 25), date(2026, 12, 28), date(2027, 1, 1)]

# ------------------------------------------------------------------ Beispiel
B_EINHEITEN = [("1", "EG links", 3.5, 82, 250), ("2", "EG rechts", 2.5, 61, 190),
               ("3", "1. OG", 4.5, 104, 320), ("4", "DG", 3.5, 78, 240)]
# Einheit, Mieter, Beginn, Ende manuell, Netto, Akonto, Kaution, Bank, Ref, LIK, Basisdatum
B_MIETER = [
    ("1", "Muster Anna", date(2019, 4, 1), None, 1650, 380, 4950, "Musterbank, Konto 11-1", 1.75, 106.2, date(2023, 12, 1)),
    ("2", "Beispiel Marco", date(2017, 10, 1), None, 1290, 400, 3870, "Musterbank, Konto 11-2", 1.50, 104.9, date(2022, 4, 1)),
    ("2", "Probst Lea", date(2025, 4, 1), None, 1350, 260, 4050, "Musterbank, Konto 11-3", 1.75, 107.1, date(2025, 4, 1)),
    ("3", "Keller Familie", date(2012, 5, 1), None, 2150, 480, 6000, "Musterbank, Konto 11-4", 1.75, 106.2, date(2023, 12, 1)),
    ("4", "Weber Tim", date(2016, 3, 1), None, 1720, 350, 4950, "Musterbank, Konto 11-5", 1.75, 106.2, date(2023, 12, 1)),
    ("4", "Frei Sara", date(2025, 12, 1), None, 1750, 200, 5250, "Musterbank, Konto 11-6", 1.25, 107.3, date(2025, 12, 1)),
]
B_KOSTEN = [
    ("Heizung und Warmwasser (Gas)", 9600, "Fläche", "Heizung"),
    ("Hauswartung", 3600, "Fläche", "Linear"),
    ("Treppenhausreinigung", 1200, "Einheiten", "Linear"),
    ("Allgemeinstrom", 640, "Anteil", "Linear"),
    ("Wasser und Abwasser", 2350, "Anteil", "Linear"),
    ("Kehrichtgrundgebühr", 480, "Einheiten", "Linear"),
    ("Service Heizanlage", 520, "Fläche", "Heizung"),
    ("Verwaltungsaufwand 3 %", "=ROUND(SUM(B10:B16)*3%,2)", "Fläche", "Linear"),
]
# ID, durch, Eingang, Frist, vereinbartes Ende, Neuer Mieter ab, erledigt (alle)
B_WECHSEL = [
    (2, "Mieter", date(2024, 12, 10), 3, None, date(2025, 4, 1), True),
    (5, "Vermieter", date(2025, 6, 30), 3, None, date(2025, 12, 1), True),
    (1, "Mieter", date(2026, 8, 20), 3, None, None, False),
]
B_FALL = dict(id=5, abnahme=date(2025, 9, 30), anwesend="T. Weber, H. Beispiel", strom=12873, wasser=None,
              waerme=None, schluessel="4 von 5", zins=9.80, miete=0, weitere=0)
# ID, Raum, Bauteil, Mangel, Ursache, Kosten, Einbau, Lebensdauer
B_MAENGEL = [
    (5, "Wohnzimmer", "Wandanstrich", "Starke Nikotinverfärbung", "Mieter", 1800, date(2019, 4, 1), 8),
    (5, "Küche", "Backofen", "Türglas gesprungen", "Mieter", 950, date(2014, 6, 1), 15),
    (5, "Wohnzimmer", "Parkett", "Wasserflecken vor Balkontür", "Mieter", 2400, date(2005, 1, 1), 40),
    (5, "Schlafzimmer", "Teppich", "Laufspuren", "Normale Abnutzung", 1200, date(2016, 3, 1), 10),
    (5, "Bad", "Spiegelschrank", "Kratzer (im Eintrittsprotokoll vermerkt)", "Vorbestehend", None, None, None),
    (5, "Eingang", "Schlüssel", "1 Wohnungsschlüssel fehlt", "Mieter", 180, None, None),
]
B_MIETZINS = dict(id=4, ref=1.25, lik=107.1, stichtag=date(2026, 10, 1), kosten=0.5, mitteilung=date(2026, 10, 15), frist=3)


# ------------------------------------------------------------------ Helfer
def schrift(g=10, fett=False, farbe=TINTE, kursiv=False):
    return Font(name="Arial", size=g, bold=fett, color=farbe, italic=kursiv)


def eingabe(c, fmt=None, wert=None):
    if wert is not None:
        c.value = wert
    c.fill = PatternFill("solid", fgColor=EINGABE)
    c.protection = Protection(locked=False)
    c.border = rahmen
    c.font = schrift()
    if fmt:
        c.number_format = fmt


def formel(c, f, fmt=None, fett=False):
    c.value = f
    c.fill = PatternFill("solid", fgColor=HELL)
    c.border = rahmen
    c.font = schrift(10, fett)
    if fmt:
        c.number_format = fmt


def kopf(ws, zeile, texte, start=1, hoehe=None):
    for i, t in enumerate(texte):
        c = ws.cell(zeile, start + i, t)
        c.font = schrift(9, True, "FFFFFF")
        c.fill = PatternFill("solid", fgColor=AKZENT)
        c.alignment = Alignment(wrap_text=True, vertical="center")
        c.border = rahmen
    if hoehe:
        ws.row_dimensions[zeile].height = hoehe


def titel(ws, t, unter=None):
    ws["A1"] = t
    ws["A1"].font = schrift(16, True, AKZENT)
    if unter:
        ws["A2"] = unter
        ws["A2"].font = schrift(9, kursiv=True, farbe="555555")


def label(ws, zelle, t, fett=True, g=10, kursiv=False, farbe=TINTE):
    ws[zelle] = t
    ws[zelle].font = schrift(g, fett, farbe, kursiv)


def leitfaden(ws, zeile, start, berechnung, ergebnis, ausgabe, spalte=1):
    """Vier Zeilen START / BERECHNUNG / ERGEBNIS / AUSGABE oben im Modul."""
    for i, (k, t) in enumerate((("START", start), ("BERECHNUNG", berechnung),
                                ("ERGEBNIS", ergebnis), ("AUSGABE", ausgabe))):
        c = ws.cell(zeile + i, spalte, f"{k}: {t}")
        c.font = schrift(9, False, "333333")


def breiten(ws, werte, start=1):
    for i, w in enumerate(werte):
        ws.column_dimensions[col(start + i)].width = w


def schuetzen(ws):
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False


def wert_oder_leer(ausdruck):
    """INDEX auf eine leere Zelle liefert 0 – hier stattdessen ""."""
    return f'IF({ausdruck}="","",{ausdruck})'


# ------------------------------------------------------------------ Blätter
def blatt_einstellungen(wb, beispiel):
    e = wb.create_sheet("Einstellungen")
    titel(e, "Einstellungen", "Gilt für alle Teile. Einmal einstellen.")
    label(e, "A3", "Kündigungstermine (Monatsende)", g=11, farbe=AKZENT)
    kopf(e, 4, ["Monat", "Termin? (x)"])
    for m in range(12):
        e.cell(5 + m, 1, MONATE[m]).font = schrift()
        eingabe(e.cell(5 + m, 2), wert="x" if (m + 1) in (3, 6, 9) else None)
    label(e, "A18", "Voreinstellung Ende März, Juni, September. Massgebend sind Gemeinde und Mietvertrag.",
          False, 9, True, "555555")
    label(e, "D3", "Feiertage", g=11, farbe=AKZENT)
    kopf(e, 4, ["Datum"], start=4)
    feiertage = FEIERTAGE
    for i in range(30):
        eingabe(e.cell(5 + i, 4), DATUM, feiertage[i] if beispiel and i < len(feiertage) else None)
    label(e, "F3", "Heizung: Monatsgewichte ‰", g=11, farbe=AKZENT)
    kopf(e, 4, ["Gewicht"], start=6)
    for m in range(12):
        eingabe(e.cell(5 + m, 6), "0", GRADTAGE[m])
        e.cell(5 + m, 5, MONATE[m][:3]).font = schrift(9)
    label(e, "E17", "Σ", g=10)
    formel(e["F17"], "=SUM(F5:F16)", "0")
    breiten(e, [14, 12, 3, 14, 6, 10])
    schuetzen(e)


def blatt_mein_haus(wb, beispiel):
    h = wb.create_sheet("Mein Haus")
    titel(h, "Mein Haus", "Alles, was mehrere Teile brauchen, steht nur hier. Gelb = Eingabe, Blau = rechnet.")
    for z, lab, wert in ((4, "Liegenschaft", "Musterstrasse 12, 8004 Zürich"),
                         (5, "Vermieter", "Hans Beispiel, Bergstrasse 1, 8004 Zürich"),
                         (6, "IBAN Vermieter", None)):
        label(h, f"A{z}", lab)
        eingabe(h.cell(z, 2), wert=wert if beispiel else None)
    label(h, "A9", "Einheiten", g=12, farbe=AKZENT)
    kopf(h, 10, ["Nr.", "Bezeichnung", "Zimmer", "Fläche m²", "Anteil (Wertquote ‰ o.ä.)", "Zähler"])
    for i in range(EIN_N):
        z = EIN0 + i
        for c_, fmt in ((1, None), (2, None), (3, "0.0"), (4, "0.00"), (5, "0.00")):
            eingabe(h.cell(z, c_), fmt)
        formel(h.cell(z, 6), f'=IF(A{z}="","",1)', "0")
        if beispiel and i < len(B_EINHEITEN):
            for c_, v in enumerate(B_EINHEITEN[i], 1):
                h.cell(z, c_).value = v
    label(h, f"B{ETOT}", "Total")
    for c_ in (4, 5, 6):
        formel(h.cell(ETOT, c_), f"=SUM({col(c_)}{EIN0}:{col(c_)}{EINL})", "0.00")

    label(h, "A43", "Mietverhältnisse", g=12, farbe=AKZENT)
    kopf(h, 44, ["ID", "Einheit Nr.", "Mieter", "Mietbeginn", "Mietende (manuell)", "Mietende wirksam",
                 "Nettomiete / Mt.", "Akonto NK / Mt.", "Kaution", "Kautionskonto (Bank)",
                 "Referenzzins Basis %", "LIK Basis (Punkte)", "Basis seit (letzte Anpassung)",
                 "Status", "Prüfen"], hoehe=42)
    dv = DataValidation(type="list", formula1=f"${'A'}${EIN0}:$A${EINL}", allow_blank=True)
    h.add_data_validation(dv)
    for i in range(MV_N):
        z = MV0 + i
        formel(h.cell(z, 1), i + 1, "0")
        for c_, fmt in ((2, None), (3, None), (4, DATUM), (5, DATUM), (7, CHF), (8, CHF), (9, CHF),
                        (10, None), (11, "0.00"), (12, "0.0"), (13, DATUM)):
            eingabe(h.cell(z, c_), fmt)
        dv.add(h.cell(z, 2))
        formel(h.cell(z, 6), f'=IF(E{z}<>"",E{z},IFERROR(INDEX(Mieterwechsel!$H${W0}:$H${WL},'
                             f'MATCH(A{z},Mieterwechsel!$A${W0}:$A${WL},0)),""))', DATUM)
        formel(h.cell(z, 14), f'=IF(C{z}="","",IF(AND(F{z}<>"",F{z}<TODAY()),"beendet",'
                              f'IF(AND(D{z}<>"",D{z}>TODAY()),"künftig","aktiv")))')
        formel(h.cell(z, 15), f'=IF(C{z}="","",IF(B{z}="","Einheit fehlt",IF(ISNA(MATCH(B{z},$A${EIN0}:$A${EINL},0)),'
                              f'"Einheit unbekannt",IF(AND(D{z}<>"",F{z}<>"",F{z}<D{z}),"Ende vor Beginn",'
                              f'IF(R{z}>0,"Überschneidung mit anderem Mieter","")))))')
        # Hilfsspalten für die Überschneidungsprüfung
        h.cell(z, 16).value = f'=IF(C{z}="","",IF(D{z}="",1,D{z}))'
        h.cell(z, 17).value = f'=IF(C{z}="","",IF(F{z}="",401768,F{z}))'
        h.cell(z, 18).value = (f'=IF(C{z}="",0,COUNTIFS($B${MV0}:$B${MVL},B{z},$P${MV0}:$P${MVL},"<="&Q{z},'
                               f'$Q${MV0}:$Q${MVL},">="&P{z})-1)')
        if beispiel and i < len(B_MIETER):
            for c_, v in zip((2, 3, 4, 5, 7, 8, 9, 10, 11, 12, 13), B_MIETER[i]):
                h.cell(z, c_).value = v
    h.conditional_formatting.add(f"O{MV0}:O{MVL}", FormulaRule(formula=[f'O{MV0}<>""'], fill=ROT))
    for c_ in (16, 17, 18):
        h.column_dimensions[col(c_)].hidden = True
    breiten(h, [16, 34, 26, 12, 13, 13, 12, 11, 11, 24, 10, 10, 12, 10, 26])
    h.freeze_panes = "D45"
    schuetzen(h)


def blatt_nebenkosten(wb, beispiel):
    n = wb.create_sheet("Nebenkosten")
    titel(n, "Nebenkosten")
    leitfaden(n, 2, "Periodenbeginn (1. eines Monats) und die Kosten der Periode. Mieter und Flächen kommen aus «Mein Haus».",
              "Verteilt jede Kostenart nach Schlüssel; Heizung nach Monatsgewichten, sonst taggenau. Leerstand bleibt beim Eigentümer.",
              "Blatt «NK Abrechnung» zeigt pro Mieter Nachzahlung oder Guthaben.",
              "«NK Abrechnung» pro Mieter drucken (eine A4-Seite).")
    label(n, "A7", "Beginn Periode")
    eingabe(n["B7"], DATUM, date(2025, 1, 1) if beispiel else None)
    dv_t = DataValidation(type="custom", formula1="DAY(B7)=1", showErrorMessage=True, error="Immer der 1. eines Monats.")
    n.add_data_validation(dv_t)
    dv_t.add("B7")
    label(n, "C7", "Ende")
    formel(n["D7"], '=IF(B7="","",EDATE(B7,12)-1)', DATUM)
    label(n, "E7", "Tage")
    formel(n["F7"], '=IF(B7="","",D7-B7+1)', "0")
    kopf(n, K0 - 1, ["Kostenart", "Betrag CHF", "Schlüssel", "Verteilung", "Total Schlüssel"])
    dv_s = DataValidation(type="list", formula1='"Fläche,Anteil,Einheiten"', allow_blank=True)
    dv_v = DataValidation(type="list", formula1='"Linear,Heizung"', allow_blank=True)
    n.add_data_validation(dv_s); n.add_data_validation(dv_v)
    for i in range(K_N):
        z = K0 + i
        eingabe(n.cell(z, 1)); eingabe(n.cell(z, 2), CHF); eingabe(n.cell(z, 3)); eingabe(n.cell(z, 4))
        dv_s.add(n.cell(z, 3)); dv_v.add(n.cell(z, 4))
        formel(n.cell(z, 5), f'=IF(C{z}="Fläche",{MH}$D${ETOT},IF(C{z}="Anteil",{MH}$E${ETOT},'
                             f'IF(C{z}="Einheiten",{MH}$F${ETOT},0)))', "0.00")
        if beispiel and i < len(B_KOSTEN):
            for c_, v in enumerate(B_KOSTEN[i], 1):
                n.cell(z, c_).value = v
    label(n, f"A{KL+1}", "Total")
    formel(n.cell(KL + 1, 2), f"=SUM(B{K0}:B{KL})", CHF, True)
    breiten(n, [38, 15, 13, 12, 14, 8])
    schuetzen(n)

    # ---------------------------------------------------- NK Verteilung
    v = wb.create_sheet("NK Verteilung")
    titel(v, "NK Verteilung", "Eine Zeile pro Mietverhältnis aus «Mein Haus». Eingabe nur in «Akonto bezahlt», falls abweichend.")
    kopf(v, 4, ["ID", "Mieter", "Einheit", "Beginn in Periode", "Ende in Periode", "Tage", "Faktor linear",
                "Faktor Heizung", "Fläche", "Anteil", "Einheiten", "Akonto Soll", "Akonto bezahlt (falls abweichend)",
                "Akonto verwendet"] + [f"Belegung M{j+1}" for j in range(12)]
         + [f"K{j+1}" for j in range(K_N)] + ["Total (5 Rp.)", "Akonto", "Saldo"], hoehe=42)
    O0 = 15                     # Belegung M1
    D0 = O0 + 12                # erste Kostenspalte
    TOT = D0 + K_N
    PS, PE, PT = "Nebenkosten!$B$7", "Nebenkosten!$D$7", "Nebenkosten!$F$7"
    v.cell(3, O0 - 1, "Gewicht ‰ →").font = schrift(8, kursiv=True)
    for j in range(12):
        formel(v.cell(3, O0 + j), f'=IF({PS}="",0,INDEX({GEW},MONTH(EDATE({PS},{j}))))', "0")
    for j in range(K_N):
        formel(v.cell(3, D0 + j), f'=IF(Nebenkosten!A{K0+j}="","",Nebenkosten!A{K0+j})')
        v.cell(3, D0 + j).alignment = Alignment(wrap_text=True, vertical="bottom")
    v.row_dimensions[3].height = 48
    for i in range(MV_N):
        z, m = NV0 + i, MV0 + i
        akt = f'AND({MH}C{m}<>"",{MH}B{m}<>"",{PS}<>"")'
        formel(v.cell(z, 1), f"={MH}A{m}", "0")
        formel(v.cell(z, 2), f'=IF({MH}C{m}="","",{MH}C{m})')
        formel(v.cell(z, 3), f'=IF({MH}B{m}="","",{MH}B{m})')
        formel(v.cell(z, 4), f'=IF({akt},MAX(IF({MH}D{m}="",{PS},{MH}D{m}),{PS}),"")', DATUM)
        formel(v.cell(z, 5), f'=IF({akt},MIN(IF({MH}F{m}="",{PE},{MH}F{m}),{PE}),"")', DATUM)
        formel(v.cell(z, 6), f'=IF(D{z}="",0,MAX(0,E{z}-D{z}+1))', "0")
        formel(v.cell(z, 7), f"=IF(F{z}=0,0,F{z}/{PT})", "0.0000")
        formel(v.cell(z, 8), f"=IF(F{z}=0,0,SUMPRODUCT({col(O0)}{z}:{col(O0+11)}{z},${col(O0)}$3:${col(O0+11)}$3)/{EIN}$F$17)", "0.0000")
        for c_, quelle in ((9, "D"), (10, "E"), (11, "F")):
            formel(v.cell(z, c_), f'=IF(F{z}=0,0,IFERROR(INDEX({MH}${quelle}${EIN0}:${quelle}${EINL},'
                                  f'MATCH(C{z},{EINH_NR},0)),0))', "0.00")
        formel(v.cell(z, 12), f"=IF(F{z}=0,0,N({MH}H{m})*SUM({col(O0)}{z}:{col(O0+11)}{z}))", CHF)
        eingabe(v.cell(z, 13), CHF)
        formel(v.cell(z, 14), f'=IF(M{z}<>"",M{z},L{z})', CHF)
        for j in range(12):
            ms = f"EDATE({PS},{j})"
            formel(v.cell(z, O0 + j), f'=IF(F{z}=0,0,MAX(0,MIN(E{z},EOMONTH({ms},0))-MAX(D{z},{ms})+1)/DAY(EOMONTH({ms},0)))', "0.00")
        for j in range(K_N):
            kr = K0 + j
            schl = f'IF(Nebenkosten!$C${kr}="Fläche",$I{z},IF(Nebenkosten!$C${kr}="Anteil",$J{z},$K{z}))'
            fak = f'IF(Nebenkosten!$D${kr}="Heizung",$H{z},$G{z})'
            formel(v.cell(z, D0 + j), f'=IF(OR(Nebenkosten!$B${kr}="",Nebenkosten!$E${kr}=0,$F{z}=0),0,'
                                      f'Nebenkosten!$B${kr}*{schl}*{fak}/Nebenkosten!$E${kr})', CHF)
        formel(v.cell(z, TOT), f"=ROUND(SUM({col(D0)}{z}:{col(TOT-1)}{z})*20,0)/20", CHF)
        formel(v.cell(z, TOT + 1), f"=IF(F{z}=0,0,N{z})", CHF)
        formel(v.cell(z, TOT + 2), f"={col(TOT)}{z}-{col(TOT+1)}{z}", CHF)
    label(v, f"B{NVL+1}", "Summe Mieter")
    label(v, f"B{NVL+2}", "Leerstand / Eigentümer")
    for j in range(K_N + 1):
        L = col(D0 + j)
        formel(v.cell(NVL + 1, D0 + j), f"=SUM({L}{NV0}:{L}{NVL})", CHF)
    for j in range(K_N):
        L = col(D0 + j)
        formel(v.cell(NVL + 2, D0 + j), f"=N(Nebenkosten!B{K0+j})-{L}{NVL+1}", CHF)
    breiten(v, [5, 22, 8, 11, 11, 6, 8, 8, 8, 8, 8, 11, 14, 11])
    v.freeze_panes = "C5"
    schuetzen(v)
    return D0, TOT

def blatt_nk_abrechnung(wb, D0, TOT):
    b = wb.create_sheet("NK Abrechnung")
    label(b, "A1", "Heiz- und Nebenkostenabrechnung", g=16, farbe=AKZENT)
    label(b, "J1", "Mietverhältnis-ID (wird nicht gedruckt):", False, 9, True)
    eingabe(b["J2"], "0", 1)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MV_N))
    b.add_data_validation(dv); dv.add("J2")
    ID = "$J$2"

    def nv(spalte):
        return f"INDEX({NV}${spalte}${NV0}:${spalte}${NVL},{ID})"

    zeilen = [("Liegenschaft", f"={MH}B4", None), ("Vermieter", f"={MH}B5", None),
              ("Mieter", f'={nv("B")}&""', None),
              ("Einheit", f'={nv("C")}&"  "&IFERROR(INDEX({MH}$B${EIN0}:$B${EINL},MATCH({nv("C")},{EINH_NR},0)),"")', None)]
    for i, (lab, f, fmt) in enumerate(zeilen, 3):
        label(b, f"A{i}", lab)
        b.cell(i, 2, f).font = schrift()
    label(b, "A7", "Abrechnungsperiode")
    b["B7"], b["C7"] = "=Nebenkosten!B7", "=Nebenkosten!D7"
    label(b, "A8", "Ihre Mietdauer in der Periode")
    b["B8"], b["C8"], b["D8"] = f"={wert_oder_leer(nv('D'))}", f"={wert_oder_leer(nv('E'))}", f'={nv("F")}&" Tage"'
    for z in "B7 C7 B8 C8".split():
        b[z].number_format = DATUM
        b[z].alignment = Alignment(horizontal="left")
        b[z].font = schrift()
    b["D8"].font = schrift()
    kopf(b, 10, ["Kostenart", "Gesamtkosten CHF", "Schlüssel", "Verteilung", "Ihr Anteil", "von Total",
                 "Zeitanteil", "Ihr Betrag CHF"])
    for j in range(K_N):
        z, kr = 11 + j, K0 + j
        da = f'Nebenkosten!A{kr}=""'
        b.cell(z, 1, f'=IF({da},"",Nebenkosten!A{kr})')
        b.cell(z, 2, f'=IF({da},"",Nebenkosten!B{kr})').number_format = CHF
        b.cell(z, 3, f'=IF({da},"",Nebenkosten!C{kr})')
        b.cell(z, 4, f'=IF({da},"",Nebenkosten!D{kr})')
        b.cell(z, 5, f'=IF({da},"",IF(Nebenkosten!C{kr}="Fläche",{nv("I")},IF(Nebenkosten!C{kr}="Anteil",{nv("J")},{nv("K")})))').number_format = "0.##"
        b.cell(z, 6, f'=IF({da},"",Nebenkosten!E{kr})').number_format = "0.##"
        b.cell(z, 7, f'=IF({da},"",IF(Nebenkosten!D{kr}="Heizung",{nv("H")},{nv("G")}))').number_format = "0.0%"
        b.cell(z, 8, f'=IF({da},"",{nv(col(D0+j))})').number_format = CHF
        for c_ in range(1, 9):
            b.cell(z, c_).font = schrift(9)
    s1 = 11 + K_N + 1
    saldo = nv(col(TOT + 2))
    for i, (lab, f) in enumerate((("Total Ihr Anteil (gerundet auf 5 Rp.)", f"={nv(col(TOT))}"),
                                  ("abzüglich Akontozahlungen", f"=-{nv(col(TOT+1))}"),
                                  (f'=IF({saldo}>=0,"Nachzahlung zu Ihren Lasten","Guthaben zu Ihren Gunsten")', f"=ABS({saldo})"))):
        b.cell(s1 + i, 5, lab).font = schrift(10 + (i == 2), i != 1)
        b.cell(s1 + i, 8, f).number_format = CHF
        b.cell(s1 + i, 8).font = schrift(10 + (i == 2), i != 1)
    b.cell(s1 + 4, 1, "Die Belege können Sie nach Vereinbarung einsehen (Art. 257b Abs. 2 OR). Einwände bitte innert 30 Tagen schriftlich.").font = schrift(8, kursiv=True)
    b.cell(s1 + 5, 1, "Nachzahlung zahlbar innert 30 Tagen. Ein Guthaben wird mit der nächsten Miete verrechnet oder überwiesen.").font = schrift(8, kursiv=True)
    breiten(b, [32, 14, 11, 10, 10, 9, 10, 14, 2, 30])
    b.print_area = f"A1:H{s1+5}"
    b.page_setup.fitToWidth = 1
    b.page_setup.fitToHeight = 1
    b.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(b)
    return s1


def blatt_mieterwechsel(wb, beispiel):
    w = wb.create_sheet("Mieterwechsel")
    titel(w, "Mieterwechsel & Fristen")
    leitfaden(w, 2, "Pro Kündigung eine Zeile: Mietverhältnis-ID, wer kündigt, Eingangsdatum.",
              "Mietende auf den nächsten zulässigen Termin, zehn Schritte mit Datum, Leerstand in Franken.",
              "Spalte «Nächste Frist»: rot überfällig, gelb innert 7 Tagen. Erledigtes mit x markieren.",
              "Blatt «Checkliste»: ein Mieterwechsel als druckbare Liste.")
    w.row_dimensions[2].height = 13
    kopf_txt = ["ID", "Mieter", "Einheit", "Kündigung durch", "Eingang Kündigung", "Frist (Mt.)",
                "Vereinbartes Ende", "Ende Mietverhältnis", "Neuer Mieter ab", "Leerstand Tage", "Leerstandskosten"]
    for name, _ in SCHRITTE:
        kopf_txt += [name, "erl."]
    kopf_txt += ["Nächste Frist", "Tage", "Nächster Schritt", "Hinweis"]
    kopf(w, W0 - 1, kopf_txt, hoehe=54)
    dv_id = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MV_N), allow_blank=True)
    dv_d = DataValidation(type="list", formula1='"Mieter,Vermieter,Ausserterminlich"', allow_blank=True)
    dv_x = DataValidation(type="list", formula1='"x"', allow_blank=True)
    for d in (dv_id, dv_d, dv_x):
        w.add_data_validation(d)
    for i, (name, _) in enumerate(SCHRITTE):
        w.cell(W0 - 2, OFFEN0 + i, name).font = schrift(7, farbe="888888")
    for i in range(W_N):
        z = W0 + i
        eingabe(w.cell(z, 1), "0"); dv_id.add(w.cell(z, 1))
        formel(w.cell(z, 2), f'=IF(A{z}="","",IFERROR(INDEX({mh("C")},A{z})&"",""))')
        formel(w.cell(z, 3), f'=IF(A{z}="","",IFERROR(INDEX({mh("B")},A{z})&"",""))')
        for c_, fmt in ((4, None), (5, DATUM), (6, "0"), (7, DATUM), (9, DATUM)):
            eingabe(w.cell(z, c_), fmt)
        dv_d.add(w.cell(z, 4))
        frist = f'IF(F{z}="",3,F{z})'
        for k in range(24):
            kc = w.cell(z, KAND0 + k)
            kc.value = (f'=IF(E{z}="","",IF(OR(COUNTIF({FLAGS},"x")=0,INDEX({FLAGS},MONTH(EOMONTH(E{z},{frist}+{k})))="x"),'
                        f'EOMONTH(E{z},{frist}+{k}),""))')
        kand = f"{col(KAND0)}{z}:{col(KAND0+23)}{z}"
        formel(w.cell(z, 8), f'=IF(G{z}<>"",G{z},IF(E{z}="","",IF(COUNT({kand})=0,"",MIN({kand}))))', DATUM)
        formel(w.cell(z, 10), f'=IF(OR(H{z}="",I{z}=""),"",MAX(0,I{z}-H{z}-1))', "0")
        formel(w.cell(z, 11), f'=IF(J{z}="","",ROUND(N(IFERROR(INDEX({mh("G")},A{z}),0))*12/365*J{z},2))', CHF)
        H, E, I = f"H{z}", f"E{z}", f"I{z}"
        ab = f"{col(STEP0 + 8)}{z}"         # Abnahme-Datum (Schritt 5)
        ps = "Nebenkosten!$B$7"
        formeln = [
            f'=IF({E}="","",{E}+7)',
            f'=IF({E}="","",{E})',
            f'=IF({H}="","",WORKDAY({H},-10,{FEIER}))',
            f'=IF({H}="","",{H}-30)',
            f'=IF({H}="","",WORKDAY({H}-1,1,{FEIER}))',
            f'=IF({ab}="","",WORKDAY({ab},2,{FEIER}))',
            f'=IF({I}="","",{I})',
            f'=IF({ab}="","",{ab}+30)',
            f'=IF(OR({H}="",{ps}=""),"",EDATE({ps},12*MAX(1,ROUNDUP(((YEAR({H})-YEAR({ps}))*12+MONTH({H})-MONTH({ps})+1)/12,0))))',
            f'=IF({H}="","",EDATE({H},12))',
        ]
        for s, f in enumerate(formeln):
            dc, ec = STEP0 + 2 * s, STEP0 + 2 * s + 1
            formel(w.cell(z, dc), f, DATUM)
            eingabe(w.cell(z, ec)); dv_x.add(w.cell(z, ec))
            w.cell(z, OFFEN0 + s).value = f'=IF(AND({col(dc)}{z}<>"",{col(ec)}{z}=""),{col(dc)}{z},"")'
        offen = f"{col(OFFEN0)}{z}:{col(OFFEN0+len(SCHRITTE)-1)}{z}"
        nx = col(NEXT_C)
        formel(w.cell(z, NEXT_C), f'=IF(COUNT({offen})=0,"",MIN({offen}))', DATUM)
        formel(w.cell(z, TAGE_C), f'=IF({nx}{z}="","",{nx}{z}-TODAY())', "0")
        formel(w.cell(z, SCHRITT_C), f'=IF({nx}{z}="","",INDEX(${col(OFFEN0)}${W0-2}:${col(OFFEN0+len(SCHRITTE)-1)}${W0-2},MATCH({nx}{z},{offen},0)))')
        formel(w.cell(z, HINW_C),
               f'=IF(A{z}="","",IF(B{z}="","ID ohne Mieter in «Mein Haus»",'
               f'IF(AND(IFERROR(INDEX({mh("E")},A{z}),"")<>"",H{z}<>"",IFERROR(INDEX({mh("E")},A{z}),"")<>H{z}),'
               f'"Mietende in «Mein Haus» weicht ab",'
               f'IF(D{z}="Vermieter","Amtliches Formular? Familienwohnung: beide Ehegatten.",'
               f'IF(D{z}="Ausserterminlich","Nachmieter zumutbar? Vereinbartes Ende eintragen.","")))))')
        w.cell(z, SCHLUESSEL_C).value = f'=IF({nx}{z}="","",{nx}{z}+ROW()/100000)'
        if beispiel and i < len(B_WECHSEL):
            id_, durch, eing, fr, vend, neu, erl = B_WECHSEL[i]
            for c_, val in zip((1, 4, 5, 6, 7, 9), (id_, durch, eing, fr, vend, neu)):
                w.cell(z, c_).value = val
            if erl:
                for s in range(len(SCHRITTE)):
                    w.cell(z, STEP0 + 2 * s + 1).value = "x"
    t = col(TAGE_C)
    rng = f"{col(NEXT_C)}{W0}:{t}{WL}"
    w.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(${t}{W0}<>"",${t}{W0}<0)'], fill=ROT))
    w.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(${t}{W0}<>"",${t}{W0}>=0,${t}{W0}<=7)'], fill=GELB))
    w.conditional_formatting.add(f"{col(HINW_C)}{W0}:{col(HINW_C)}{WL}",
                                 FormulaRule(formula=[f'LEFT({col(HINW_C)}{W0},2)="ID"'], fill=ROT))
    breiten(w, [5, 18, 7, 13, 11, 6, 11, 11, 11, 8, 11])
    for s in range(len(SCHRITTE)):
        w.column_dimensions[col(STEP0 + 2 * s)].width = 11
        w.column_dimensions[col(STEP0 + 2 * s + 1)].width = 4
    breiten(w, [11, 6, 26, 40], NEXT_C)
    for c_ in range(KAND0, SCHLUESSEL_C + 1):
        w.column_dimensions[col(c_)].hidden = True
    w.freeze_panes = f"D{W0}"
    schuetzen(w)

    # ---------------------------------------------------- Checkliste
    c = wb.create_sheet("Checkliste")
    label(c, "A1", "Checkliste Mieterwechsel", g=16, farbe=AKZENT)
    label(c, "F1", "Zeile im Blatt «Mieterwechsel» (wird nicht gedruckt):", False, 9, True)
    eingabe(c["F2"], "0", 1)
    dv_n = DataValidation(type="whole", operator="between", formula1="1", formula2=str(W_N))
    c.add_data_validation(dv_n); dv_n.add("F2")
    sel = "$F$2"

    def mw(spalte):
        return f"INDEX(Mieterwechsel!${spalte}${W0}:${spalte}${WL},{sel})"

    for i, (lab, f, fmt) in enumerate((("Liegenschaft", f"={MH}B4", None), ("Mieter", f"={mw('B')}&\"\"", None),
                                       ("Einheit", f"={mw('C')}&\"\"", None),
                                       ("Ende Mietverhältnis", f"={wert_oder_leer(mw('H'))}", DATUM),
                                       ("Neuer Mieter ab", f"={wert_oder_leer(mw('I'))}", DATUM)), 3):
        label(c, f"A{i}", lab)
        c.cell(i, 2, f).font = schrift()
        if fmt:
            c.cell(i, 2).number_format = fmt
            c.cell(i, 2).alignment = Alignment(horizontal="left")
    kopf(c, 9, ["Schritt", "Datum", "Status", "Was zu tun ist"])
    for s, (name, was) in enumerate(SCHRITTE):
        z = 10 + s
        dc, ec = col(STEP0 + 2 * s), col(STEP0 + 2 * s + 1)
        c.cell(z, 1, f"{s+1}. {name}").font = schrift(10, True)
        c.cell(z, 2, f"={wert_oder_leer(mw(dc))}").number_format = DATUM
        c.cell(z, 3, f'=IF({mw(ec)}="x","erledigt",IF({mw(dc)}="","–",IF({mw(dc)}<TODAY(),"überfällig","offen")))')
        c.cell(z, 4, was).alignment = Alignment(wrap_text=True, vertical="top")
        for c_ in range(1, 5):
            c.cell(z, c_).border = rahmen
            if c_ > 1:
                c.cell(z, c_).font = schrift(9)
        c.row_dimensions[z].height = 30
    c.conditional_formatting.add("C10:C19", FormulaRule(formula=['C10="überfällig"'], fill=ROT))
    c.conditional_formatting.add("C10:C19", FormulaRule(formula=['C10="erledigt"'], fill=GRUEN))
    breiten(c, [38, 12, 11, 70, 2, 30])
    c.print_area = "A1:D20"
    c.page_setup.orientation = "landscape"
    c.page_setup.fitToWidth = 1
    c.page_setup.fitToHeight = 1
    c.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(c)


def blatt_mietzins(wb, beispiel):
    m = wb.create_sheet("Mietzins")
    titel(m, "Mietzinsrechner")
    leitfaden(m, 2, "Mietverhältnis wählen; Basiswerte kommen aus «Mein Haus». Neuen Referenzzinssatz, aktuellen LIK und Stichtag eintragen.",
              "Referenzzins nach Art. 13 VMWG, Teuerung zu 40 % (Art. 16 VMWG), pauschale Kostensteigerung pro Jahr.",
              "Senkung: Anspruch des Mieters auf Begehren. Erhöhung: nur mit amtlichem Formular, frühestens auf den angezeigten Termin.",
              "Werte für das amtliche Formular bzw. die Antwort auf ein Senkungsbegehren.")
    label(m, "A7", "Mietverhältnis-ID")
    eingabe(m["B7"], "0", B_MIETZINS["id"] if beispiel else 1)
    ID = "$B$7"

    def h(spalte):
        return f"INDEX({mh(spalte)},{ID})"

    basis = [(8, "Mieter", f'=IFERROR({h("C")}&"","")', None),
             (9, "Nettomiete heute / Mt.", f'=IFERROR(N({h("G")}),0)', CHF),
             (10, "Referenzzins Basis (%)", f'=IFERROR(N({h("K")}),0)', "0.00"),
             (11, "LIK Basis", f'=IFERROR({wert_oder_leer(h("L"))},"")', "0.0"),
             (12, "Basis seit", f'=IFERROR({wert_oder_leer(h("M"))},"")', DATUM)]
    for z, lab, f, fmt in basis:
        label(m, f"A{z}", lab, False)
        formel(m.cell(z, 2), f, fmt)
    label(m, "A14", "Eingaben", g=11, farbe=AKZENT)
    einga = [(15, "Referenzzins neu (%)", B_MIETZINS["ref"], "0.00"),
             (16, "LIK aktuell (Punkte)", B_MIETZINS["lik"] if beispiel else None, "0.0"),
             (17, "Stichtag der Berechnung", B_MIETZINS["stichtag"] if beispiel else None, DATUM),
             (18, "Allgemeine Kostensteigerung pro Jahr (%)", B_MIETZINS["kosten"] if beispiel else None, "0.00"),
             (19, "Mitteilung zugestellt am (nur Erhöhung)", B_MIETZINS["mitteilung"] if beispiel else None, DATUM),
             (20, "Kündigungsfrist (Monate)", 3, "0")]
    for z, lab, wert, fmt in einga:
        label(m, f"A{z}", lab, False)
        eingabe(m.cell(z, 2), fmt, wert)
    label(m, "C15", "seit 2.9.2025: 1.25 % (BWO); vor jeder Verwendung prüfen", False, 8, True, "555555")
    label(m, "C16", "Landesindex der Konsumentenpreise, BFS", False, 8, True, "555555")
    label(m, "C18", "Gerichtspraxis, kantonal verschieden; leer = nicht berücksichtigt", False, 8, True, "555555")
    label(m, "A22", "Berechnung", g=11, farbe=AKZENT)
    rechnung = [
        (23, "Schritte zu 0.25 Prozentpunkten", "=ROUND((B15-B10)/0.25,6)", "0.00"),
        (24, "Anpassung Referenzzins", "=IF(B23>=0,0.03*B23,-(1-1/(1+0.03*-B23)))", PROZ),
        (25, "Teuerung (40 % der LIK-Veränderung)", '=IF(OR(B11="",B16="",N(B11)=0),0,(B16/B11-1)*0.4)', PROZ),
        (26, "Allgemeine Kostensteigerung", '=IF(OR(B18="",B12="",B17=""),0,B18/100*(B17-B12)/365.25)', PROZ),
        (27, "Total Anpassung", "=B24+B25+B26", PROZ),
        (28, "Neue Nettomiete / Mt.", "=ROUND(B9*(1+B27)*20,0)/20", CHF),
        (29, "Veränderung / Mt.", "=B28-B9", CHF),
    ]
    for z, lab, f, fmt in rechnung:
        label(m, f"A{z}", lab, z in (27, 28))
        formel(m.cell(z, 2), f, fmt, z in (27, 28))
    label(m, "A31", "Ergebnis")
    formel(m["B31"], '=IF(B9=0,"Mietverhältnis ohne Nettomiete",IF(B27<0,"Senkung: Der Mieter kann eine Herabsetzung auf diesen Betrag verlangen.",'
                     'IF(B27>0,"Erhöhung möglich – nur mit amtlichem Formular","Keine Anpassung")))', None, True)
    m.merge_cells("B31:C31")
    label(m, "A32", "Frühestens wirksam ab (Erhöhung)")
    kand = f"C40:{col(3+23)}40"
    formel(m["B32"], f'=IF(OR(B27<=0,B19=""),"",IF(COUNT({kand})=0,"",MIN({kand})+1))', DATUM)
    label(m, "A33", "Prüfen")
    formel(m["B33"], '=IF(B9=0,"",IF(ABS(B23-ROUND(B23,0))>0.000001,"Basis-Referenzzins prüfen (keine 0.25-Schritte)",'
                     'IF(OR(B10>5,B15>5),"Über 5 %: andere Sätze nach Art. 13 VMWG – nicht abgebildet",'
                     'IF(AND(B25<>0,B11=""),"LIK-Basis fehlt",""))))')
    m.conditional_formatting.add("B33", FormulaRule(formula=['B33<>""'], fill=ROT))
    m["A39"] = "Hilfszeile: mögliche Termine (Mitteilung + 10 Tage vor Beginn der Kündigungsfrist, konservativ)"
    m["A39"].font = schrift(7, farbe="888888")
    for k in range(24):
        c = m.cell(40, 3 + k, f'=IF($B$19="","",IF(OR(COUNTIF({FLAGS},"x")=0,INDEX({FLAGS},MONTH(EOMONTH($B$19+10,$B$20+{k})))="x"),'
                              f'EOMONTH($B$19+10,$B$20+{k}),""))')
        c.number_format = DATUM
        c.font = schrift(7, farbe="888888")
    hinweise = ["Rechtliche Hinweise (ohne Gewähr):",
                "• Referenzzins unter 5 %: +3 % pro 0.25 Prozentpunkte Erhöhung; Senkung entsprechend 2.91 % pro Schritt (1 − 1/(1 + 3 % × Schritte)).",
                "• Teuerung: 40 % der Veränderung des Landesindex seit der letzten Anpassung (Art. 16 VMWG).",
                "• Allgemeine Kostensteigerung: Pauschale nach Gerichtspraxis; Höhe kantonal verschieden, im Streitfall Nachweis nötig.",
                "• Erhöhungen nur mit amtlichem Formular des Kantons, mindestens 10 Tage vor Beginn der Kündigungsfrist (Art. 269d OR).",
                "• Senkungen schuldet der Vermieter auf Begehren des Mieters (Art. 270a OR). Nicht abgebildet: absolute Methode, Mietzinsreserven, wertvermehrende Investitionen."]
    for i, t in enumerate(hinweise):
        m.cell(42 + i, 1, t).font = schrift(9, i == 0, kursiv=i > 0)
        m.merge_cells(start_row=42 + i, start_column=1, end_row=42 + i, end_column=3)
        m.cell(42 + i, 1).alignment = Alignment(wrap_text=True, vertical="top")
        m.row_dimensions[42 + i].height = 24 if i else 14
    m.merge_cells("B33:C33")
    for z in range(2, 6):
        m.merge_cells(start_row=z, start_column=1, end_row=z, end_column=3)
    breiten(m, [40, 18, 50])
    m.print_area = "A1:C47"
    m.page_setup.fitToWidth = 1
    m.page_setup.fitToHeight = 1
    m.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(m)


def blatt_rueckgabe(wb, beispiel):
    r = wb.create_sheet("Rückgabe")
    titel(r, "Wohnungsrückgabe: Protokoll, Schaden & Kaution")
    label(r, "A3", "Auszug für Protokoll, Abrechnung und Freigabe: Mietverhältnis-ID")
    eingabe(r["E3"], "0", B_FALL["id"] if beispiel else None)
    leitfaden(r, 2, "Pro Auszug eine Zeile «Fälle»; Mängel unten mit derselben ID. Lebensdauer aus der paritätischen Tabelle HEV/MV.",
              "Mieteranteil = Kosten × (1 − Alter ÷ Lebensdauer); normale Abnutzung und Vorbestehendes = 0. Kaution wird aufgeteilt.",
              "Spalte «An Vermieter» / «An Mieter»; «Prüfen» meldet fehlende Angaben.",
              "Blätter «Protokoll», «Kautionsabrechnung», «Freigabe Kaution» für die oben gewählte ID.", spalte=8)
    kopf(r, F0 - 1, ["ID", "Mieter", "Mietende", "Abnahme am", "Anwesend", "Zähler Strom", "Zähler Wasser",
                     "Zähler Wärme", "Schlüssel erhalten", "Kaution", "Kautionszins", "Ausstehende Mietzinse",
                     "NK-Saldo (aus Nebenkosten)", "NK-Saldo manuell", "Weitere Forderungen", "Schäden zulasten Mieter",
                     "Total Forderungen", "Kaution inkl. Zins", "An Vermieter", "An Mieter", "Restforderung",
                     "Zusätzlich vom Vermieter", "Prüfen"], hoehe=42)
    dv_id = DataValidation(type="whole", operator="between", formula1="1", formula2=str(MV_N), allow_blank=True)
    r.add_data_validation(dv_id); dv_id.add("E3")
    for i in range(F_N):
        z = F0 + i
        eingabe(r.cell(z, 1), "0"); dv_id.add(r.cell(z, 1))
        formel(r.cell(z, 2), f'=IF(A{z}="","",IFERROR(INDEX({mh("C")},A{z})&"",""))')
        formel(r.cell(z, 3), f'=IF(A{z}="","",IFERROR({wert_oder_leer(f"INDEX({mh(chr(70))},A{z})")},""))', DATUM)
        for c_, fmt in ((4, DATUM), (5, None), (6, "0"), (7, "0"), (8, "0"), (9, None), (11, CHF), (12, CHF),
                        (14, CHF), (15, CHF)):
            eingabe(r.cell(z, c_), fmt)
        formel(r.cell(z, 10), f'=IF(A{z}="","",IFERROR(N(INDEX({mh("I")},A{z})),0))', CHF)
        formel(r.cell(z, 13), f'=IF(A{z}="","",IFERROR(INDEX({NV}${SALDO}${NV0}:${SALDO}${NVL},A{z}),0))', CHF)
        formel(r.cell(z, 16), f'=IF(A{z}="","",SUMIF($A${M0}:$A${ML},A{z},$K${M0}:$K${ML}))', CHF)
        formel(r.cell(z, 17), f'=IF(A{z}="","",P{z}+N(L{z})+IF(N{z}<>"",N{z},N(M{z}))+N(O{z}))', CHF, True)
        formel(r.cell(z, 18), f'=IF(A{z}="","",N(J{z})+N(K{z}))', CHF)
        formel(r.cell(z, 19), f'=IF(A{z}="","",MAX(0,MIN(Q{z},R{z})))', CHF, True)
        formel(r.cell(z, 20), f'=IF(A{z}="","",R{z}-S{z})', CHF, True)
        formel(r.cell(z, 21), f'=IF(A{z}="","",MAX(0,Q{z}-R{z}))', CHF)
        formel(r.cell(z, 22), f'=IF(A{z}="","",MAX(0,-Q{z}))', CHF)
        formel(r.cell(z, 23), f'=IF(A{z}="","",IF(B{z}="","ID ohne Mieter",IF(C{z}="","Mietende fehlt",'
                              f'IF(AND(D{z}<>"",D{z}-C{z}>5),"Abnahme > 5 Tage nach Mietende: Rügefrist",""))))')
    if beispiel:
        f = B_FALL
        for c_, val in zip((1, 4, 5, 6, 7, 8, 9, 11, 12, 15),
                           (f["id"], f["abnahme"], f["anwesend"], f["strom"], f["wasser"], f["waerme"],
                            f["schluessel"], f["zins"], f["miete"], f["weitere"])):
            r.cell(F0, c_).value = val
    r.conditional_formatting.add(f"W{F0}:W{FL}", FormulaRule(formula=[f'W{F0}<>""'], fill=ROT))

    label(r, f"A{M0-2}", "Mängel (Abnahmeprotokoll)", g=11, farbe=AKZENT)
    kopf(r, M0 - 1, ["ID", "Raum", "Bauteil", "Mangel / Zustand", "Ursache", "Kosten laut Offerte",
                     "Einbau / letzte Erneuerung", "Lebensdauer Jahre", "Alter bei Mietende", "Restwert",
                     "Zulasten Mieter", "Prüfen", "Nr. im Ausdruck", "Mietende"], hoehe=42)
    dv_u = DataValidation(type="list", formula1='"Mieter,Normale Abnutzung,Vorbestehend"', allow_blank=True)
    r.add_data_validation(dv_u)
    sel = "$E$3"
    for i in range(M_N):
        z = M0 + i
        eingabe(r.cell(z, 1), "0"); dv_id.add(r.cell(z, 1))
        for c_, fmt in ((2, None), (3, None), (4, None), (5, None), (6, CHF), (7, DATUM), (8, "0")):
            eingabe(r.cell(z, c_), fmt)
        dv_u.add(r.cell(z, 5))
        formel(r.cell(z, 14), f'=IF(A{z}="","",IFERROR({wert_oder_leer(f"INDEX({mh(chr(70))},A{z})")},""))', DATUM)
        formel(r.cell(z, 9), f'=IF(OR(G{z}="",N{z}=""),"",(N{z}-G{z})/365.25)', "0.0")
        formel(r.cell(z, 10), f'=IF(F{z}="","",IF(E{z}<>"Mieter",0,IF(H{z}="",1,IF(OR(G{z}="",N{z}=""),"",MAX(0,1-I{z}/H{z})))))', "0%")
        formel(r.cell(z, 11), f'=IF(OR(J{z}="",F{z}=""),0,ROUND(F{z}*J{z}*20,0)/20)', CHF)
        formel(r.cell(z, 12), f'=IF(AND(A{z}="",B{z}="",F{z}=""),"",IF(A{z}="","ID fehlt",IF(F{z}="","",IF(F{z}<0,"Kosten negativ",'
                              f'IF(E{z}="","Ursache wählen",IF(N{z}="","Mietende fehlt",IF(AND(G{z}<>"",G{z}>N{z}),"Einbau nach Mietende",'
                              f'IF(AND(E{z}="Mieter",H{z}<>"",G{z}=""),"Einbaudatum fehlt",'
                              f'IF(AND(E{z}="Mieter",H{z}<>"",H{z}<=0),"Lebensdauer ungültig","")))))))))')
        formel(r.cell(z, 13), f'=IF(AND(A{z}<>"",A{z}={sel}),COUNTIF($A${M0}:A{z},{sel}),"")', "0")
        if beispiel and i < len(B_MAENGEL):
            for c_, val in enumerate(B_MAENGEL[i], 1):
                r.cell(z, c_).value = val
    r.conditional_formatting.add(f"L{M0}:L{ML}", FormulaRule(formula=[f'L{M0}<>""'], fill=ROT))
    breiten(r, [5, 18, 15, 30, 17, 12, 12, 10, 10, 9, 12, 20, 9, 11, 12, 12, 13, 11, 12, 12, 11, 11, 26])
    r.freeze_panes = f"B{F0}"
    schuetzen(r)


def fall(spalte):
    return f'IFERROR(INDEX(Rückgabe!${spalte}${F0}:${spalte}${FL},MATCH(Rückgabe!$E$3,Rückgabe!$A${F0}:$A${FL},0)),"")'


def mangel(spalte, k):
    return (f'IFERROR(INDEX(Rückgabe!${spalte}${M0}:${spalte}${ML},'
            f'MATCH({k},Rückgabe!$M${M0}:$M${ML},0)),"")')


def kopfblock(ws, start):
    zeilen = [("Liegenschaft", f"={MH}B4", None), ("Mieter", f'={fall("B")}&""', None),
              ("Einheit", f'=IFERROR(INDEX({mh("B")},Rückgabe!$E$3)&"  "&INDEX({MH}$B${EIN0}:$B${EINL},'
                          f'MATCH(INDEX({mh("B")},Rückgabe!$E$3),{EINH_NR},0)),"")', None),
              ("Mietende", f"={fall('C')}", DATUM), ("Abnahme am", f"={fall('D')}", DATUM)]
    for i, (lab, f, fmt) in enumerate(zeilen):
        label(ws, f"A{start+i}", lab)
        c = ws.cell(start + i, 2, f)
        c.font = schrift()
        if fmt:
            c.number_format = fmt
            c.alignment = Alignment(horizontal="left")
    return start + len(zeilen)


def blatt_protokoll(wb):
    p = wb.create_sheet("Protokoll")
    label(p, "A1", "Wohnungsabnahmeprotokoll", g=16, farbe=AKZENT)
    label(p, "A2", "Auswahl über die ID oben im Blatt «Rückgabe». Leere Zeilen eignen sich für Notizen vor Ort.", False, 8, True, "555555")
    z = kopfblock(p, 4)
    for i, (lab, f) in enumerate((("Anwesend", f'={fall("E")}&""'),
                                  ("Zähler Strom / Wasser / Wärme", f'={fall("F")}&" / "&{fall("G")}&" / "&{fall("H")}'),
                                  ("Schlüssel erhalten", f'={fall("I")}&""'))):
        label(p, f"A{z+i}", lab)
        p.cell(z + i, 2, f).font = schrift()
    t0 = z + 4
    kopf(p, t0, ["Raum", "Bauteil", "Mangel / Zustand", "Ursache"])
    for k in range(1, DRUCK + 1):
        zz = t0 + k
        for c_, sp in enumerate("BCDE", 1):
            c = p.cell(zz, c_, f"={mangel(sp, k)}")
            c.border = rahmen
            c.font = schrift(9)
    e = t0 + DRUCK + 2
    p.cell(e, 1, "Versteckte Mängel werden nach ihrer Entdeckung gerügt (Art. 267a Abs. 2 OR). Der Mieter erhält eine Kopie dieses Protokolls.").font = schrift(8, kursiv=True)
    p.cell(e + 3, 1, "Mieter: ____________________________").font = schrift()
    p.cell(e + 3, 3, "Vermieter: ____________________________").font = schrift()
    breiten(p, [30, 18, 44, 18])
    p.print_area = f"A1:D{e+4}"
    p.page_setup.fitToWidth = 1
    p.page_setup.fitToHeight = 1
    p.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(p)


def blatt_kautionsabrechnung(wb):
    a = wb.create_sheet("Kautionsabrechnung")
    label(a, "A1", "Abrechnung Wohnungsrückgabe und Mietkaution", g=15, farbe=AKZENT)
    z = kopfblock(a, 3)
    t0 = z + 1
    kopf(a, t0, ["Raum · Bauteil", "Mangel", "Ursache", "Kosten CHF", "Alter J.", "Lebensdauer J.", "Restwert", "Zulasten Mieter CHF"])
    for k in range(1, DRUCK + 1):
        zz = t0 + k
        a.cell(zz, 1, f'=IF({mangel("B", k)}="","",{mangel("B", k)}&" · "&{mangel("C", k)})')
        for c_, sp, fmt in ((2, "D", None), (3, "E", None), (4, "F", CHF), (5, "I", "0.0"), (6, "H", "0"),
                            (7, "J", "0%"), (8, "K", CHF)):
            c = a.cell(zz, c_, f'=IF({mangel("B", k)}="","",{mangel(sp, k)})')
            if fmt:
                c.number_format = fmt
        for c_ in range(1, 9):
            a.cell(zz, c_).font = schrift(9)
    s = t0 + DRUCK + 2
    posten = [("Schäden zulasten Mieter", "P"), ("Ausstehende Mietzinse", "L"), ("Nebenkosten-Saldo", None),
              ("Weitere Forderungen", "O"), ("Total Forderungen", "Q"), ("Mietkaution inkl. Zins", "R"),
              ("Aus der Kaution an den Vermieter", "S"), ("Aus der Kaution an den Mieter", "T"),
              ("Restforderung über die Kaution hinaus", "U"), ("Zusätzlich vom Vermieter zu bezahlen", "V")]
    for i, (lab, sp) in enumerate(posten):
        fett = lab.startswith(("Total", "Aus der"))
        a.cell(s + i, 4, lab).font = schrift(10, fett)
        f = (f'=IF({fall("N")}<>"",{fall("N")},N({fall("M")}))' if sp is None else f"=N({fall(sp)})")
        c = a.cell(s + i, 8, f)
        c.number_format = CHF
        c.font = schrift(10, fett)
    a.cell(s + 11, 1, "Berechnung: Kosten × (1 − Alter ÷ Lebensdauer) gemäss paritätischer Lebensdauertabelle HEV/MV. "
                      "Normale Abnutzung geht nicht zulasten des Mieters (Art. 267 OR).").font = schrift(8, kursiv=True)
    a.cell(s + 12, 1, "Offerten und Rechnungen können eingesehen werden. Einwände bitte innert 30 Tagen schriftlich.").font = schrift(8, kursiv=True)
    breiten(a, [28, 30, 15, 12, 8, 11, 9, 15])
    a.print_area = f"A1:H{s+12}"
    a.page_setup.fitToWidth = 1
    a.page_setup.fitToHeight = 1
    a.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(a)
    return s


def blatt_freigabe(wb):
    f = wb.create_sheet("Freigabe Kaution")
    label(f, "A1", "Freigabe der Mietkaution", g=15, farbe=AKZENT)
    zeilen = [(3, "An", f'=IFERROR(INDEX({mh("J")},Rückgabe!$E$3)&"","")', None),
              (5, "Mietobjekt", f'={MH}B4', None),
              (6, "Mietende", f"={fall('C')}", DATUM)]
    for z, lab, form, fmt in zeilen:
        label(f, f"A{z}", lab)
        c = f.cell(z, 2, form)
        c.font = schrift()
        if fmt:
            c.number_format = fmt
            c.alignment = Alignment(horizontal="left")
    f["A8"] = "Die Unterzeichnenden ersuchen Sie, das Mietkautionskonto zu saldieren und wie folgt auszuzahlen:"
    f["A8"].font = schrift()
    label(f, "A10", "An den Mieter")
    f["B10"] = f'={fall("B")}&""'
    f["C10"] = f"=N({fall('T')})"
    label(f, "A11", "IBAN Mieter")
    eingabe(f["B11"])
    label(f, "A12", "An den Vermieter")
    f["B12"] = f"={MH}B5"
    f["C12"] = f"=N({fall('S')})"
    label(f, "A13", "IBAN Vermieter")
    f["B13"] = f'={MH}B6&""'
    for z in (10, 12):
        f.cell(z, 2).font = schrift()
        f.cell(z, 3).number_format = CHF
        f.cell(z, 3).font = schrift(11, True)
    f["B13"].font = schrift()
    f["A14"] = "Allfälliger Zins bis zur Saldierung geht an den Mieter."
    f["A14"].font = schrift(9, kursiv=True)
    f["A17"] = "Ort, Datum:"
    f["B17"] = "______________________________"
    f["A20"] = "Mieter:"
    f["B20"] = "______________________________"
    f["A23"] = "Vermieter:"
    f["B23"] = "______________________________"
    for z in (17, 20, 23):
        f.cell(z, 1).font = schrift(10, True)
        f.cell(z, 2).font = schrift()
    breiten(f, [22, 46, 16])
    f.print_area = "A1:C24"
    schuetzen(f)


def blatt_kontrolle(wb, D0):
    k = wb.create_sheet("Kontrolle")
    titel(k, "Kontrolle", "Jede Zeile muss ihren Sollwert haben. Sonst zeigt die Spalte «Wo» die betroffene Stelle.")
    kopf(k, 3, ["Prüfung", "Ist", "Soll", "Wo"])
    dl = col(D0 + K_N - 1)
    pr = [
        ("Mietverhältnisse mit Hinweis (Einheit, Ende vor Beginn, Überschneidung)", f'=COUNTIF({MH}O{MV0}:O{MVL},"?*")', 0, "Mein Haus, Spalte «Prüfen»"),
        ("Monatsgewichte Heizung (Summe)", f"={EIN}F17", 1000, "Einstellungen"),
        ("Kosten ohne Schlüssel oder Verteilung",
         f'=SUMPRODUCT((Nebenkosten!B{K0}:B{KL}<>"")*((Nebenkosten!C{K0}:C{KL}="")+(Nebenkosten!D{K0}:D{KL}="")))', 0, "Nebenkosten"),
        ("Kostenarten mit Schlüssel-Total 0",
         f'=SUMPRODUCT((Nebenkosten!B{K0}:B{KL}<>"")*(Nebenkosten!C{K0}:C{KL}<>"")*(Nebenkosten!E{K0}:E{KL}=0))', 0, "Nebenkosten / Mein Haus"),
        ("Kosten erfasst, aber kein Periodenbeginn", f'=IF(AND(COUNT(Nebenkosten!B{K0}:B{KL})>0,Nebenkosten!B7=""),1,0)', 0, "Nebenkosten"),
        ("Nebenkosten: verteilt + Leerstand − Total",
         f"=ROUND(SUM({NV}{col(D0)}{NVL+1}:{dl}{NVL+1})+SUM({NV}{col(D0)}{NVL+2}:{dl}{NVL+2})-Nebenkosten!B{KL+1},2)", 0, "NK Verteilung"),
        ("Mieterwechsel: ID ohne Mieter oder abweichendes Mietende",
         f'=COUNTIF(Mieterwechsel!{col(HINW_C)}{W0}:{col(HINW_C)}{WL},"ID*")+COUNTIF(Mieterwechsel!{col(HINW_C)}{W0}:{col(HINW_C)}{WL},"Mietende*")', 0, "Mieterwechsel, Spalte «Hinweis»"),
        ("Mietzins: Hinweis", '=IF(Mietzins!B33="",0,1)', 0, "Mietzins, Zeile «Prüfen»"),
        ("Rückgabe: Fälle mit Hinweis", f'=COUNTIF(Rückgabe!W{F0}:W{FL},"?*")', 0, "Rückgabe, Spalte «Prüfen»"),
        ("Rückgabe: Mängel mit Hinweis", f'=COUNTIF(Rückgabe!L{M0}:L{ML},"?*")', 0, "Rückgabe, Mängel «Prüfen»"),
        ("Rückgabe: Kaution aufgeteilt − Kaution",
         f"=ROUND(SUM(Rückgabe!S{F0}:T{FL})-SUM(Rückgabe!R{F0}:R{FL}),2)", 0, "Rückgabe"),
    ]
    for i, (lab, f, soll, wo) in enumerate(pr, 4):
        k.cell(i, 1, lab).font = schrift(10)
        formel(k.cell(i, 2), f, "0.##")
        k.cell(i, 3, soll).font = schrift(9, farbe="555555")
        k.cell(i, 4, wo).font = schrift(9, kursiv=True, farbe="555555")
        k.conditional_formatting.add(f"B{i}", FormulaRule(formula=[f"B{i}<>C{i}"], fill=ROT))
    last = 4 + len(pr) - 1
    label(k, f"A{last+2}", "Status", g=12)
    formel(k.cell(last + 2, 2), f'=IF(SUMPRODUCT(--(B4:B{last}<>C4:C{last}))=0,"OK","Bitte prüfen")', None, True)
    k.conditional_formatting.add(f"B{last+2}", FormulaRule(formula=[f'B{last+2}="OK"'], fill=GRUEN))
    k.conditional_formatting.add(f"B{last+2}", FormulaRule(formula=[f'B{last+2}<>"OK"'], fill=ROT))
    breiten(k, [62, 12, 8, 34])
    schuetzen(k)
    return f"B{last+2}"


def blatt_start(wb, status_zelle):
    s = wb["Start"]
    label(s, "A1", "Vermieter-Ordner Schweiz", g=20, farbe=AKZENT)
    label(s, "A2", "Die wichtigsten Aufgaben einer privaten Mietliegenschaft – nach Schweizer Mietrecht, ohne Abo, ohne Makros.",
          False, 10, True, "555555")
    label(s, "A4", "So arbeiten Sie", g=12, farbe=AKZENT)
    kacheln = [("1 · Mein Haus", "Mein Haus", "Wohnungen und Mieter einmal erfassen"),
               ("2 · Mieterwechsel", "Mieterwechsel", "Kündigung → Mietende und zehn Schritte mit Datum"),
               ("3 · Rückgabe", "Rückgabe", "Protokoll, Schaden nach Lebensdauer, Kaution"),
               ("4 · Nebenkosten", "Nebenkosten", "Kosten verteilen, Abrechnung pro Mieter"),
               ("5 · Mietzins", "Mietzins", "Referenzzins, Teuerung, Kostensteigerung"),
               ("6 · Kontrolle", "Kontrolle", "Alles plausibel? Hier steht es.")]
    for i, (t, blatt, was) in enumerate(kacheln):
        z = 5 + i
        c = s.cell(z, 1, t)
        c.hyperlink = Hyperlink(ref=f"A{z}", location=f"'{blatt}'!A1", display=t)
        c.font = Font(name="Arial", size=11, bold=True, color="1F4E79", underline="single")
        s.cell(z, 2, was).font = schrift()
    label(s, "A12", "Was steht an?", g=12, farbe=AKZENT)
    kopf(s, 13, ["Datum", "Mieter", "Schritt", "Tage"])
    key = f"Mieterwechsel!${col(SCHLUESSEL_C)}${W0}:${col(SCHLUESSEL_C)}${WL}"
    for k in range(1, 6):
        z = 13 + k
        sm = f"SMALL({key},{k})"
        zeile = f"MATCH({sm},{key},0)"
        s.cell(z, 1, f'=IFERROR(INT({sm}),"")').number_format = DATUM
        s.cell(z, 2, f'=IFERROR(INDEX(Mieterwechsel!$B${W0}:$B${WL},{zeile}),"")')
        s.cell(z, 3, f'=IFERROR(INDEX(Mieterwechsel!${col(SCHRITT_C)}${W0}:${col(SCHRITT_C)}${WL},{zeile}),"")')
        s.cell(z, 4, f'=IFERROR(INT({sm})-TODAY(),"")').number_format = "0"
        for c_ in range(1, 5):
            s.cell(z, c_).border = rahmen
            s.cell(z, c_).font = schrift()
            s.cell(z, c_).alignment = Alignment(horizontal="left")
    s.conditional_formatting.add("A14:D18", FormulaRule(formula=['AND($D14<>"",$D14<0)'], fill=ROT))
    s.conditional_formatting.add("A14:D18", FormulaRule(formula=['AND($D14<>"",$D14>=0,$D14<=7)'], fill=GELB))
    label(s, "A21", "Kennzahlen", g=12, farbe=AKZENT)
    zahlen = [("Laufende Mietverhältnisse", f'=COUNTIF({MH}N{MV0}:N{MVL},"aktiv")', "0"),
              ("Nettomieten pro Monat", f'=SUMIF({MH}N{MV0}:N{MVL},"aktiv",{MH}G{MV0}:G{MVL})', CHF),
              ("Kautionen laufender Mietverhältnisse", f'=SUMIF({MH}N{MV0}:N{MVL},"aktiv",{MH}I{MV0}:I{MVL})', CHF),
              ("Einheiten ohne laufendes Mietverhältnis", f"=MAX(0,N({MH}F{ETOT})-B22)", "0"),
              ("Kontrolle", f"=Kontrolle!{status_zelle}", None)]
    for i, (lab, f, fmt) in enumerate(zahlen):
        z = 22 + i
        s.cell(z, 1, lab).font = schrift()
        formel(s.cell(z, 2), f, fmt, lab == "Kontrolle")
    s.conditional_formatting.add("B26", FormulaRule(formula=['B26="OK"'], fill=GRUEN))
    s.conditional_formatting.add("B26", FormulaRule(formula=['B26<>"OK"'], fill=ROT))
    hinweise = ["Gelb = Eingabe · Blau = rechnet automatisch (Blattschutz ohne Passwort).",
                "Ersetzt keine Rechtsberatung. Bei Streit: Schlichtungsbehörde Ihres Kantons oder Ihr Verband."]
    for i, t in enumerate(hinweise):
        s.cell(29 + i, 1, t).font = schrift(9, kursiv=True, farbe="555555")
    breiten(s, [34, 48, 40, 8])
    s.print_area = "A1:D30"
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.page_setup.fitToHeight = 1
    s.sheet_properties.pageSetUpPr.fitToPage = True
    schuetzen(s)


def baue(beispiel=True):
    wb = openpyxl.Workbook()
    wb.active.title = "Start"
    blatt_mein_haus(wb, beispiel)
    blatt_mieterwechsel(wb, beispiel)
    blatt_rueckgabe(wb, beispiel)
    blatt_protokoll(wb)
    s_abr = blatt_kautionsabrechnung(wb)
    blatt_freigabe(wb)
    D0, TOT = blatt_nebenkosten(wb, beispiel)
    assert col(TOT + 2) == SALDO
    blatt_nk_abrechnung(wb, D0, TOT)
    blatt_mietzins(wb, beispiel)
    blatt_einstellungen(wb, beispiel)
    status = blatt_kontrolle(wb, D0)
    blatt_start(wb, status)
    wb.active = 0
    name = "Vermieter-Ordner-CH-Beispiel.xlsx" if beispiel else "Vermieter-Ordner-CH.xlsx"
    wb.save(name)
    print("geschrieben:", name)
    return dict(abr=s_abr)


if __name__ == "__main__":
    art = sys.argv[1] if len(sys.argv) > 1 else "beide"
    if art in ("beispiel", "beide"):
        baue(True)
    if art in ("leer", "beide"):
        baue(False)
