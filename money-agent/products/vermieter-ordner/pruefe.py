"""Prüfstrecke für den Vermieter-Ordner.

Baut die Arbeitsmappe, erzeugt Varianten (Beispiel, Mietzins-Erhöhung,
Akonto-Abweichung, provozierte Fehler), lässt LibreOffice rechnen und
vergleicht jedes Modul mit einer unabhängigen Python-Rechnung –
inklusive der Datenflüsse zwischen den Modulen.

Aufruf: python3 pruefe.py
"""
import calendar
import os
import subprocess
import sys
import tempfile
from datetime import date, datetime, timedelta

import openpyxl

import build as B

HIER = os.path.dirname(os.path.abspath(__file__))
fehler = 0


def pruefe(ok, text):
    global fehler
    fehler += not ok
    print(("OK  " if ok else "ERR ") + text)


def d(x):
    return x.date() if isinstance(x, datetime) else x


# ------------------------------------------------------------ Referenzrechnungen
def monatsende(t, plus):
    m = t.month - 1 + plus
    j, m = t.year + m // 12, m % 12 + 1
    return date(j, m, calendar.monthrange(j, m)[1])


def mietende(eingang, frist, termine=(3, 6, 9)):
    for k in range(24):
        e = monatsende(eingang, frist + k)
        if e.month in termine:
            return e


def werktag(t, n):
    s = 1 if n > 0 else -1
    while n:
        t += timedelta(days=s)
        if t.weekday() < 5 and t not in B.FEIERTAGE:
            n -= s
    return t


def edate(t, monate):
    m = t.month - 1 + monate
    j, m = t.year + m // 12, m % 12 + 1
    return date(j, m, min(t.day, calendar.monthrange(j, m)[1]))


def schritte(eingang, ende, neu, ps):
    ab = werktag(ende - timedelta(days=1), 1)
    n = max(1, -(-((ende.year - ps.year) * 12 + ende.month - ps.month + 1) // 12))
    return [eingang + timedelta(days=7), eingang, werktag(ende, -10), ende - timedelta(days=30), ab,
            werktag(ab, 2), neu, ab + timedelta(days=30), edate(ps, 12 * n), edate(ende, 12)]


def mietenden():
    """Wirksames Mietende je ID: manuell, sonst aus dem Mieterwechsel."""
    aus_wechsel = {w[0]: (w[4] or mietende(w[2], w[3])) for w in B.B_WECHSEL}
    return {i + 1: (m[3] or aus_wechsel.get(i + 1)) for i, m in enumerate(B.B_MIETER)}


def nebenkosten(start=date(2025, 1, 1), akonto_override=None):
    ende_p = edate(start, 12) - timedelta(days=1)
    tage = (ende_p - start).days + 1
    einh = {nr: {"Fläche": fl, "Anteil": an, "Einheiten": 1} for nr, _, _, fl, an in B.B_EINHEITEN}
    tot = {k: sum(e[k] for e in einh.values()) for k in ("Fläche", "Anteil", "Einheiten")}
    kosten = [list(k) for k in B.B_KOSTEN]
    kosten[-1][1] = round(sum(k[1] for k in kosten[:-1]) * 0.03, 2)
    monate = [edate(start, j) for j in range(12)]
    gew = [B.GRADTAGE[m.month - 1] for m in monate]
    enden = mietenden()
    erg = {}
    for i, (nr, name, beg, _, _, akonto, *_r) in enumerate(B.B_MIETER):
        idn = i + 1
        b, e = max(beg or start, start), min(enden[idn] or ende_p, ende_p)
        t = max(0, (e - b).days + 1)
        if t == 0:
            erg[idn] = (0, 0, 0)
            continue
        beleg = []
        for m in monate:
            me = date(m.year, m.month, calendar.monthrange(m.year, m.month)[1])
            beleg.append(max(0, (min(e, me) - max(b, m)).days + 1) / me.day)
        lin, heiz = t / tage, sum(x * w for x, w in zip(beleg, gew)) / sum(B.GRADTAGE)
        s = sum(betr * einh[nr][schl] * (heiz if vert == "Heizung" else lin) / tot[schl]
                for _, betr, schl, vert in kosten)
        total = round(s * 20) / 20
        ak = akonto_override.get(idn) if akonto_override and idn in akonto_override else akonto * sum(beleg)
        erg[idn] = (total, round(ak, 2), round(total - ak, 2))
    return erg


def mietzins(netto, r0, r1, lik0, lik1, basis, stichtag, kosten_pa, mitteilung, frist=3):
    n = round((r1 - r0) / 0.25, 6)
    ref = 0.03 * n if n >= 0 else -(1 - 1 / (1 + 0.03 * -n))
    teu = (lik1 / lik0 - 1) * 0.4 if lik0 and lik1 else 0
    kos = kosten_pa / 100 * (stichtag - basis).days / 365.25 if kosten_pa else 0
    tot = ref + teu + kos
    neu = round(netto * (1 + tot) * 20) / 20
    termin = mietende(mitteilung + timedelta(days=10), frist) + timedelta(days=1) if tot > 0 and mitteilung else None
    return n, ref, teu, kos, tot, neu, termin


def rueckgabe(id_, maengel, kaution, zins, miete, nk, weitere):
    ende = mietenden()[id_]
    schaden = 0
    for m in maengel:
        if m[0] != id_ or m[5] is None or m[4] != "Mieter":
            continue
        rest = 1 if m[7] is None else max(0, 1 - ((ende - m[6]).days / 365.25) / m[7])
        schaden += round(m[5] * rest * 20) / 20
    ford = schaden + miete + nk + weitere
    k = kaution + zins
    an_v = max(0, min(ford, k))
    return [round(x, 2) for x in (schaden, ford, k, an_v, k - an_v, max(0, ford - k), max(0, -ford))]


# ------------------------------------------------------------ Varianten
def variante(name, aenderung, t):
    wb = openpyxl.load_workbook(os.path.join(HIER, "Vermieter-Ordner-CH-Beispiel.xlsx"))
    if aenderung:
        aenderung(wb)
    p = os.path.join(t, f"{name}.xlsx")
    wb.save(p)
    return p


def v_erhoehung(wb):
    m = wb["Mietzins"]
    m["B7"].value, m["B15"].value, m["B16"].value = 6, 1.75, 108.0
    m["B17"].value, m["B18"].value, m["B19"].value = date(2026, 12, 1), None, date(2026, 11, 20)


def v_akonto(wb):
    wb["NK Verteilung"].cell(B.NV0 + 3, 13).value = 5000      # ID 4


FEHLERFAELLE = {
    "Überschneidung zweier Mieter": lambda wb: [wb["Mein Haus"].cell(B.MV0 + 6, c, v) for c, v in
                                              ((2, "1"), (3, "Doppelt Max"), (4, date(2024, 1, 1)))],
    "Mieter mit unbekannter Einheit": lambda wb: [wb["Mein Haus"].cell(B.MV0 + 6, c, v) for c, v in
                                                ((2, "9"), (3, "Irrtum Ida"), (4, date(2024, 1, 1)))],
    "Mieterwechsel ohne Mieter": lambda wb: [wb["Mieterwechsel"].cell(B.W0 + 5, c, v) for c, v in
                                           ((1, 40), (5, date(2026, 1, 5)))],
    "Manuelles Mietende widerspricht": lambda wb: wb["Mein Haus"].cell(B.MV0 + 1, 5, date(2025, 6, 30)),
    "Einbau nach Mietende": lambda wb: wb["Rückgabe"].cell(B.M0, 7, date(2026, 1, 1)),
    "Mietzins-Basis nicht in 0.25-Schritten": lambda wb: wb["Mein Haus"].cell(B.MV0 + 3, 11, 1.6),
    "Kostenart ohne Schlüssel": lambda wb: wb["Nebenkosten"].cell(B.K0 + 10, 1, "Lift") and
                                           [wb["Nebenkosten"].cell(B.K0 + 10, c, v) for c, v in ((2, 500), (4, "Linear"))],
}


def main():
    os.chdir(HIER)
    subprocess.run([sys.executable, "build.py"], check=True, capture_output=True)
    with tempfile.TemporaryDirectory() as t:
        pfade = [variante("beispiel", None, t), variante("erhoehung", v_erhoehung, t), variante("akonto", v_akonto, t)]
        pfade += [variante(f"f{i}", f, t) for i, f in enumerate(FEHLERFAELLE.values())]
        pfade.append(os.path.join(HIER, "Vermieter-Ordner-CH.xlsx"))
        out = os.path.join(t, "out")
        subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out] + pfade,
                       check=True, capture_output=True)
        lade = lambda n: openpyxl.load_workbook(os.path.join(out, n), data_only=True)
        wb = lade("beispiel.xlsx")

        print("── Mein Haus: wirksames Mietende (Datenfluss aus Mieterwechsel)")
        enden = mietenden()
        for i in range(len(B.B_MIETER)):
            ist = d(wb["Mein Haus"].cell(B.MV0 + i, 6).value) or None
            pruefe(ist == enden[i + 1], f"ID {i+1}: {ist} / soll {enden[i+1]}")

        print("── Nebenkosten (Mieter und Flächen aus Mein Haus)")
        ref = nebenkosten()
        v = wb["NK Verteilung"]
        for idn, (tot, ak, saldo) in ref.items():
            z = B.NV0 + idn - 1
            ist = (round(v.cell(z, 52).value, 2), round(v.cell(z, 53).value, 2), round(v.cell(z, 54).value, 2))
            pruefe(ist == (tot, ak, saldo), f"ID {idn}: Total/Akonto/Saldo {ist} / soll {(tot, ak, saldo)}")
        a = wb["NK Abrechnung"]
        pruefe(a["B5"].value == "Muster Anna" and round(a.cell(37, 8).value, 2) == ref[1][0],
               f"Abrechnungsblatt ID 1: {a['B5'].value}, Total {a.cell(37, 8).value}")

        print("── Mieterwechsel: Mietende und zehn Schritte")
        w = wb["Mieterwechsel"]
        for i, (idn, _, eing, fr, vend, neu, _) in enumerate(B.B_WECHSEL):
            z = B.W0 + i
            ende = vend or mietende(eing, fr)
            soll = schritte(eing, ende, neu, date(2025, 1, 1))
            ist = [d(w.cell(z, B.STEP0 + 2 * s).value) or None for s in range(10)]
            pruefe(d(w.cell(z, 8).value) == ende and ist == soll, f"Zeile {i+1} (ID {idn}): Ende {ende}, Schritte "
                   + ("stimmen" if ist == soll else f"IST {ist} SOLL {soll}"))
        pruefe(w.cell(B.W0 + 1, 10).value == 61 and round(w.cell(B.W0 + 1, 11).value, 2) == round(1720 * 12 / 365 * 61, 2),
               f"Leerstand Einheit 4: {w.cell(B.W0+1, 10).value} Tage, CHF {w.cell(B.W0+1, 11).value}")
        s = wb["Start"]
        offen = [d(w.cell(B.W0 + 2, B.STEP0 + 2 * k).value) for k in range(10)]
        naechste = min(x for x in offen if x)
        pruefe(d(s["A14"].value) == naechste and s["B14"].value == "Muster Anna",
               f"Start «Was steht an»: {d(s['A14'].value)} {s['B14'].value} – {s['C14'].value}")

        print("── Mietzins")
        for name, wbx, args in (
                ("Senkung (Beispiel ID 4)", wb, (2150, 1.75, 1.25, 106.2, 107.1, date(2023, 12, 1), date(2026, 10, 1), 0.5, date(2026, 10, 15))),
                ("Erhöhung (ID 6)", lade("erhoehung.xlsx"), (1750, 1.25, 1.75, 107.3, 108.0, date(2025, 12, 1), date(2026, 12, 1), None, date(2026, 11, 20)))):
            m = wbx["Mietzins"]
            n, rf, te, ko, tot, neu, termin = mietzins(*args)
            ist = [m[f"B{z}"].value for z in (23, 24, 25, 26, 27, 28)]
            ok = all(abs(a_ - b_) < 1e-9 for a_, b_ in zip(ist[:5], (n, rf, te, ko, tot))) and round(ist[5], 2) == neu
            ok = ok and (d(m["B32"].value) or None) == termin
            pruefe(ok, f"{name}: Anpassung {tot:+.4%}, neue Miete {neu:.2f}, Termin {termin}")

        print("── Rückgabe (Mietende aus Mieterwechsel, NK-Saldo aus Nebenkosten)")
        f = B.B_FALL
        soll = rueckgabe(5, B.B_MAENGEL, 4950, f["zins"], 0, ref[5][2], 0)
        r = wb["Rückgabe"]
        ist = [round(r.cell(B.F0, c).value, 2) for c in (16, 17, 18, 19, 20, 21, 22)]
        pruefe(ist == soll and round(r.cell(B.F0, 13).value, 2) == ref[5][2],
               f"Fall ID 5: Schaden {ist[0]}, Forderungen {ist[1]}, an Vermieter {ist[3]}, an Mieter {ist[4]}")
        ka, pr = wb["Kautionsabrechnung"], wb["Protokoll"]
        zeilen = [ka.cell(z, 1).value for z in range(10, 10 + len(B.B_MAENGEL))]
        pruefe(zeilen[0] == "Wohnzimmer · Wandanstrich" and zeilen[-1] == "Eingang · Schlüssel",
               f"Kautionsabrechnung listet {len([z for z in zeilen if z])} Mängel von ID 5")
        pruefe(pr["B5"].value == "Weber Tim" and pr.cell(14, 2).value == "Wandanstrich",
               f"Protokoll: {pr['B5'].value}, erste Position {pr.cell(14, 2).value}")
        fr = wb["Freigabe Kaution"]
        pruefe(round(fr["C10"].value, 2) == soll[4] and round(fr["C12"].value, 2) == soll[3],
               f"Freigabe: Mieter {fr['C10'].value}, Vermieter {fr['C12'].value}")

        print("── Akonto-Abweichung")
        ref_a = nebenkosten(akonto_override={4: 5000})
        va = lade("akonto.xlsx")["NK Verteilung"]
        pruefe(round(va.cell(B.NV0 + 3, 54).value, 2) == ref_a[4][2], f"ID 4 Saldo mit Akonto 5000: {va.cell(B.NV0+3, 54).value}")

        print("── Kontrolle")
        pruefe(wb["Kontrolle"]["B16"].value == "OK", f"Beispiel: {wb['Kontrolle']['B16'].value}")
        leer = lade("Vermieter-Ordner-CH.xlsx")
        werte = [c.value for ws in leer for row in ws.iter_rows() for c in row
                 if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
        pruefe(not werte and leer["Kontrolle"]["B16"].value == "OK", f"Leere Vorlage: {len(werte)} Fehlerwerte, Status OK")
        for i, name in enumerate(FEHLERFAELLE):
            st = lade(f"f{i}.xlsx")["Kontrolle"]["B16"].value
            pruefe(st == "Bitte prüfen", f"Fehlerfall erkannt: {name} -> «{st}»")
        for n in ["beispiel.xlsx", "erhoehung.xlsx", "akonto.xlsx"] + [f"f{i}.xlsx" for i in range(len(FEHLERFAELLE))]:
            x = lade(n)
            werte = [c.value for ws in x for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith(("#", "Err:"))]
            if werte:
                pruefe(False, f"{n}: Fehlerwerte {werte[:3]}")
    print("FEHLER:", fehler)
    sys.exit(1 if fehler else 0)


if __name__ == "__main__":
    main()
