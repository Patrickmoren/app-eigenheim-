# -*- coding: utf-8 -*-
"""Erzeugt das Finanzmodell «bestanden» (3 Szenarien, 2027–2031) als Excel.
Eingaben stehen im Blatt «Annahmen» (blaue Schrift), alle Ergebnisse sind Formeln.
Aufruf: python3 build.py  →  Finanzmodell-bestanden.xlsx"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gl
from annahmen import A as EING, JAHRE, SZENARIEN, LERNENDE_CH, QV_KANDIDATEN_CH

DATEI = os.environ.get("ZIEL", "Finanzmodell-bestanden.xlsx")
ORANGE, TINTE, BLAU, HELL = "C2410C", "1A1A1A", "1F4FB4", "FFF1E8"
def f(g=10, b=False, c=TINTE, i=False):
    return Font(name="Arial", size=g, bold=b, color=c, italic=i)
def band(ws, r, c1, c2, farbe=ORANGE):
    for c in range(c1, c2 + 1):
        ws.cell(r, c).fill = PatternFill("solid", fgColor=farbe)

# Spalten je Szenario im Blatt Annahmen: Vorsichtig C–G, Basis I–M, Wachstum O–S
SPA = {s: [gl(3 + 6 * k + i) for i in range(5)] for k, s in enumerate(SZENARIEN)}
wb = openpyxl.Workbook()
An = wb.active; An.title = "Annahmen"
An.column_dimensions["A"].width = 50; An.column_dimensions["B"].width = 11
for s in SZENARIEN:
    for c in SPA[s]:
        An.column_dimensions[c].width = 9.5
An.column_dimensions["U"].width = 90
An["A1"] = "bestanden – Finanzmodell · Annahmen"; An["A1"].font = f(14, True, ORANGE)
An["A2"] = ("Blaue Werte sind Eingaben, alles andere ist berechnet. Einzelwerte stehen im ersten Jahr des Szenarios "
            "und gelten für alle Jahre. CHF ohne MWST. Hypothesen sind in Spalte U so bezeichnet.")
An["A2"].font = f(9, i=True)
for k, s in enumerate(SZENARIEN):
    c = SPA[s][0]; An[f"{c}4"] = s; An[f"{c}4"].font = f(11, True, ORANGE)
    for i, j in enumerate(JAHRE):
        An[f"{SPA[s][i]}5"] = j; An[f"{SPA[s][i]}5"].font = f(9, True)
An["U5"] = "Quelle / Begründung"; An["U5"].font = f(9, True)

R, SKALAR = {}, set()
r = 6
for e in EING:
    if e[0] == "_t":
        An[f"A{r}"] = e[1]; An[f"A{r}"].font = f(10, True, "FFFFFF"); band(An, r, 1, 21); r += 1; continue
    key, text, einheit, werte, fmt, quelle = e
    R[key] = r
    An[f"A{r}"] = text; An[f"B{r}"] = einheit; An[f"B{r}"].font = f(8, c="555555")
    for s in SZENARIEN:
        w = werte[s]
        if not isinstance(w, list):
            SKALAR.add(key); w = [w]
        for i, x in enumerate(w):
            c = An[f"{SPA[s][i]}{r}"]; c.value = x; c.font = f(10, c=BLAU); c.number_format = fmt
            c.fill = PatternFill("solid", fgColor="FFFBEA")
    An[f"U{r}"] = quelle; An[f"U{r}"].font = f(8, c="555555")
    An[f"U{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
An.freeze_panes = "C6"

def a(key, s, i):
    col = SPA[s][0] if key in SKALAR else SPA[s][i]
    return f"Annahmen!${col}${R[key]}"

# ------------------------------------------------------------- Modellblätter
SP = ["C", "D", "E", "F", "G"]
ZEILEN = [   # Schlüssel, Text, Formel(i, a, m, mp) , Format, Stil
  ("_t", "Mengen"),
  ("lern", "Lernende in abgedeckten Berufen", lambda i, a, m, mp: f"={a('lern', i)}", "#,##0", ""),
  ("qv", "QV-Kandidaten pro Jahr", lambda i, a, m, mp: f"={m('lern')}*{a('qv', i)}", "#,##0", ""),
  ("b2bl", "Lernende mit Betriebslizenz", lambda i, a, m, mp: f"={m('lern')}*{a('pen', i)}", "#,##0", "fett"),
  ("betr", "Lehrbetriebe mit Lizenz", lambda i, a, m, mp: f"={m('b2bl')}/{a('lpb', i)}", "#,##0", ""),
  ("neub", "davon neu gewonnen", lambda i, a, m, mp:
      f"={m('betr')}" if i == 0 else f"=MAX(0,{m('betr')}-{mp('betr')}*(1-{a('churn', i)}))", "#,##0", ""),
  ("cbas", "QV-Kandidaten ohne Betriebslizenz", lambda i, a, m, mp: f"={m('qv')}*(1-{a('pen', i)})", "#,##0", ""),
  ("ck", "B2C-Käufer/innen QV-Pass", lambda i, a, m, mp: f"={m('cbas')}*{a('conv', i)}", "#,##0", "fett"),
  ("oda", "Verbandslizenzen", lambda i, a, m, mp: f"={a('oda', i)}", "0", ""),
  ("akt", "Aktive Nutzer/innen (bezahlt)", lambda i, a, m, mp: f"={m('b2bl')}*{a('aktiv', i)}+{m('ck')}", "#,##0", ""),
  ("_t", "Umsatz"),
  ("ub", "Betriebslizenzen (B2B)", lambda i, a, m, mp: f"={m('b2bl')}*{a('preis_b', i)}", "#,##0", ""),
  ("uc", "QV-Pass (B2C)", lambda i, a, m, mp: f"={m('ck')}*{a('preis_c', i)}", "#,##0", ""),
  ("uo", "Verbandslizenzen (B2B2C)", lambda i, a, m, mp: f"={m('oda')}*{a('preis_oda', i)}", "#,##0", ""),
  ("u", "Umsatz total", lambda i, a, m, mp: f"={m('ub')}+{m('uc')}+{m('uo')}", "#,##0", "summe"),
  ("anteil_b", "Anteil B2B inkl. Verbände", lambda i, a, m, mp: f"=IF({m('u')}=0,0,({m('ub')}+{m('uo')})/{m('u')})", "0%", ""),
  ("_t", "Variable Kosten und Deckungsbeitrag"),
  ("vz", "Zahlungsgebühren", lambda i, a, m, mp: f"={m('uc')}*{a('zahl_c', i)}+({m('ub')}+{m('uo')})*{a('zahl_b', i)}", "#,##0", ""),
  ("vk", "KI-Kosten", lambda i, a, m, mp: f"={m('akt')}*{a('ki', i)}", "#,##0", ""),
  ("vh", "Hosting variabel", lambda i, a, m, mp: f"={m('akt')}*{a('host', i)}", "#,##0", ""),
  ("v", "Variable Kosten total", lambda i, a, m, mp: f"={m('vz')}+{m('vk')}+{m('vh')}", "#,##0", ""),
  ("db", "Deckungsbeitrag", lambda i, a, m, mp: f"={m('u')}-{m('v')}", "#,##0", "summe"),
  ("dbq", "Deckungsbeitrag in % des Umsatzes", lambda i, a, m, mp: f"=IF({m('u')}=0,0,{m('db')}/{m('u')})", "0%", ""),
  ("_t", "Fixkosten"),
  ("kp", "Personal", lambda i, a, m, mp: f"={a('fte', i)}*{a('k_fte', i)}", "#,##0", ""),
  ("kc", "Inhalte (neue Berufe + Pflege)", lambda i, a, m, mp:
      f"={a('berufe', i)}*{a('c_neu', i)}" if i == 0 else
      f"=MAX(0,{a('berufe', i)}-{a('berufe', i - 1)})*{a('c_neu', i)}+{a('berufe', i - 1)}*{a('c_pfl', i)}", "#,##0", ""),
  ("kt", "Technologie fix", lambda i, a, m, mp: f"={a('tech', i)}", "#,##0", ""),
  ("km", "Marketing und Vertrieb", lambda i, a, m, mp:
      f"={m('ck')}*(1-{a('org', i)})*{a('cac_c', i)}+{m('neub')}*{a('cac_b', i)}+{a('mk_fix', i)}", "#,##0", ""),
  ("ka", "Administration", lambda i, a, m, mp: f"={a('admin', i)}", "#,##0", ""),
  ("k", "Fixkosten total", lambda i, a, m, mp: f"={m('kp')}+{m('kc')}+{m('kt')}+{m('km')}+{m('ka')}", "#,##0", ""),
  ("ebit", "Operatives Ergebnis", lambda i, a, m, mp: f"={m('db')}-{m('k')}", "#,##0", "summe"),
  ("marge", "Operative Marge", lambda i, a, m, mp: f"=IF({m('u')}=0,0,{m('ebit')}/{m('u')})", "0%", ""),
  ("_t", "Kasse (ohne Steuern, ohne Fremdkapital)"),
  ("kasse", "Kasse Ende Jahr", lambda i, a, m, mp:
      f"={a('kasse0', i)}+{m('ebit')}" if i == 0 else f"={mp('kasse')}+{m('ebit')}", "#,##0", "fett"),
  ("bedarf", "Finanzierungsbedarf bis hier (tiefster Kassenstand)", lambda i, a, m, mp:
      f"=MAX(0,-{m('kasse')})" if i == 0 else f"=MAX({mp('bedarf')},-{m('kasse')})", "#,##0", ""),
  ("_t", "Kennzahlen"),
  ("mant", "Marktanteil an allen Lernenden CH (Betriebslizenz)", lambda i, a, m, mp: f"={m('b2bl')}/{LERNENDE_CH}", "0.0%", ""),
  ("ltv", "LTV je Lehrbetrieb (Deckungsbeitrag über Lebensdauer)", lambda i, a, m, mp:
      f"={a('lpb', i)}*{a('preis_b', i)}*(1-{a('zahl_b', i)})/{a('churn', i)}-{a('lpb', i)}*{a('aktiv', i)}*({a('ki', i)}+{a('host', i)})/{a('churn', i)}", "#,##0", ""),
  ("ltvcac", "LTV : Akquisekosten Lehrbetrieb", lambda i, a, m, mp: f"={m('ltv')}/{a('cac_b', i)}", "0.0", ""),
  ("dbc", "Deckungsbeitrag je B2C-Käufer/in", lambda i, a, m, mp:
      f"={a('preis_c', i)}*(1-{a('zahl_c', i)})-{a('ki', i)}-{a('host', i)}", "#,##0.00", ""),
  ("cacc", "Ø bezahlte Akquise je B2C-Käufer/in (inkl. Gratiskanäle)", lambda i, a, m, mp:
      f"={a('cac_c', i)}*(1-{a('org', i)})", "#,##0.00", ""),
]
MZ = {}
for s in SZENARIEN:
    M = wb.create_sheet(f"Modell {s}")
    M.column_dimensions["A"].width = 56
    for c in SP: M.column_dimensions[c].width = 13
    M["A1"] = f"bestanden – Modell · Szenario {s}"; M["A1"].font = f(14, True, ORANGE)
    M["A2"] = "Berechnet aus «Annahmen». Nicht von Hand ändern."; M["A2"].font = f(9, i=True)
    rows, r = {}, 4
    for z in ZEILEN:
        if z[0] != "_t": rows[z[0]] = None
    # Zeilennummern zuerst vergeben (Vorwärtsbezüge)
    for z in ZEILEN:
        if z[0] == "_t": r += 1; continue
        rows[z[0]] = r; r += 1
    MZ[s] = rows
    r = 4
    for z in ZEILEN:
        if z[0] == "_t":
            M[f"A{r}"] = z[1]; M[f"A{r}"].font = f(10, True, "FFFFFF"); band(M, r, 1, 7)
            for i, j in enumerate(JAHRE):
                M[f"{SP[i]}{r}"] = j; M[f"{SP[i]}{r}"].font = f(10, True, "FFFFFF")
            r += 1; continue
        key, text, formel, fmt, stil = z
        M[f"A{r}"] = text; M[f"A{r}"].font = f(10, stil in ("fett", "summe"))
        for i in range(5):
            c = M[f"{SP[i]}{r}"]
            c.value = formel(i, lambda k, i2, s=s: a(k, s, i2),
                             lambda k, i=i: f"{SP[i]}{rows[k]}",
                             lambda k, i=i: f"{SP[i - 1]}{rows[k]}")
            c.number_format = fmt; c.font = f(10, stil in ("fett", "summe"))
            if stil == "summe":
                c.fill = PatternFill("solid", fgColor=HELL); c.border = Border(top=Side(style="thin", color=TINTE))
        r += 1
    M.freeze_panes = "C4"

# ------------------------------------------------------------- Übersicht
U = wb.create_sheet("Übersicht", 0)
U.column_dimensions["A"].width = 46; U.column_dimensions["B"].width = 12
for c in SP: U.column_dimensions[c].width = 13
U["A1"] = "bestanden – Übersicht der drei Szenarien"; U["A1"].font = f(14, True, ORANGE)
U["A2"] = "CHF, ohne MWST und Steuern. Alle Werte sind Formeln auf die Modellblätter."; U["A2"].font = f(9, i=True)
ZEIGE = [("lern", "Lernende abgedeckt", "#,##0"), ("b2bl", "Lernende mit Betriebslizenz", "#,##0"),
         ("betr", "Lehrbetriebe", "#,##0"), ("ck", "B2C-Käufer/innen", "#,##0"), ("u", "Umsatz", "#,##0"),
         ("anteil_b", "Anteil B2B", "0%"), ("db", "Deckungsbeitrag", "#,##0"), ("kp", "Personal", "#,##0"),
         ("km", "Marketing und Vertrieb", "#,##0"), ("ebit", "Operatives Ergebnis", "#,##0"),
         ("kasse", "Kasse Ende Jahr", "#,##0"), ("bedarf", "Finanzierungsbedarf (kumuliert)", "#,##0"),
         ("mant", "Marktanteil Lernende CH", "0.0%")]
r = 4
UZ = {}
for s in SZENARIEN:
    U[f"A{r}"] = s; U[f"A{r}"].font = f(11, True, "FFFFFF"); band(U, r, 1, 7)
    for i, j in enumerate(JAHRE):
        U[f"{SP[i]}{r}"] = j; U[f"{SP[i]}{r}"].font = f(10, True, "FFFFFF")
    r += 1
    for key, text, fmt in ZEIGE:
        U[f"A{r}"] = text; UZ[(s, key)] = r
        for i in range(5):
            c = U[f"{SP[i]}{r}"]; c.value = f"='Modell {s}'!{SP[i]}{MZ[s][key]}"; c.number_format = fmt
            if key in ("u", "ebit"): c.font = f(10, True)
        r += 1
    e = MZ[s]["ebit"]; ms = f"'Modell {s}'"
    U[f"A{r}"] = "Dauerhaft positives Ergebnis ab"; U[f"A{r}"].font = f(10, True); UZ[(s, "be")] = r
    U[f"C{r}"] = (f"=IF({ms}!G{e}>0,IF({ms}!F{e}>0,IF({ms}!E{e}>0,IF({ms}!D{e}>0,IF({ms}!C{e}>0,2027,2028),"
                  f"2029),2030),2031),\"nach 2031\")")
    r += 2

# ------------------------------------------------------------- Zielrechnung
Z = wb.create_sheet("Zielrechnung CHF 7–10 Mio.")
Z.column_dimensions["A"].width = 70; Z.column_dimensions["B"].width = 16; Z.column_dimensions["C"].width = 16
Z.column_dimensions["D"].width = 60
Z["A1"] = "Was braucht es für CHF 7,5 Mio. oder CHF 10 Mio. Umsatz – nur Schweiz?"; Z["A1"].font = f(14, True, ORANGE)
Z["B3"] = "CHF 7,5 Mio."; Z["C3"] = "CHF 10 Mio."
for c in ("B3", "C3"): Z[c].font = f(10, True)
def zz(r, text, b, c, fmt="#,##0", blau=False, note=""):
    Z[f"A{r}"] = text; Z[f"B{r}"] = b; Z[f"C{r}"] = c; Z[f"D{r}"] = note
    Z[f"D{r}"].font = f(8, c="555555")
    for k in ("B", "C"):
        Z[f"{k}{r}"].number_format = fmt
        if blau: Z[f"{k}{r}"].font = f(10, c=BLAU)
zz(4, "Umsatzziel", 7500000, 10000000, blau=True)
zz(5, "Lernende in der Schweiz (alle Berufe)", LERNENDE_CH, LERNENDE_CH, blau=True, note="BFS, Lehrverhältnisse 2025")
zz(6, "QV-Kandidaten pro Jahr (alle Berufe)", QV_KANDIDATEN_CH, QV_KANDIDATEN_CH, blau=True, note="EHB Trendbericht 6, 2024")
zz(7, "Ø Preis Betriebslizenz je Lernende/r und Jahr", 110, 110, blau=True)
zz(8, "Preis QV-Pass", 59, 59, blau=True)
zz(10, "Nur Betriebslizenzen: nötige lizenzierte Lernende", "=B4/B7", "=C4/C7")
zz(11, "… in % aller Lernenden der Schweiz", "=B10/B5", "=C10/C5", "0%")
zz(12, "Nur QV-Pass: nötige Käufer/innen pro Jahr", "=B4/B8", "=C4/C8")
zz(13, "… in % aller QV-Kandidaten (über 100 % = unmöglich)", "=B12/B6", "=C12/C6", "0%")
Z["A15"] = "Obere plausible Grenze mit heutigem Modell (Annahmen blau, Hypothesen)"; Z["A15"].font = f(10, True)
zz(16, "Anteil aller Lernenden mit Betriebslizenz", 0.25, 0.25, "0%", True, "Hypothese: jede/r vierte Lernende der Schweiz")
zz(17, "Kaufquote QV-Pass unter den übrigen Kandidaten", 0.08, 0.08, "0%", True, "Hypothese")
zz(18, "Verbandslizenzen × CHF 40'000", 20, 20, "0", True, "Hypothese: 20 Verbände")
zz(19, "Umsatz Betriebslizenzen", "=B5*B16*B7", "=C5*C16*C7")
zz(20, "Umsatz QV-Pass", "=B6*(1-B16)*B17*B8", "=C6*(1-C16)*C17*C8")
zz(21, "Umsatz Verbände", "=B18*40000", "=C18*40000")
zz(22, "Umsatz total bei oberer Grenze", "=B19+B20+B21", "=C19+C20+C21")
zz(23, "Lücke zum Ziel (muss aus neuen Segmenten kommen)", "=MAX(0,B4-B22)", "=MAX(0,C4-C22)")
for c in ("A22", "B22", "C22", "A23", "B23", "C23"): Z[c].font = f(10, True)
Z["A25"] = ("Mögliche Segmente für die Lücke (nur Schweiz, alle zu prüfen): Erwachsene mit Berufsabschluss "
            "(Art. 32 BBV), Berufsprüfungen und höhere Fachprüfungen, üK-Zentren und Berufsfachschulen als Kunden, "
            "Prüfungs- und Aufgabenplattform für Verbände (eQV).")
Z["A25"].alignment = Alignment(wrap_text=True); Z.row_dimensions[25].height = 45
wb.save(DATEI)
print("geschrieben:", DATEI)
